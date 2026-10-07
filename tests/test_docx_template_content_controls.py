"""Content controls through a registered dynamic DOCX tool, end to end.

Covers the wiring in ``_sync_impl`` (conditionals first, then controls, then
placeholders) and the ``type: list`` + ``enum`` multi-choice argument that a
group of check boxes binds to.
"""
import asyncio
import io
import json
import sys
from pathlib import Path
from unittest.mock import patch

project_root = Path(__file__).parent.parent
sys.path.insert(0, str(project_root))

import pytest  # noqa: E402
from docx import Document  # noqa: E402
from fastmcp import Client, FastMCP  # noqa: E402

import docx_tools.dynamic_docx_tools as dd  # noqa: E402

from . import docx_content_control_helpers as cc  # noqa: E402

W = cc.W
SPEC_ARGS = [
    {"name": "full_name", "type": "string", "required": False, "description": "Name"},
    {"name": "size", "type": "string", "enum": ["small", "medium"], "required": False,
     "description": "Size"},
    {"name": "channels", "type": "list", "enum": ["email", "phone"],
     "required": False, "description": "Channels"},
    {"name": "has_agent", "type": "bool", "required": False, "default": False,
     "description": "Has an agent"},
]


@pytest.fixture
def tool(tmp_path):
    doc = Document()
    p = doc.add_paragraph("Name: ")
    cc.add_inline(p, cc.text("full_name", sdt_id=1))
    p2 = doc.add_paragraph()
    cc.add_inline(p2, cc.checkbox("size=small", sdt_id=2))
    cc.add_inline(p2, cc.checkbox("size=medium", sdt_id=3))
    cc.add_inline(p2, cc.checkbox("channels=email", sdt_id=4))
    cc.add_inline(p2, cc.checkbox("channels=phone", sdt_id=5))
    doc.add_paragraph("{{#if has_agent}}")
    cc.add_inline(doc.add_paragraph(), cc.text("full_name", sdt_id=6, prompt="inside the condition"))
    doc.add_paragraph("{{/if}}")
    doc.add_paragraph("Placeholder: {{full_name}}")
    template = tmp_path / "t.docx"
    doc.save(template)

    out = {}

    def fake_upload(buf, kind, filename=None, add_unique_prefix=None, **kw):
        out["doc"] = Document(io.BytesIO(buf.getvalue()))
        return "ok"

    mcp = FastMCP("t")
    spec = {"name": "cc_probe", "description": "probe", "docx_path": "t.docx", "args": SPEC_ARGS}
    with patch.object(dd, "find_docx_template_by_name", lambda f: str(template)), \
            patch.object(dd, "upload_file", fake_upload):
        assert dd._register_single_template(mcp, spec)
        yield mcp, out


def _call(mcp, data):
    async def go():
        async with Client(mcp) as c:
            return await c.call_tool("cc_probe", {"data": data})
    return asyncio.run(go())


def _schema(mcp):
    async def go():
        async with Client(mcp) as c:
            return (await c.list_tools())[0].inputSchema
    return asyncio.run(go())


def _controls(doc):
    return list(doc.element.body.iter(f"{{{W}}}sdt"))


def test_list_with_enum_is_published_as_an_array_of_choices(tool):
    mcp, _ = tool
    prop = _schema(mcp)["properties"]["data"]["properties"]["channels"]
    assert prop["type"] == "array"
    assert prop["items"]["enum"] == ["email", "phone"]
    assert prop["description"] == "Channels"
    assert "anyOf" not in json.dumps(prop)


def test_a_list_choice_outside_the_enum_is_refused(tool):
    mcp, _ = tool
    with pytest.raises(Exception):
        _call(mcp, {"channels": ["nonexistent"]})


def test_controls_are_filled_and_placeholders_still_work(tool):
    mcp, out = tool
    _call(mcp, {"full_name": "Jan", "size": "medium", "channels": ["email", "phone"]})
    doc = out["doc"]
    controls = _controls(doc)
    assert cc.content_text(controls[0]) == "Jan"
    assert [cc.checked(b)[0] for b in controls[1:5]] == [False, True, True, True]
    assert "Placeholder: Jan" in [p.text for p in doc.paragraphs]


def test_a_control_in_a_pruned_block_is_gone_not_filled(tool):
    mcp, out = tool
    _call(mcp, {"full_name": "Jan"})
    assert len(_controls(out["doc"])) == 5
    assert "inside the condition" not in out["doc"].element.xml


def test_unsent_arguments_leave_the_controls_as_the_template_has_them(tool):
    mcp, out = tool
    _call(mcp, {})
    controls = _controls(out["doc"])
    assert cc.content_text(controls[0]) == "Click here."
    assert [cc.checked(b)[0] for b in controls[1:5]] == [False] * 4


def test_a_plain_list_without_enum_is_unchanged():
    mcp = FastMCP("t2")
    spec = {"name": "plain_list", "docx_path": "x.docx",
            "args": [{"name": "tags", "type": "list", "required": False, "description": "d"}]}
    with patch.object(dd, "find_docx_template_by_name", lambda f: "/tmp/x.docx"):
        assert dd._register_single_template(mcp, spec)

    async def go():
        async with Client(mcp) as c:
            return (await c.list_tools())[0].inputSchema
    prop = asyncio.run(go())["properties"]["data"]["properties"]["tags"]
    assert prop["type"] == "array" and prop["items"] == {"type": "string"}
