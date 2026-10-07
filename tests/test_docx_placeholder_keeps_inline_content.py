"""Replacing a placeholder rebuilds only the runs it spans (PR #200 review).

It used to remove every run of the paragraph and append the rebuilt text at
its end, so anything in the paragraph that is not a run — a content control,
a hyperlink — moved to the front. A form line like ``Name: {{name}} [box]``
came out as ``[box] Name: Jane``.
"""
import sys
from pathlib import Path

project_root = Path(__file__).parent.parent
sys.path.insert(0, str(project_root))

import pytest  # noqa: E402
from docx import Document  # noqa: E402

from docx_tools.dynamic_docx_tools import _replace_placeholders_in_document  # noqa: E402

from . import docx_content_control_helpers as cc  # noqa: E402

W = cc.W


def _sequence(element):
    """Visible text in document order, with each content control as ``[CC]``."""
    out = []
    for el in element.iter(f"{{{W}}}t", f"{{{W}}}sdt"):
        if el.tag == f"{{{W}}}sdt":
            out.append("[CC]")
        elif not any(a.tag == f"{{{W}}}sdt" for a in el.iterancestors()):
            out.append(el.text)  # the box's own glyph is part of [CC]
    return "".join(out)


@pytest.mark.parametrize("value", ["Jane", "**Jane** Doe", "Jane\\nDoe"])
def test_a_control_after_the_placeholder_stays_after_it(value):
    doc = Document()
    p = doc.add_paragraph("Name: {{name}}, consent ")
    cc.add_inline(p, cc.checkbox("consent"))
    p.add_run(" end")
    _replace_placeholders_in_document(doc, {"name": value})
    text = _sequence(doc.element.body)
    assert text.startswith("Name: ") and "[CC]" in text
    assert text.index("[CC]") > text.index("Jane")
    assert text.endswith("[CC] end")


def test_a_control_before_the_placeholder_stays_before_it():
    doc = Document()
    p = doc.add_paragraph("Box ")
    cc.add_inline(p, cc.checkbox("consent"))
    p.add_run(" then {{name}} end")
    _replace_placeholders_in_document(doc, {"name": "Jane"})
    assert _sequence(doc.element.body) == "Box [CC] then Jane end"


def test_two_placeholders_around_a_control_keep_the_order():
    doc = Document()
    p = doc.add_paragraph("{{a}} ")
    cc.add_inline(p, cc.checkbox("consent"))
    p.add_run(" {{b}}")
    _replace_placeholders_in_document(doc, {"a": "A", "b": "B"})
    assert _sequence(doc.element.body) == "A [CC] B"


def test_in_a_table_cell_too():
    doc = Document()
    p = doc.add_table(rows=1, cols=1).cell(0, 0).paragraphs[0]
    p.add_run("{{name}} ")
    cc.add_inline(p, cc.checkbox("consent"))
    _replace_placeholders_in_document(doc, {"name": "Jane"})
    assert _sequence(doc.tables[0]._tbl) == "Jane [CC]"


def test_a_block_value_beside_a_control_does_not_delete_the_paragraph():
    """A whole-paragraph block value removes the placeholder's paragraph; with a
    control in it, the paragraph is not 'whole' and must survive."""
    doc = Document()
    p = doc.add_paragraph("{{items}}")
    cc.add_inline(p, cc.checkbox("consent"))
    _replace_placeholders_in_document(doc, {"items": "- one\n- two"})
    text = _sequence(doc.element.body)
    assert "[CC]" in text and "one" in text and "two" in text


def test_text_around_the_placeholder_in_other_runs_keeps_its_formatting():
    doc = Document()
    p = doc.add_paragraph()
    p.add_run("Bold ").bold = True
    p.add_run("{{name}}")
    p.add_run(" italic").italic = True
    _replace_placeholders_in_document(doc, {"name": "Jane"})
    runs = doc.paragraphs[0].runs
    assert "".join(r.text for r in runs) == "Bold Jane italic"
    assert runs[0].bold and runs[-1].italic
