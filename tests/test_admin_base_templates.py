"""The Base templates page: the five files that style every generated document.

These are not dynamic templates — no spec, no name, exactly one of each — so
they get their own page and their own tests. The properties that matter are
that the page agrees with what the tools actually resolve, that replacing one
takes effect without a restart, and that reverting is offered only when there
is something to revert.
"""
import io
import re
import sys
from pathlib import Path

project_root = Path(__file__).parent.parent
sys.path.insert(0, str(project_root))

import pytest
from docx import Document
from openpyxl import Workbook
from openpyxl.styles import Font, NamedStyle
from starlette.testclient import TestClient

import admin.store as store_mod
import template_utils as tu
import metrics
from config import Config


@pytest.fixture
def admin_client(tmp_path, monkeypatch):
    custom = tmp_path / "custom_templates"
    cfg = tmp_path / "config"
    custom.mkdir()
    cfg.mkdir()
    monkeypatch.setattr(store_mod, "_APP_CUSTOM_DIR", tmp_path / "nx1")
    monkeypatch.setattr(store_mod, "_APP_CONFIG_DIR", tmp_path / "nx2")
    monkeypatch.setattr(store_mod, "_LOCAL_CUSTOM_DIR", custom)
    monkeypatch.setattr(store_mod, "_LOCAL_CONFIG_DIR", cfg)
    monkeypatch.setattr(tu, "APP_CUSTOM_DIR", tmp_path / "nx1")
    monkeypatch.setattr(tu, "LOCAL_CUSTOM_DIR", custom)
    monkeypatch.setenv("ADMIN_ENABLED", "true")
    monkeypatch.setenv("ADMIN_PASSWORD", "pw")
    monkeypatch.delenv("API_KEY", raising=False)

    from fastmcp import FastMCP
    from admin.app import build_combined_app

    client = TestClient(build_combined_app(FastMCP("t"), Config.from_env()))
    metrics.reset()
    client.__enter__()
    client.post("/admin/login", data={"password": "pw"})
    yield client, custom
    client.__exit__(None, None, None)
    metrics.reset()


def _csrf(client) -> str:
    html = client.get("/admin/base").text
    tag = re.search(r'<input[^>]*name="csrf"[^>]*>', html)
    return re.search(r'value="([^"]*)"', tag.group(0)).group(1)


def _post(client, url, data=None, files=None):
    payload = dict(data or {})
    payload["csrf"] = _csrf(client)
    return client.post(url, data=payload, files=files)


def _docx_bytes(text="hello"):
    doc = Document()
    doc.add_paragraph(text)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _xlsx_bytes(*style_names):
    wb = Workbook()
    for name in style_names:
        style = NamedStyle(name=name)
        style.font = Font(bold=True)
        wb.add_named_style(style)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_every_slot_is_listed(admin_client):
    from admin.base_templates import SLOTS

    client, _ = admin_client
    html = client.get("/admin/base").text
    for s in SLOTS:
        assert s.label in html, f"{s.key} missing from the page"
        assert f"/admin/base/{s.key}/upload" in html


def test_page_is_reachable_from_the_nav(admin_client):
    client, _ = admin_client
    nav = re.search(r"<nav>(.*?)</nav>", client.get("/admin/").text, re.S)
    assert nav and "/admin/base" in nav.group(1)


def test_download_route_is_not_swallowed_by_the_generic_pattern(admin_client):
    """/base/docx/download also fits /{kind}/{name}/download.

    Starlette matches in registration order, so the base routes are declared
    first. If that ordering is lost this redirects home instead of serving.
    """
    client, _ = admin_client
    r = client.get("/admin/base/docx/download", follow_redirects=False)
    assert r.status_code == 200, "the generic /{kind}/{name}/download won the match"
    assert r.content[:2] == b"PK"  # a real .docx, not a redirect body


def test_bundled_default_is_served_when_nothing_is_installed(admin_client):
    client, custom = admin_client
    assert not (custom / "custom_docx_template.docx").exists()
    html = client.get("/admin/base").text
    assert "Bundled default" in html
    r = client.get("/admin/base/docx/download")
    assert r.content == (project_root / "default_templates"
                         / "default_docx_template.docx").read_bytes()


def test_upload_installs_and_takes_effect_immediately(admin_client):
    client, custom = admin_client
    data = _docx_bytes("brand")
    r = _post(client, "/admin/base/docx/upload",
              files={"file": ("mine.docx", data, "application/octet-stream")})
    assert r.status_code == 200
    installed = custom / "custom_docx_template.docx"
    assert installed.read_bytes() == data
    # The tools resolve it without a restart.
    from template_utils import find_docx_template
    assert Path(find_docx_template()) == installed
    # And the page says so.
    assert "● Custom" in r.text


def test_revert_is_offered_only_once_something_is_installed(admin_client):
    client, _ = admin_client
    assert "/revert" not in client.get("/admin/base").text
    _post(client, "/admin/base/docx/upload",
          files={"file": ("m.docx", _docx_bytes(), "application/octet-stream")})
    assert "/admin/base/docx/revert" in client.get("/admin/base").text


def test_revert_restores_the_bundled_default(admin_client):
    client, custom = admin_client
    _post(client, "/admin/base/docx/upload",
          files={"file": ("m.docx", _docx_bytes(), "application/octet-stream")})
    r = _post(client, "/admin/base/docx/revert")
    assert not (custom / "custom_docx_template.docx").exists()
    assert "bundled default is in use again" in r.text
    from template_utils import find_docx_template
    assert "default_docx_template.docx" in find_docx_template()


def test_revert_without_a_custom_file_changes_nothing(admin_client):
    client, _ = admin_client
    r = _post(client, "/admin/base/docx/revert")
    assert "no custom file to remove" in r.text.lower()


def test_upload_rejects_the_wrong_extension(admin_client):
    client, custom = admin_client
    r = _post(client, "/admin/base/xlsx/upload",
              files={"file": ("nope.docx", _docx_bytes(), "application/octet-stream")})
    assert "expects a .xlsx file" in r.text
    assert not (custom / "custom_xlsx_template.xlsx").exists()


def test_upload_rejects_a_file_it_cannot_open(admin_client):
    client, custom = admin_client
    r = _post(client, "/admin/base/docx/upload",
              files={"file": ("broken.docx", b"not a docx", "application/octet-stream")})
    assert "could not open" in r.text.lower()
    assert not (custom / "custom_docx_template.docx").exists(), "a reject must not install"


def test_upload_requires_csrf(admin_client):
    client, custom = admin_client
    r = client.post("/admin/base/docx/upload",
                    files={"file": ("m.docx", _docx_bytes(), "application/octet-stream")})
    assert r.status_code == 403
    assert not (custom / "custom_docx_template.docx").exists()


def test_revert_requires_csrf(admin_client):
    client, custom = admin_client
    _post(client, "/admin/base/docx/upload",
          files={"file": ("m.docx", _docx_bytes(), "application/octet-stream")})
    r = client.post("/admin/base/docx/revert")
    assert r.status_code == 403
    assert (custom / "custom_docx_template.docx").exists()


def test_excel_slot_lists_the_named_styles(admin_client):
    """#170: the set of valid `style:<Name>` values had no way to be read."""
    client, _ = admin_client
    r = _post(client, "/admin/base/xlsx/upload",
              files={"file": ("styles.xlsx", _xlsx_bytes("Callout", "BrandTotal"),
                              "application/octet-stream")})
    assert r.status_code == 200
    assert "Callout" in r.text and "BrandTotal" in r.text
    # And the tool agrees with what the page showed.
    from xlsx_tools.styles import load_template_styles
    assert load_template_styles().names == {"Callout", "BrandTotal"}


def test_excel_has_no_bundled_default(admin_client):
    """Reverting Excel removes the feature; it does not restore a fallback."""
    client, custom = admin_client
    _post(client, "/admin/base/xlsx/upload",
          files={"file": ("s.xlsx", _xlsx_bytes("Callout"), "application/octet-stream")})
    r = _post(client, "/admin/base/xlsx/revert")
    assert not (custom / "custom_xlsx_template.xlsx").exists()
    assert "Nothing ships in its place" in r.text
    assert client.get("/admin/base/xlsx/download", follow_redirects=False).status_code == 200
    assert "Not found" in client.get("/admin/base/xlsx/download").text


def test_source_is_decided_by_identity_not_directory_name(tmp_path):
    """Whether a slot shows "Custom" must not hinge on a directory's name.

    template_utils._classify_template_source() answers this by looking for a
    path part literally called "custom_templates". That is right in the normal
    layout and wrong when the bundled defaults happen to sit *under* such a
    directory — it would report the default as custom, and the page would then
    offer to revert a file it cannot delete. source_of() compares against the
    exact path a replacement is written to instead.
    """
    from template_utils import _classify_template_source
    from admin.base_templates import slot, source_of

    s = slot("docx")
    custom_dir = tmp_path / "custom_templates"
    custom_dir.mkdir()

    # A bundled default that lives under a path containing "custom_templates".
    nested_default = custom_dir / "default_templates"
    nested_default.mkdir()
    default_file = nested_default / s.default_name
    default_file.write_bytes(b"PK-default")

    # The heuristic gets this wrong...
    assert _classify_template_source(default_file) == "custom"
    # ...and source_of does not.
    assert source_of(custom_dir, s, default_file) == "default"

    # The real replacement still reads as custom.
    installed = custom_dir / s.custom_name
    installed.write_bytes(b"PK-mine")
    assert source_of(custom_dir, s, installed) == "custom"
    assert source_of(custom_dir, s, None) == "none"


def test_unknown_slot_redirects_home(admin_client):
    client, _ = admin_client
    r = client.get("/admin/base/nonsense/download", follow_redirects=False)
    assert r.status_code == 303


def test_base_page_requires_authentication(tmp_path, monkeypatch):
    monkeypatch.setattr(store_mod, "_LOCAL_CUSTOM_DIR", tmp_path / "c")
    monkeypatch.setattr(store_mod, "_LOCAL_CONFIG_DIR", tmp_path / "cfg")
    monkeypatch.setenv("ADMIN_ENABLED", "true")
    monkeypatch.setenv("ADMIN_PASSWORD", "pw")
    from fastmcp import FastMCP
    from admin.app import build_combined_app
    with TestClient(build_combined_app(FastMCP("t"), Config.from_env())) as c:
        for path in ("/admin/base", "/admin/base/docx/download"):
            r = c.get(path, follow_redirects=False)
            assert r.status_code == 303
            assert r.headers["location"].endswith("/admin/login")
