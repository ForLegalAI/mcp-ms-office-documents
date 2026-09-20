"""What the builder draws itself takes its colour from the template.

Anything the builder draws rather than places in a placeholder — KPI figures,
timeline detail lines, a quote, blank-slide text, chart axis labels, table
fills — was pinned to a literal or left inheriting the presentation's default
text style (`tx1`, black). On the dark template in #195 the KPI figures came
out black on near-black, and every table header was Office's old default blue
whatever the template's own accent was.
"""

import sys
from pathlib import Path

project_root = Path(__file__).parent.parent
sys.path.insert(0, str(project_root))

import pytest
from pptx import Presentation as PptxReader
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_THEME_COLOR
from pptx.oxml.ns import qn

from pptx_tools.placeholder_style import read_body_color
from pptx_tools.slide_builder import PowerpointPresentation
from pptx_tools.templates import TemplateSpec

BASE_16_9 = project_root / "default_templates" / "default_pptx_template_16_9.pptx"


def _spec(path=BASE_16_9) -> TemplateSpec:
    return TemplateSpec(name="t", path=Path(path), description="", is_default=True,
                        layouts={}, defaults={}, strip_slides=True, aspect="16:9")


def _dark_template(tmp_path, scheme="bg1") -> Path:
    """The shipped template with its master body colour changed."""
    prs = PptxReader(str(BASE_16_9))
    bodyStyle = (prs.slide_masters[0]._element
                 .find(qn("p:txStyles")).find(qn("p:bodyStyle")))
    lvl1 = bodyStyle.find(qn("a:lvl1pPr"))
    solidFill = lvl1.find(qn("a:defRPr")).find(qn("a:solidFill"))
    solidFill.find(qn("a:schemeClr")).set("val", scheme)
    path = tmp_path / f"body_{scheme}.pptx"
    prs.save(str(path))
    return path


def _build(slides, path=BASE_16_9):
    deck = PowerpointPresentation(slides, "16:9", template_spec=_spec(path))
    return deck.presentation.slides[0]


def _run_colors(shape):
    return [run.font.color for paragraph in shape.text_frame.paragraphs
            for run in paragraph.runs]


def _box_with(slide, needle):
    """The shape whose text contains *needle*."""
    for shape in slide.shapes:
        if shape.has_text_frame and needle in shape.text_frame.text:
            return shape
    raise AssertionError(f"no shape contains {needle!r}")


# ---------------------------------------------------------------------------
# reading the colour
# ---------------------------------------------------------------------------

def test_the_shipped_template_body_colour_is_read():
    prs = PptxReader(str(BASE_16_9))
    assert read_body_color(prs.slide_masters[0]) == MSO_THEME_COLOR.TEXT_1


@pytest.mark.parametrize("scheme,expected", [
    ("bg1", MSO_THEME_COLOR.BACKGROUND_1),
    ("accent1", MSO_THEME_COLOR.ACCENT_1),
    ("lt1", MSO_THEME_COLOR.LIGHT_1),
])
def test_a_scheme_colour_stays_a_scheme_colour(tmp_path, scheme, expected):
    """Resolving it to RGB here would stop it following the theme."""
    prs = PptxReader(str(_dark_template(tmp_path, scheme)))
    assert read_body_color(prs.slide_masters[0]) == expected


# ---------------------------------------------------------------------------
# applying it
# ---------------------------------------------------------------------------

KPI = {"type": "kpi", "title": "t",
       "items": [{"value": "42", "label": "answers"}]}
TIMELINE = {"type": "timeline", "title": "t",
            "steps": [{"label": "Step", "detail": "two weeks"},
                      {"label": "Next", "detail": "four weeks"}]}
QUOTE = {"type": "quote", "title": "t", "text": "words", "attribution": "someone"}
BLANK = {"type": "blank", "title": "t",
         "elements": [{"kind": "text", "x": "10%", "y": "40%", "w": "50%",
                       "text": "positioned"}]}


# The text each type draws for itself. A timeline's chevron label and a blank
# slide's drawn title are deliberately left out: the label sits on an accent
# fill and takes its colour from the shape style, and the title is painted from
# the template's title style, not its body style.
@pytest.mark.parametrize("slide,needles", [
    (KPI, ["42", "answers"]),
    (TIMELINE, ["two weeks"]),
    (QUOTE, ["words", "someone"]),
    (BLANK, ["positioned"]),
], ids=["kpi", "timeline", "quote", "blank"])
def test_drawn_text_follows_the_template_body_colour(tmp_path, slide, needles):
    path = _dark_template(tmp_path, "bg1")
    built = _build([slide], path)

    for needle in needles:
        colors = _run_colors(_box_with(built, needle))
        assert colors, f"{needle!r} has no runs"
        assert all(color.theme_color == MSO_THEME_COLOR.BACKGROUND_1
                   for color in colors), f"{needle!r} was not painted"


def test_drawn_text_is_left_alone_when_the_template_states_no_colour(tmp_path):
    prs = PptxReader(str(BASE_16_9))
    bodyStyle = (prs.slide_masters[0]._element
                 .find(qn("p:txStyles")).find(qn("p:bodyStyle")))
    defRPr = bodyStyle.find(qn("a:lvl1pPr")).find(qn("a:defRPr"))
    defRPr.remove(defRPr.find(qn("a:solidFill")))
    path = tmp_path / "nocolour.pptx"
    prs.save(str(path))

    assert read_body_color(PptxReader(str(path)).slide_masters[0]) is None
    built = _build([KPI], path)          # must not raise, and must not paint
    for color in _run_colors(_box_with(built, "42")):
        assert color.type is None


def test_chart_text_follows_the_template_body_colour(tmp_path):
    path = _dark_template(tmp_path, "bg1")
    built = _build([{"type": "chart", "title": "t", "chart_type": "bar",
                     "categories": ["a", "b"],
                     "series": [{"name": "s", "values": [1, 2]}]}], path)

    charts = [shape.chart for shape in built.shapes if shape.has_chart]
    assert len(charts) == 1
    assert charts[0].font.color.theme_color == MSO_THEME_COLOR.BACKGROUND_1


def test_scatter_text_follows_the_template_body_colour(tmp_path):
    path = _dark_template(tmp_path, "bg1")
    built = _build([{"type": "scatter", "title": "t",
                     "series": [{"name": "s", "points": [[1, 2], [3, 4]]}]}], path)

    charts = [shape.chart for shape in built.shapes if shape.has_chart]
    assert len(charts) == 1
    assert charts[0].font.color.theme_color == MSO_THEME_COLOR.BACKGROUND_1


def test_table_header_uses_the_template_accent(tmp_path):
    built = _build([{"type": "table", "title": "t",
                     "rows": [["h1", "h2"], ["a", "b"], ["c", "d"]]}])
    table = next(shape.table for shape in built.shapes if shape.has_table)

    assert table.cell(0, 0).fill.fore_color.theme_color == MSO_THEME_COLOR.ACCENT_1
    assert '<a:schemeClr val="accent1"/>' in table.cell(0, 0)._tc.xml


def test_an_explicit_header_colour_still_wins():
    built = _build([{"type": "table", "title": "t", "header_color": "#FF0000",
                     "rows": [["h1", "h2"], ["a", "b"]]}])
    table = next(shape.table for shape in built.shapes if shape.has_table)
    assert table.cell(0, 0).fill.fore_color.rgb == RGBColor(0xFF, 0x00, 0x00)
