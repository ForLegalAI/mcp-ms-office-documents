"""Bullets drawn in a text box, not a placeholder (#123).

A text box inherits its paragraph formatting from the presentation's default
text style, which has no bullets — so the body beside a picture or a chart
rendered as plain lines, while the identical markdown bulleted correctly in a
content placeholder. The master's body style is now applied explicitly.
"""
import base64
import sys
from pathlib import Path

project_root = Path(__file__).parent.parent
sys.path.insert(0, str(project_root))

import pytest
from pptx import Presentation as PptxReader
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.oxml.ns import qn
from pptx.util import Pt

from pptx_tools.placeholder_style import apply_list_style, body_list_levels
from pptx_tools.slide_builder import PowerpointPresentation
from pptx_tools.templates import TemplateSpec

BASE_16_9 = project_root / "default_templates" / "default_pptx_template_16_9.pptx"
BULLET = "•"

PNG_DATA_URI = "data:image/png;base64," + base64.b64encode(base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8z8BQDwAEhQGAhKmMIQAAAABJRU5ErkJggg=="
)).decode()

IMAGE_SLIDE = {"type": "image", "title": "I", "source": PNG_DATA_URI,
               "body": "- one\n  - two", "layout": "Nadpis a obsah"}
CHART_SLIDE = {"type": "chart", "title": "C", "chart_type": "column",
               "categories": ["a"], "series": [{"name": "s", "values": [1]}],
               "body": "- one\n  - two"}


def build(slides, spec=None):
    pres = PowerpointPresentation(slides, "16:9", template_spec=spec)
    return pres, PptxReader(pres.save()).slides[0]


def body_box(slide):
    return [shape for shape in slide.shapes
            if shape.shape_type == MSO_SHAPE_TYPE.TEXT_BOX
            and "one" in shape.text_frame.text][0]


def properties(paragraph):
    return paragraph._p.find(qn('a:pPr'))


class TestTheGlyphsAreThere:

    @pytest.mark.parametrize("slide_data", [IMAGE_SLIDE, CHART_SLIDE])
    def test_bullets_beside_a_picture_or_a_chart(self, slide_data):
        _, slide = build([slide_data])
        first = body_box(slide).text_frame.paragraphs[0]

        assert properties(first).find(qn('a:buChar')).get('char') == BULLET
        assert properties(first).find(qn('a:buFont')) is not None

    def test_each_level_takes_its_own_indent(self):
        """Level 2 is indented further, as it is in a body placeholder."""
        _, slide = build([IMAGE_SLIDE])
        master = PptxReader(str(BASE_16_9)).slide_master
        levels = body_list_levels(master)

        paragraphs = body_box(slide).text_frame.paragraphs
        assert properties(paragraphs[0]).get('marL') == levels[0].get('marL')
        assert properties(paragraphs[1]).get('marL') == levels[1].get('marL')
        assert properties(paragraphs[0]).get('marL') != properties(paragraphs[1]).get('marL')

    def test_the_builder_keeps_its_own_sizing(self):
        """The body style's 28pt first level is for a full-width placeholder."""
        _, slide = build([IMAGE_SLIDE])
        pPr = properties(body_box(slide).text_frame.paragraphs[0])

        assert pPr.find(qn('a:defRPr')) is None

    def test_a_placeholder_body_is_left_to_inherit(self):
        """A content slide already bulleted; it must not gain explicit copies."""
        _, slide = build([{"type": "content", "title": "C", "body": "- one"}])
        placeholder = [s for s in slide.placeholders
                       if s.placeholder_format.idx == 1][0]
        pPr = properties(placeholder.text_frame.paragraphs[0])

        assert pPr is None or pPr.find(qn('a:buChar')) is None


class TestApplyingTheStyle:

    def test_it_keeps_the_children_of_pPr_in_schema_order(self):
        """buChar must precede defRPr, or PowerPoint calls the file damaged."""
        presentation = PptxReader(str(BASE_16_9))
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        frame = slide.shapes.add_textbox(0, 0, 100, 100).text_frame
        frame.text = "one"
        paragraph = frame.paragraphs[0]
        paragraph.font.size = Pt(18)           # writes <a:defRPr>
        paragraph.line_spacing = 1.0           # writes <a:lnSpc>

        apply_list_style(frame, presentation.slide_master)

        tags = [child.tag for child in properties(paragraph)]
        assert tags.index(qn('a:lnSpc')) < tags.index(qn('a:buChar'))
        assert tags.index(qn('a:buChar')) < tags.index(qn('a:defRPr'))
        assert properties(paragraph).find(qn('a:defRPr')).get('sz') == '1800'

    def test_an_existing_bullet_setting_is_replaced_not_doubled(self):
        presentation = PptxReader(str(BASE_16_9))
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        frame = slide.shapes.add_textbox(0, 0, 100, 100).text_frame
        frame.text = "one"
        pPr = frame.paragraphs[0]._p.get_or_add_pPr()
        pPr.append(pPr.makeelement(qn('a:buNone'), {}))

        apply_list_style(frame, presentation.slide_master)

        assert properties(frame.paragraphs[0]).find(qn('a:buNone')) is None
        assert len(properties(frame.paragraphs[0]).findall(qn('a:buChar'))) == 1

    def test_a_master_without_a_body_style_is_left_alone(self, tmp_path):
        presentation = PptxReader(str(BASE_16_9))
        txStyles = presentation.slide_master._element.find(qn('p:txStyles'))
        txStyles.remove(txStyles.find(qn('p:bodyStyle')))
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        frame = slide.shapes.add_textbox(0, 0, 100, 100).text_frame
        frame.text = "one"

        assert apply_list_style(frame, presentation.slide_master) is False
        assert body_list_levels(presentation.slide_master) == {}

    def test_a_deck_still_builds_on_a_template_without_a_body_style(self, tmp_path):
        presentation = PptxReader(str(BASE_16_9))
        txStyles = presentation.slide_master._element.find(qn('p:txStyles'))
        txStyles.remove(txStyles.find(qn('p:bodyStyle')))
        path = tmp_path / "no_body_style.pptx"
        presentation.save(str(path))

        pres, slide = build([IMAGE_SLIDE],
                            spec=TemplateSpec(name="bare", path=path, aspect="16:9"))

        assert "one" in body_box(slide).text_frame.text
        assert not [w for w in pres.warnings if w.severity == "error"]
