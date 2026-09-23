"""Column widths, cell fills and merged cells on a table slide (#124).

`TableSlide` used to offer alignment, a header colour, zebra shading and a
font size, so a two-character column got as much room as a description, a
risk table could not flag its own severity, and a grouped header was
impossible.
"""
import sys
from pathlib import Path

project_root = Path(__file__).parent.parent
sys.path.insert(0, str(project_root))

import pytest
from pptx import Presentation as PptxReader
from pptx.dml.color import MSO_THEME_COLOR

from pptx_tools import warnings as W
from pptx_tools.schema import coerce_slides
from pptx_tools.slide_builder import PowerpointPresentation

ROWS = [["Risk", "Impact", "Owner"],
        ["Data loss", "High", "Ops"],
        ["Latency", "Low", "Eng"]]


def build(**table):
    pres = PowerpointPresentation([{"type": "table", "title": "T", "rows": ROWS, **table}], "16:9")
    shape = [s for s in PptxReader(pres.save()).slides[0].shapes if s.has_table][0]
    return pres, shape.table


def codes(pres):
    return [warning.code for warning in pres.warnings]


class TestColumnWidths:

    def test_weights_become_proportions_of_the_table(self):
        _, table = build(widths=[3, 1, 1])
        widths = [column.width for column in table.columns]

        assert widths[0] == pytest.approx(widths[1] * 3, rel=0.001)
        assert widths[1] == pytest.approx(widths[2], rel=0.001)

    def test_the_columns_still_add_up_to_the_table(self):
        """Rounding three shares must not leave or take a sliver."""
        _, table = build(widths=[1, 1, 1])
        shape_width = sum(column.width for column in table.columns)

        _, equal = build()
        assert shape_width == sum(column.width for column in equal.columns)

    def test_any_scale_of_the_same_ratio_is_the_same_table(self):
        _, tenths = build(widths=[0.6, 0.2, 0.2])
        _, hundreds = build(widths=[60, 20, 20])

        assert [c.width for c in tenths.columns] == [c.width for c in hundreds.columns]

    def test_the_wrong_number_of_widths_is_reported_and_ignored(self):
        pres, table = build(widths=[3, 1])
        widths = [column.width for column in table.columns]

        assert widths[0] == widths[1] == widths[2]
        assert W.TABLE_WIDTHS_IGNORED in codes(pres)
        assert "2 value(s) for 3 column(s)" in pres.warning_messages[0]

    def test_a_zero_width_is_refused_by_the_schema(self):
        with pytest.raises(ValueError):
            coerce_slides([{"type": "table", "rows": ROWS, "widths": [1, 0, 1]}])


class TestCellFills:

    def test_one_cell_takes_the_colour(self):
        _, table = build(fills=[{"row": 1, "col": 1, "color": "C00000"}])

        assert str(table.cell(1, 1).fill.fore_color.rgb) == "C00000"

    def test_a_row_without_a_column_fills_the_whole_row(self):
        _, table = build(fills=[{"row": 2, "color": "00B050"}])

        assert all(str(table.cell(2, col).fill.fore_color.rgb) == "00B050"
                   for col in range(3))

    def test_a_theme_name_follows_the_palette(self):
        _, table = build(fills=[{"row": 1, "color": "accent2"}])

        assert table.cell(1, 0).fill.fore_color.theme_color == MSO_THEME_COLOR.ACCENT_2

    @pytest.mark.parametrize("name, member", [
        ("accent1", MSO_THEME_COLOR.ACCENT_1), ("accent6", MSO_THEME_COLOR.ACCENT_6),
        ("dark1", MSO_THEME_COLOR.DARK_1), ("dark2", MSO_THEME_COLOR.DARK_2),
        ("light1", MSO_THEME_COLOR.LIGHT_1), ("light2", MSO_THEME_COLOR.LIGHT_2),
    ])
    def test_every_theme_name_writes_a_valid_scheme_colour(self, name, member):
        """Regression: "dark1" and friends went into <a:schemeClr val=…> as
        spelled, but DrawingML only knows "dk1"/"lt1"/"dk2"/"lt2". The file
        could not be read back and PowerPoint offered to repair it. Reading
        the colour back is what failed, so that is the assertion."""
        _, table = build(header_color=name, fills=[{"row": 1, "color": name}])

        assert table.cell(0, 0).fill.fore_color.theme_color == member
        assert table.cell(1, 0).fill.fore_color.theme_color == member

    def test_an_explicit_fill_beats_zebra_and_the_header(self):
        _, table = build(zebra=True, header_color="accent1",
                         fills=[{"row": 0, "col": 0, "color": "FFFF00"},
                                {"row": 2, "col": 0, "color": "FFFF00"}])

        assert str(table.cell(0, 0).fill.fore_color.rgb) == "FFFF00"
        assert str(table.cell(2, 0).fill.fore_color.rgb) == "FFFF00"

    def test_a_cell_outside_the_table_is_reported_and_skipped(self):
        pres, table = build(fills=[{"row": 9, "color": "FFFF00"},
                                   {"row": 1, "col": 1, "color": "FFFF00"}])

        assert W.TABLE_FILL_IGNORED in codes(pres)
        assert str(table.cell(1, 1).fill.fore_color.rgb) == "FFFF00"


class TestMergedCells:

    def test_a_block_spans_and_keeps_only_its_own_text(self):
        """python-pptx concatenates on merge; the block must not repeat cells."""
        _, table = build(merges=[{"row": 0, "col": 1, "col_span": 2}])

        assert table.cell(0, 1).text == "Impact"
        assert table.cell(0, 2).is_spanned
        assert table.cell(0, 1).span_width == 2

    def test_a_vertical_block(self):
        _, table = build(merges=[{"row": 1, "col": 0, "row_span": 2}])

        assert table.cell(1, 0).span_height == 2
        assert table.cell(2, 0).is_spanned

    def test_the_rest_of_the_table_is_untouched(self):
        _, table = build(merges=[{"row": 0, "col": 1, "col_span": 2}])

        assert [table.cell(1, col).text for col in range(3)] == ["Data loss", "High", "Ops"]

    def test_overlapping_blocks_are_reported_not_raised(self):
        """python-pptx raises mid-merge, which would take the whole deck."""
        pres, table = build(merges=[{"row": 0, "col": 0, "col_span": 2},
                                    {"row": 0, "col": 1, "col_span": 2}])

        assert W.TABLE_MERGE_IGNORED in codes(pres)
        assert table.cell(0, 1).is_spanned
        assert not table.cell(0, 2).is_spanned

    def test_a_huge_block_is_rejected_before_it_is_expanded(self):
        """Regression: spans have no upper bound, and every cell of a block
        went into a set before the bounds check, so one merge on a 2x2 table
        could allocate millions of cells. The bounds check now runs first."""
        import time

        start = time.perf_counter()
        pres, table = build(merges=[{"row": 0, "col": 0,
                                     "row_span": 100_000, "col_span": 100_000}])

        assert time.perf_counter() - start < 5.0
        assert W.TABLE_MERGE_IGNORED in codes(pres)
        assert not table.cell(0, 1).is_spanned

    def test_blocks_that_only_touch_both_apply(self):
        pres, table = build(merges=[{"row": 1, "col": 0, "row_span": 2},
                                    {"row": 1, "col": 1, "col_span": 2}])

        assert W.TABLE_MERGE_IGNORED not in codes(pres)
        assert table.cell(2, 0).is_spanned and table.cell(1, 2).is_spanned

    def test_a_block_running_past_the_edge_is_reported_and_skipped(self):
        pres, table = build(merges=[{"row": 0, "col": 2, "col_span": 2}])

        assert W.TABLE_MERGE_IGNORED in codes(pres)
        assert not table.cell(0, 2).is_spanned

    def test_a_single_cell_merge_is_refused_by_the_schema(self):
        with pytest.raises(ValueError, match="more than one cell"):
            coerce_slides([{"type": "table", "rows": ROWS,
                            "merges": [{"row": 0, "col": 0}]}])


class TestTogether:

    def test_widths_fills_and_merges_on_one_table(self):
        pres, table = build(
            widths=[3, 1, 1],
            fills=[{"row": 1, "col": 1, "color": "C00000"}],
            merges=[{"row": 0, "col": 1, "col_span": 2}],
        )

        assert pres.warnings == []
        assert table.columns[0].width > table.columns[1].width
        assert str(table.cell(1, 1).fill.fore_color.rgb) == "C00000"
        assert table.cell(0, 2).is_spanned

    def test_none_of_them_is_required(self):
        pres, table = build()

        assert pres.warnings == []
        assert table.cell(0, 0).text == "Risk"
