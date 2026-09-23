"""Only http, https, mailto and tel links become clickable, in both tools.

From the hardening review: any ``[label](target)`` became an external
hyperlink relationship — file: and UNC paths, custom protocol handlers,
javascript: — which the reader's Office resolves on click. A model steered by
injected content could plant a link that leaks credentials or launches a local
handler. The shared grammar now decides which schemes are allowed; a refused
link keeps its label as plain text and the build reports it.
"""
import sys
from pathlib import Path

project_root = Path(__file__).parent.parent
sys.path.insert(0, str(project_root))

import pytest
from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE
from pptx import Presentation as PptxReader

from inline_markdown import is_safe_link_target, refused_link_targets

SAFE = ["https://example.com/a?b=c", "http://x.y", "HTTPS://EXAMPLE.COM",
        "mailto:someone@example.com", "tel:+420123456789"]
REFUSED = ["file:///etc/passwd", "\\\\server\\share\\doc", "C:\\Users\\x",
           "javascript:void0", "ms-word:ofe|u|https://x", "smb://host/share",
           "example.com", "../relative", "#anchor"]


class TestTheRule:

    @pytest.mark.parametrize("target", SAFE)
    def test_safe(self, target):
        assert is_safe_link_target(target)

    @pytest.mark.parametrize("target", REFUSED)
    def test_refused(self, target):
        assert not is_safe_link_target(target)

    def test_the_scan_skips_code_spans(self):
        text = "`[c](file:x)` [bad](file:y) [ok](https://z)"
        assert refused_link_targets(text) == ["file:y"]

    def test_the_scan_counts_image_shaped_text(self):
        """Review finding on #198: the scan skipped every "![alt](src)", but
        the inline renderers have no image branch — they draw "!" and a link,
        refused or not — so a refusal there went unreported. Only Word's
        whole-line image is an image, and Word removes those lines itself."""
        assert refused_link_targets("see ![x](file:y) here") == ["file:y"]


def _pptx_links(title):
    from pptx_tools import warnings as W
    from pptx_tools.slide_builder import PowerpointPresentation

    pres = PowerpointPresentation([{"type": "content", "title": title, "body": "- x"}], "16:9")
    runs = [run for shape in PptxReader(pres.save()).slides[0].shapes if shape.has_text_frame
            for para in shape.text_frame.paragraphs for run in para.runs]
    links = {run.text: run.hyperlink.address for run in runs}
    refused = [w for w in pres.warnings if w.code == W.LINK_REFUSED]
    return links, refused


class TestPowerPoint:

    @pytest.mark.parametrize("target", SAFE)
    def test_a_safe_link_is_clickable(self, target):
        links, refused = _pptx_links(f"see [here]({target})")
        assert links["here"] == target
        assert refused == []

    @pytest.mark.parametrize("target", REFUSED)
    def test_a_refused_link_keeps_its_label_as_text_and_is_reported(self, target):
        links, refused = _pptx_links(f"see [here]({target})")
        assert links["here"] is None
        assert len(refused) == 1
        assert refused[0].severity == "warning"
        assert refused[0].slide is None


def _docx_links(markdown):
    from docx_tools import warnings as W
    from docx_tools.base_docx_tool import _markdown_to_word_buffer

    buffer, warnings = _markdown_to_word_buffer(markdown)
    doc = Document(buffer)
    targets = [rel.target_ref for rel in doc.part.rels.values()
               if rel.reltype == RELATIONSHIP_TYPE.HYPERLINK]
    refused = [w for w in warnings if w.code == W.LINK_REFUSED]
    return doc, targets, refused


class TestImageShapedLinks:
    """Review finding on #198: see test_the_scan_counts_image_shaped_text."""

    def test_pptx_reports_one_in_a_bullet(self):
        _, refused = _pptx_links("see ![here](file:///x)")
        assert len(refused) == 1

    def test_word_reports_one_mid_paragraph(self):
        _, targets, refused = _docx_links("See ![here](file:///x) for details.")
        assert targets == []
        assert len(refused) == 1

    def test_word_does_not_report_a_whole_line_image(self):
        _, _, refused = _docx_links("Text\n\n![alt](file:///x)\n\nMore")
        assert refused == []

    def test_word_does_not_report_a_link_inside_a_code_fence(self):
        _, targets, refused = _docx_links("```\n[a](file:///x)\n```\n\n[b](file:///y)")
        assert targets == []
        assert len(refused) == 1
        assert "file:///y" in refused[0].message and "file:///x" not in refused[0].message


class TestWord:

    @pytest.mark.parametrize("target", SAFE)
    def test_a_safe_link_is_clickable(self, target):
        _, targets, refused = _docx_links(f"see [here]({target})")
        assert targets == [target]
        assert refused == []

    @pytest.mark.parametrize("target", REFUSED)
    def test_a_refused_link_keeps_its_label_as_text_and_is_reported(self, target):
        doc, targets, refused = _docx_links(f"see [here]({target}) now")
        assert targets == []
        assert "see here now" in "\n".join(p.text for p in doc.paragraphs)
        assert len(refused) == 1
        assert refused[0].severity == "warning"
