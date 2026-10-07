"""Word content controls in dynamic DOCX templates are filled by their Tag.

Phase 1 kinds: check box, drop-down list, plain text. A tag names an argument
(``full_name``); a check box may name an option too (``size=small``). Untagged
controls, foreign tags, unsupported kinds and unset values leave the control
exactly as the template has it.
"""
import sys
from pathlib import Path

project_root = Path(__file__).parent.parent
sys.path.insert(0, str(project_root))

import pytest  # noqa: E402
from docx import Document  # noqa: E402

from docx_tools.content_controls import (  # noqa: E402
    CHECKBOX, DROPDOWN, RICH_TEXT, TEXT, classify, describe_content_controls,
    parse_tag, resolve_content_controls,
)

from . import docx_content_control_helpers as cc  # noqa: E402

W = cc.W


def _doc(*controls):
    doc = Document()
    for control in controls:
        cc.add_inline(doc.add_paragraph(), control)
    return doc


# --- tags and kinds -----------------------------------------------------------

@pytest.mark.parametrize("tag,expected", [
    ("full_name", ("full_name", None)), ("size=small", ("size", "small")),
    (" size = small ", ("size", "small")), ("", ("", None)),
])
def test_parse_tag(tag, expected):
    assert parse_tag(tag) == expected


def test_classify_reads_the_type_element():
    def kind(control):
        return classify(control.find(f"{{{W}}}sdtPr"))
    assert kind(cc.checkbox("a")) == CHECKBOX
    assert kind(cc.dropdown("a")) == DROPDOWN
    assert kind(cc.text("a")) == TEXT
    assert kind(cc.rich_text("a")) == RICH_TEXT


def test_describe_lists_every_control_with_its_items_and_alias():
    doc = _doc(cc.checkbox("size=small", alias="Small"), cc.dropdown("plan"), cc.rich_text())
    found = describe_content_controls(doc)
    assert [(f["name"], f["option"], f["kind"], f["supported"]) for f in found] == [
        ("size", "small", CHECKBOX, True), ("plan", None, DROPDOWN, True),
        ("", None, RICH_TEXT, False)]
    assert found[0]["alias"] == "Small"
    assert found[1]["items"] == ["basic", "pro"]


# --- check boxes ---------------------------------------------------------------

def test_option_tag_ticks_only_the_matching_box():
    small, medium = cc.checkbox("size=small"), cc.checkbox("size=medium", True)
    resolve_content_controls(_doc(small, medium), {"size": "SMALL"})
    assert cc.checked(small) == (True, "☒")
    assert cc.checked(medium) == (False, "☐")


def test_a_list_value_ticks_every_listed_option():
    boxes = [cc.checkbox(f"channels={o}") for o in ("email", "phone", "post")]
    resolve_content_controls(_doc(*boxes), {"channels": ["email", "post"]})
    assert [cc.checked(b)[0] for b in boxes] == [True, False, True]


def test_bare_tag_follows_a_boolean_and_false_unticks():
    yes, no = cc.checkbox("consent"), cc.checkbox("opt_out", True)
    resolve_content_controls(_doc(yes, no), {"consent": True, "opt_out": False})
    assert cc.checked(yes)[0] is True
    assert cc.checked(no)[0] is False


def test_an_empty_list_is_a_value_and_unticks():
    box = cc.checkbox("channels=email", True)
    resolve_content_controls(_doc(box), {"channels": []})
    assert cc.checked(box)[0] is False


def test_an_unset_value_leaves_the_box_as_the_template_has_it():
    box = cc.checkbox("size=small", True)
    resolve_content_controls(_doc(box), {"size": None})
    assert cc.checked(box) == (True, "☒")


# --- drop-down lists -----------------------------------------------------------

def test_dropdown_selects_by_value_and_clears_the_prompt():
    dd = cc.dropdown("plan")
    assert resolve_content_controls(_doc(dd), {"plan": "pro"}) == 1
    assert cc.content_text(dd) == "Pro plan"
    pr = dd.find(f"{{{W}}}sdtPr")
    assert pr.find(f"{{{W}}}dropDownList").get(f"{{{W}}}lastValue") == "pro"
    assert pr.find(f"{{{W}}}showingPlcHdr") is None
    rpr = dd.find(f".//{{{W}}}rPr")
    assert rpr.find(f"{{{W}}}rStyle") is None      # no longer grey prompt text
    assert rpr.find(f"{{{W}}}b") is not None       # the control's own formatting stays


def test_dropdown_also_matches_the_display_text():
    dd = cc.dropdown("plan")
    resolve_content_controls(_doc(dd), {"plan": "pro PLAN"})
    assert cc.content_text(dd) == "Pro plan"


def test_dropdown_without_a_matching_item_is_left_alone():
    dd = cc.dropdown("plan")
    assert resolve_content_controls(_doc(dd), {"plan": "enterprise"}) == 0
    assert cc.content_text(dd) == "Choose an item."
    assert dd.find(f".//{{{W}}}showingPlcHdr") is not None


# --- plain text ----------------------------------------------------------------

def test_text_replaces_the_prompt_in_the_controls_formatting():
    t = cc.text("full_name")
    resolve_content_controls(_doc(t), {"full_name": "Jane Doe"})
    assert cc.content_text(t) == "Jane Doe"
    runs = t.findall(f".//{{{W}}}r")
    assert len(runs) == 1
    rpr = runs[0].find(f"{{{W}}}rPr")
    assert rpr.find(f"{{{W}}}rStyle") is None and rpr.find(f"{{{W}}}i") is not None


def test_multiline_text_gets_line_breaks_single_line_gets_spaces():
    multi, single = cc.text("address", multiline=True), cc.text("full_name")
    resolve_content_controls(_doc(multi, single), {"address": "1 Main St\nSpringfield",
                                                   "full_name": "Jane\nDoe"})
    assert cc.content_text(multi) == "1 Main St\nSpringfield"
    assert cc.content_text(single) == "Jane Doe"


def test_block_level_text_control_writes_into_its_paragraph():
    doc = Document()
    t = cc.add_block(doc, cc.text("notes", block=True))
    resolve_content_controls(doc, {"notes": "Done"})
    assert cc.content_text(t) == "Done"
    assert t.find(f"{{{W}}}sdtContent").find(f"{{{W}}}p") is not None


def test_text_values_that_are_blank_keep_the_prompt():
    t = cc.text("full_name")
    resolve_content_controls(_doc(t), {"full_name": "  "})
    assert cc.content_text(t) == "Click here."


def test_a_list_value_in_a_text_control_is_comma_joined():
    t = cc.text("channels")
    resolve_content_controls(_doc(t), {"channels": ["email", "phone"]})
    assert cc.content_text(t) == "email, phone"


def test_a_filled_control_loses_its_data_binding():
    t = cc.text("full_name", binding=True)
    resolve_content_controls(_doc(t), {"full_name": "Jan"})
    assert t.find(f".//{{{W}}}dataBinding") is None


def test_an_unfilled_control_keeps_its_data_binding():
    t = cc.text("full_name", binding=True)
    resolve_content_controls(_doc(t), {"full_name": None})
    assert t.find(f".//{{{W}}}dataBinding") is not None


# --- what is left alone ----------------------------------------------------------

def test_untagged_foreign_and_unsupported_controls_are_untouched():
    untagged, foreign, rich = cc.text(None), cc.text("unknown"), cc.rich_text("full_name")
    assert resolve_content_controls(_doc(untagged, foreign, rich), {"full_name": "Jan"}) == 0
    assert cc.content_text(untagged) == "Click here."
    assert cc.content_text(foreign) == "Click here."
    assert cc.content_text(rich) == "rich"


def test_an_option_tag_on_a_non_checkbox_is_ignored():
    t = cc.text("size=small")
    assert resolve_content_controls(_doc(t), {"size": "small"}) == 0


# --- where controls are found -------------------------------------------------------

def test_controls_in_tables_and_headers_are_filled():
    doc = Document()
    in_cell = cc.add_inline(doc.add_table(rows=1, cols=1).cell(0, 0).paragraphs[0],
                            cc.text("full_name"))
    in_header = cc.add_inline(doc.sections[0].header.paragraphs[0], cc.checkbox("consent"))
    assert resolve_content_controls(doc, {"full_name": "Jan", "consent": True}) == 2
    assert cc.content_text(in_cell) == "Jan"
    assert cc.checked(in_header)[0] is True


def test_the_document_still_saves_and_reopens_with_its_controls(tmp_path):
    doc = _doc(cc.checkbox("consent"), cc.dropdown("plan"), cc.text("full_name"))
    resolve_content_controls(doc, {"consent": True, "plan": "basic", "full_name": "Jan"})
    out = tmp_path / "filled.docx"
    doc.save(out)
    xml = Document(out).element.xml
    assert xml.count("<w:sdt>") == 3
    assert "Basic plan" in xml and "Jan" in xml


def test_controls_in_distinct_headers_of_two_sections_are_all_filled():
    """Each unlinked header is its own part; none may be skipped (PR #200 review)."""
    from docx.enum.section import WD_SECTION

    doc = Document()
    first = cc.add_inline(doc.sections[0].header.paragraphs[0], cc.text("full_name", sdt_id=1))
    second_section = doc.add_section(WD_SECTION.NEW_PAGE)
    second_section.header.is_linked_to_previous = False
    second = cc.add_inline(second_section.header.paragraphs[0], cc.text("full_name", sdt_id=2))
    assert resolve_content_controls(doc, {"full_name": "Jane"}) == 2
    assert cc.content_text(first) == cc.content_text(second) == "Jane"


def test_the_choose_an_item_prompt_is_not_a_choice():
    dd = cc.dropdown("plan", items=(("", "Choose an item."), ("basic", "Basic plan")))
    doc = _doc(dd)
    assert describe_content_controls(doc)[0]["items"] == ["basic"]
    assert resolve_content_controls(doc, {"plan": "Choose an item."}) == 0
    assert resolve_content_controls(doc, {"plan": "basic"}) == 1
    assert cc.content_text(dd) == "Basic plan"


# --- what the caller is told ------------------------------------------------------

def _channel():
    from docx_tools.warnings import channel
    return channel()


def test_a_dropdown_without_the_item_is_reported_with_its_items():
    warnings = _channel()
    resolve_content_controls(_doc(cc.dropdown("plan")), {"plan": "enterprise"}, warnings)
    [w] = warnings.as_dicts()
    assert (w["code"], w["severity"], w["tag"]) == ("control_item_missing", "error", "plan")
    assert "'enterprise'" in w["message"] and "basic, pro" in w["message"]


def test_a_value_for_a_kind_that_is_not_filled_is_reported():
    warnings = _channel()
    resolve_content_controls(_doc(cc.rich_text("notes"), cc.text("size=small", sdt_id=9)),
                             {"notes": "x", "size": "small"}, warnings)
    assert [(w["code"], w["tag"]) for w in warnings.as_dicts()] == [
        ("control_not_filled", "notes"), ("control_not_filled", "size=small")]


def test_nothing_is_reported_for_unset_values_or_successful_fills():
    warnings = _channel()
    resolve_content_controls(_doc(cc.dropdown("plan"), cc.rich_text("notes", sdt_id=8)),
                             {"plan": None, "notes": ""}, warnings)
    resolve_content_controls(_doc(cc.dropdown("plan")), {"plan": "pro"}, warnings)
    assert warnings.as_dicts() == []
