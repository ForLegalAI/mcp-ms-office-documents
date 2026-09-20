"""Cloning a template (#164).

"Same as the formal letter, but for the Prague office" was the commonest
real task the UI could not do: it meant creating one from scratch and
re-typing every argument name, type, default and description by hand, where
a dozen arguments each with an AI-facing description is normal.

The property worth guarding is that the asset is *copied*, not shared. Two
specs pointing at one file would make "Replace document" on either silently
change the other — a duplication is the cheaper mistake.
"""
import io
import re
import sys
from pathlib import Path

project_root = Path(__file__).parent.parent
sys.path.insert(0, str(project_root))

import pytest
from docx import Document
from starlette.testclient import TestClient

import admin.store as store_mod
import metrics
import template_utils as tu
from config import Config

PPTX_SOURCE = project_root / "default_templates" / "default_pptx_template_16_9.pptx"


@pytest.fixture
def admin_client(tmp_path, monkeypatch):
    custom = tmp_path / "custom"
    cfg = tmp_path / "config"
    custom.mkdir()
    cfg.mkdir()
    monkeypatch.setattr(store_mod, "_APP_CUSTOM_DIR", tmp_path / "nx1")
    monkeypatch.setattr(store_mod, "_APP_CONFIG_DIR", tmp_path / "nx2")
    monkeypatch.setattr(store_mod, "_LOCAL_CUSTOM_DIR", custom)
    monkeypatch.setattr(store_mod, "_LOCAL_CONFIG_DIR", cfg)
    monkeypatch.setattr(tu, "APP_CUSTOM_DIR", tmp_path / "nx1")
    monkeypatch.setattr(tu, "LOCAL_CUSTOM_DIR", custom)
    import pptx_tools.templates as templates_mod
    monkeypatch.setattr(templates_mod, "_APP_CONFIG_DIR", tmp_path / "nx2")
    monkeypatch.setattr(templates_mod, "_LOCAL_CONFIG_DIR", cfg)
    templates_mod.clear_cache()
    monkeypatch.setenv("ADMIN_ENABLED", "true")
    monkeypatch.setenv("ADMIN_PASSWORD", "pw")
    monkeypatch.delenv("API_KEY", raising=False)

    from fastmcp import FastMCP
    from admin.app import build_combined_app

    mcp = FastMCP("test-clone")
    client = TestClient(build_combined_app(mcp, Config.from_env()))
    metrics.reset()
    client.__enter__()
    client.post("/admin/login", data={"password": "pw"})
    yield client, mcp, custom
    client.__exit__(None, None, None)
    metrics.reset()


def _docx_bytes(text="Dear {{recipient}}"):
    doc = Document()
    doc.add_paragraph(text)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _first_paragraph(path: Path) -> str:
    return Document(str(path)).paragraphs[0].text


def _csrf(client) -> str:
    html = client.get("/admin/new/docx").text
    tag = re.search(r'<input[^>]*name="csrf"[^>]*>', html)
    return re.search(r'value="([^"]*)"', tag.group(0)).group(1)


def _post(client, url, data=None, files=None):
    payload = dict(data or {})
    payload["csrf"] = _csrf(client)
    return client.post(url, data=payload, files=files)


def _store():
    return store_mod.FileTemplateStore.from_config()


def _letter(client, name="formal_letter", body="Dear {{recipient}}"):
    """A saved docx template with two described arguments."""
    _post(client, "/admin/docx/draft", data={"name": name},
          files={"file": (f"{name}.docx", _docx_bytes(body),
                          "application/octet-stream")})
    _post(client, "/admin/docx/save", data={
        "kind": "docx", "asset_filename": f"{name}.docx", "name": name,
        "original_name": name, "title": "Formal Letter",
        "description": "A formal business letter",
        "arg_name": ["recipient", "subject"],
        "arg_type": ["string", "string"],
        "arg_required": ["true", "false"],
        "arg_default": ["", "Re: your enquiry"],
        "arg_desc": ["Full name of the recipient", "What the letter is about"],
    })
    return name


def _deck(client, name="brand", default=True):
    data = {"kind": "pptx", "asset_filename": f"{name}.pptx", "name": name,
            "original_name": name, "description": "Brand deck"}
    if default:
        data["is_default"] = "on"
    _post(client, "/admin/pptx/draft", data={"name": name},
          files={"file": (f"{name}.pptx", PPTX_SOURCE.read_bytes(),
                          "application/octet-stream")})
    _post(client, "/admin/pptx/save", data=data)
    return name


async def _tool_names(mcp):
    return [t.name for t in await mcp.list_tools()]


# ---------------------------------------------------------------------------
# What a clone carries
# ---------------------------------------------------------------------------


def test_the_clone_keeps_every_argument_and_its_description(admin_client):
    """The whole point: re-typing these by hand is the task being removed."""
    client, _mcp, _custom = admin_client
    _letter(client)

    _post(client, "/admin/docx/formal_letter/clone", data={"name": "prague_letter"})

    copy = _store().get_spec("docx", "prague_letter")
    assert copy is not None
    assert [a["name"] for a in copy["args"]] == ["recipient", "subject"]
    assert [a["description"] for a in copy["args"]] == [
        "Full name of the recipient", "What the letter is about"]
    assert [a["required"] for a in copy["args"]] == [True, False]
    assert copy["args"][1]["default"] == "Re: your enquiry"
    assert copy["description"] == "A formal business letter"


@pytest.mark.anyio
async def test_the_clone_is_registered_and_the_original_survives(admin_client):
    client, mcp, _custom = admin_client
    _letter(client)

    _post(client, "/admin/docx/formal_letter/clone", data={"name": "prague_letter"})

    names = await _tool_names(mcp)
    assert "prague_letter" in names
    assert "formal_letter" in names, "cloning must not disturb the original"


def test_the_clone_lands_on_its_own_edit_page(admin_client):
    """Not the index: the description almost always needs changing at once."""
    client, _mcp, _custom = admin_client
    _letter(client)

    r = _post(client, "/admin/docx/formal_letter/clone",
              data={"name": "prague_letter"})

    assert "Edit prague_letter" in r.text
    assert "Copied from formal_letter" in r.text


# ---------------------------------------------------------------------------
# The asset is copied, not shared
# ---------------------------------------------------------------------------


def test_the_asset_is_copied_under_the_new_name(admin_client):
    client, _mcp, custom = admin_client
    _letter(client)

    _post(client, "/admin/docx/formal_letter/clone", data={"name": "prague_letter"})

    copy = _store().get_spec("docx", "prague_letter")
    assert copy["docx_path"] == "prague_letter.docx"
    assert (custom / "prague_letter.docx").exists()
    assert (custom / "formal_letter.docx").exists(), "the original file stays"


def test_replacing_the_clones_document_does_not_touch_the_original(admin_client):
    """The reason the file is copied rather than shared.

    Sharing one file between two specs would make "Replace document" on
    either one silently change the other — the failure this design exists to
    prevent, and one nothing in the UI would report.
    """
    client, _mcp, custom = admin_client
    _letter(client, body="ORIGINAL WORDING {{recipient}}")
    _post(client, "/admin/docx/formal_letter/clone", data={"name": "prague_letter"})

    _post(client, "/admin/docx/prague_letter/reupload",
          files={"file": ("prague_letter.docx",
                          _docx_bytes("REPLACED WORDING {{recipient}}"),
                          "application/octet-stream")})

    assert "REPLACED" in _first_paragraph(custom / "prague_letter.docx")
    assert "ORIGINAL" in _first_paragraph(custom / "formal_letter.docx"), \
        "replacing the clone's document must not rewrite the original's"


def test_the_source_extension_is_kept(admin_client):
    """A clone of a .potx stays a .potx rather than being renamed .pptx."""
    client, _mcp, custom = admin_client
    _post(client, "/admin/pptx/draft", data={"name": "designer"},
          files={"file": ("designer.potx", PPTX_SOURCE.read_bytes(),
                          "application/octet-stream")})
    _post(client, "/admin/pptx/save", data={
        "kind": "pptx", "asset_filename": "designer.potx", "name": "designer",
        "original_name": "designer", "description": "d"})
    assert _store().get_spec("pptx", "designer")["pptx_path"] == "designer.potx"

    _post(client, "/admin/pptx/designer/clone", data={"name": "designer_copy"})

    assert _store().get_spec("pptx", "designer_copy")["pptx_path"] == "designer_copy.potx"
    assert (custom / "designer_copy.potx").exists()


# ---------------------------------------------------------------------------
# What a clone must not carry
# ---------------------------------------------------------------------------


def test_the_pptx_default_flag_is_not_carried(admin_client):
    """Two defaults is a state the registry resolves silently by picking one."""
    client, _mcp, _custom = admin_client
    _deck(client, "brand", default=True)

    _post(client, "/admin/pptx/brand/clone", data={"name": "brand_copy"})

    assert _store().get_spec("pptx", "brand").get("default") is True
    assert _store().get_spec("pptx", "brand_copy").get("default") is None


def test_exactly_one_pptx_template_stays_the_default(admin_client):
    """The property the dropped flag exists to preserve, checked end to end."""
    import pptx_tools.templates as templates_mod

    client, _mcp, _custom = admin_client
    _deck(client, "brand", default=True)
    _post(client, "/admin/pptx/brand/clone", data={"name": "brand_copy"})
    templates_mod.clear_cache()

    defaults = [s.name for s in templates_mod.load_specs(force=True) if s.is_default]
    assert defaults == ["brand"], f"expected one default, got {defaults}"


@pytest.mark.anyio
async def test_cloning_a_disabled_template_gives_a_disabled_clone(admin_client):
    """A copy of something switched off should not arrive switched on.

    The description is almost always wrong until edited, so landing a live
    tool from a template its owner had deliberately taken out of service is
    the wrong default.
    """
    client, mcp, _custom = admin_client
    _letter(client)
    _post(client, "/admin/docx/formal_letter/enabled", data={"enabled": ""})

    _post(client, "/admin/docx/formal_letter/clone", data={"name": "prague_letter"})

    assert _store().get_spec("docx", "prague_letter")["enabled"] is False
    assert "prague_letter" not in await _tool_names(mcp)


# ---------------------------------------------------------------------------
# Refusals
# ---------------------------------------------------------------------------


def test_cloning_onto_an_existing_name_is_refused(admin_client):
    client, _mcp, _custom = admin_client
    _letter(client, "formal_letter")
    _letter(client, "occupied")

    r = _post(client, "/admin/docx/formal_letter/clone", data={"name": "occupied"})

    assert "already exists" in r.text
    assert _store().get_spec("docx", "occupied")["docx_path"] == "occupied.docx", \
        "the occupant must be untouched"


def test_cloning_to_an_invalid_name_is_refused(admin_client):
    client, _mcp, _custom = admin_client
    _letter(client)

    r = _post(client, "/admin/docx/formal_letter/clone", data={"name": "not a name!"})

    assert r.status_code == 200
    assert "Clone formal_letter" in r.text, "it should come back to the form"
    assert _store().get_spec("docx", "formal_letter") is not None


def test_cloning_an_unknown_template_is_not_found(admin_client):
    client, _mcp, _custom = admin_client
    r = client.get("/admin/docx/ghost/clone")
    assert "not found" in r.text.lower()


def test_cloning_needs_the_csrf_token(admin_client):
    client, _mcp, _custom = admin_client
    _letter(client)

    client.post("/admin/docx/formal_letter/clone", data={"name": "sneaky"})
    assert _store().get_spec("docx", "sneaky") is None


def test_a_missing_source_file_is_reported_not_crashed(admin_client):
    client, _mcp, custom = admin_client
    _letter(client)
    (custom / "formal_letter.docx").unlink()

    r = _post(client, "/admin/docx/formal_letter/clone", data={"name": "doomed"})

    assert r.status_code == 200
    assert "Asset not found: formal_letter.docx" in r.text, \
        "the message should name the file that is missing"
    assert "Clone formal_letter" in r.text, "it should come back to the form"
    assert _store().get_spec("docx", "doomed") is None


# ---------------------------------------------------------------------------
# Where the action is offered
# ---------------------------------------------------------------------------


def test_clone_is_offered_on_the_row_and_the_edit_page(admin_client):
    client, _mcp, _custom = admin_client
    _letter(client)

    assert "/admin/docx/formal_letter/clone" in client.get("/admin/").text
    assert "/admin/docx/formal_letter/clone" in \
        client.get("/admin/docx/formal_letter/edit").text


def test_clone_is_not_offered_while_creating(admin_client):
    """There is nothing saved to clone yet on the configure form."""
    client, _mcp, _custom = admin_client
    r = _post(client, "/admin/docx/draft", data={"name": "brand_new"},
              files={"file": ("brand_new.docx", _docx_bytes(),
                              "application/octet-stream")})
    assert "/clone" not in r.text


def test_the_clone_form_says_what_comes_across(admin_client):
    client, _mcp, _custom = admin_client
    _letter(client)

    html = client.get("/admin/docx/formal_letter/clone").text
    assert "formal_letter.docx" in html, "it should name the file being copied"
    assert "copy" in html.lower()
