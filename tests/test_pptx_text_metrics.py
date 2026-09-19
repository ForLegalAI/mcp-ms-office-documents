"""Measuring text instead of counting characters (#125).

The fill estimate decided how full a text box was from the number of
characters in it, with one mean glyph width for every face and every letter —
so "WWW WWW" and "iii iii" were the same length. It now measures against a
real font file where one can be loaded, and keeps the old arithmetic only
when none can.
"""
import sys
from pathlib import Path
from unittest.mock import patch

project_root = Path(__file__).parent.parent
sys.path.insert(0, str(project_root))

import pytest
from pptx import Presentation as PptxReader
from pptx.util import Inches

import pptx_tools.helpers as helpers
from pptx_tools import text_metrics
from pptx_tools.schema import Bullet
from pptx_tools.slide_builder import PowerpointPresentation
from pptx_tools.text_metrics import (
    _candidate_names, font_name, load_font, measure_lines, theme_body_typeface,
)

BASE_16_9 = project_root / "default_templates" / "default_pptx_template_16_9.pptx"

WIDTH, HEIGHT = Inches(6), Inches(3)


def fill(text, typeface="Arial", **kwargs):
    return helpers.estimate_text_fill([Bullet(text=text)], WIDTH, HEIGHT, 18.0,
                                      typeface=typeface, **kwargs)


class TestItMeasuresGlyphs:

    def test_the_same_characters_are_not_the_same_width(self):
        """The bug in one assertion: identical lengths, different text."""
        thin, wide = "iii " * 80, "WWW " * 80
        assert len(thin) == len(wide)

        assert fill(thin) < 0.6
        assert fill(wide) > 1.0

    def test_character_counting_could_not_tell_them_apart(self):
        thin, wide = "iii " * 80, "WWW " * 80
        with patch.object(helpers, "measure_lines", return_value=None):
            assert fill(thin) == fill(wide)

    def test_more_text_is_still_more_full(self):
        assert fill("word " * 20) < fill("word " * 200)

    def test_an_empty_box_is_empty(self):
        assert helpers.estimate_text_fill([], WIDTH, HEIGHT, 18.0) == 0.0

    def test_a_deeper_level_has_less_room(self):
        text = "word " * 60
        shallow = helpers.estimate_text_fill([Bullet(text=text, level=1)],
                                             WIDTH, HEIGHT, 18.0, typeface="Arial")
        deep = helpers.estimate_text_fill([Bullet(text=text, level=3)],
                                          WIDTH, HEIGHT, 18.0, typeface="Arial")
        assert deep > shallow


class TestWrapping:

    def test_lines_are_counted_by_wrapping_at_spaces(self):
        one_line = measure_lines([("short", 400)], "Arial", 18)
        several = measure_lines([("word " * 40, 400)], "Arial", 18)

        assert one_line == 1
        assert several > 1

    def test_a_word_wider_than_the_box_breaks_rather_than_counting_as_one(self):
        """Counting it as one line would under-measure a long unbroken string."""
        lines = measure_lines([("W" * 200, 100)], "Arial", 18)

        assert lines > 1

    def test_an_empty_bullet_still_occupies_a_line(self):
        assert measure_lines([("", 400)], "Arial", 18) == 1

    def test_a_zero_width_box_does_not_loop_forever(self):
        assert measure_lines([("some text", 0)], "Arial", 18) == 1


class TestChoosingAFace:

    def test_a_metric_compatible_substitute_is_preferred_over_a_generic_one(self):
        """Carlito's advances are Calibri's, so measuring it is exact."""
        candidates = _candidate_names("Calibri")

        assert candidates.index("Carlito") < candidates.index("DejaVuSans")

    def test_the_theme_reference_form_is_left_to_the_fallback(self):
        assert _candidate_names("+mj-lt")[0] in text_metrics.GENERIC_SANS

    def test_an_unknown_face_still_measures_with_something(self):
        font = load_font("Aptos", 18)

        assert font is not None
        assert font_name(font).endswith((".ttf", ".otf"))

    def test_fonts_are_cached_per_face_and_size(self):
        assert load_font("Arial", 18) is load_font("Arial", 18)

    def test_no_font_at_all_means_no_measurement(self):
        with patch.object(text_metrics, "GENERIC_SANS", ()):
            load_font.cache_clear()
            assert measure_lines([("text", 400)], "NoSuchFaceAnywhere", 18) is None
        load_font.cache_clear()


class TestTheImageCarriesWhatTheTablePromises:
    """A substitute the image lacks is a promise the fallback quietly breaks.

    METRIC_COMPATIBLE says "measuring this face measures the real thing". That
    only holds where the file exists, so the runtime image installs one package
    per family it can — and the two drifted apart once already (#134 review).

    Two families it cannot: Alpine packages neither Caladea nor Gelasio, so
    naming them in the Dockerfile failed the image build outright and
    v4.0-beta.4 published no image at all. They are listed as unpackaged here
    rather than dropped, because the table is right about the faces and a
    Debian host does install them.
    """

    # The Alpine package that provides each substitute family; None where Alpine
    # has none, which means the container reaches GENERIC_SANS for that family.
    PACKAGES = {
        "Carlito": "font-carlito",
        "Caladea": None,
        "Gelasio": None,
        "LiberationSans": "font-liberation",
        "Liberation Sans": "font-liberation",
        "LiberationSerif": "font-liberation",
        "Liberation Serif": "font-liberation",
        "LiberationMono": "font-liberation",
        "Liberation Mono": "font-liberation",
        "Arimo": "font-croscore",
        "Tinos": "font-croscore",
        "Cousine": "font-croscore",
    }

    # Typefaces the image cannot measure exactly, for the reason above. Adding
    # to this set is a deliberate loss of precision, not a passing test.
    UNPACKAGED = {"cambria", "georgia"}

    @pytest.fixture
    def installed(self):
        return set(_dockerfile_font_line().split())

    def test_every_substitute_family_is_a_package_we_know(self):
        unknown = {
            face
            for faces in text_metrics.METRIC_COMPATIBLE.values()
            for face in faces
            if face not in self.PACKAGES
        }
        assert not unknown, f"no package mapping for {unknown}"

    def test_every_metric_compatible_face_is_installed(self, installed):
        for typeface, faces in text_metrics.METRIC_COMPATIBLE.items():
            packages = {self.PACKAGES[face] for face in faces} - {None}
            if not packages:
                assert typeface in self.UNPACKAGED, (
                    f"{typeface} claims {sorted(faces)}, none of which Alpine packages"
                )
                continue
            assert packages & installed, (
                f"{typeface} claims {sorted(faces)}, none of which the image installs"
            )

    def test_the_unmeasurable_families_are_the_documented_two(self):
        # Pins the cost of the missing packages: if Alpine gains font-caladea,
        # installing it is a real improvement, and this test asks for the edit.
        unpackaged = {
            typeface
            for typeface, faces in text_metrics.METRIC_COMPATIBLE.items()
            if not {self.PACKAGES[face] for face in faces} - {None}
        }
        assert unpackaged == self.UNPACKAGED

    def test_the_image_installs_nothing_the_table_does_not_name(self, installed):
        # The other direction: a package nobody measures with is dead weight in
        # the image, and a typo lands here rather than in a failed release build.
        known = {package for package in self.PACKAGES.values() if package}
        named = {token for token in installed if token.startswith("font-")}
        assert named <= known, f"the image installs {sorted(named - known)}, used by nothing"


def _dockerfile_font_line():
    """The Dockerfile's `apk add` line that installs the font packages."""
    dockerfile = (project_root / "Dockerfile").read_text(encoding="utf-8")
    line = [row for row in dockerfile.splitlines() if "apk add" in row and "font-" in row]
    assert line, "the runtime image installs no fonts"
    return line[0]


@pytest.mark.network
class TestThePackagesExist:
    """A package name Alpine does not have fails the image BUILD, not a test.

    That is how v4.0-beta.4 shipped no Docker image: `font-caladea` and
    `font-gelasio` do not exist, the whole `apk add` exited 2, and the release
    workflow was the first thing to find out. CI cannot run a build, but it can
    read Alpine's package index — so this test does, against every font package
    the Dockerfile names.
    """

    INDEXES = (
        "https://dl-cdn.alpinelinux.org/alpine/latest-stable/main/x86_64/APKINDEX.tar.gz",
        "https://dl-cdn.alpinelinux.org/alpine/latest-stable/community/x86_64/APKINDEX.tar.gz",
    )

    @pytest.fixture
    def available(self):
        import io
        import tarfile

        import requests

        names = set()
        for url in self.INDEXES:
            response = requests.get(url, timeout=30)
            response.raise_for_status()
            with tarfile.open(fileobj=io.BytesIO(response.content), mode="r:gz") as archive:
                index = archive.extractfile("APKINDEX").read().decode("utf-8", "replace")
            names |= {
                row[2:] for row in index.splitlines() if row.startswith("P:")
            }
        return names

    def test_every_font_package_exists_in_alpine(self, available):
        named = {token for token in _dockerfile_font_line().split()
                 if token.startswith("font-")}
        assert named, "the runtime image installs no fonts"
        assert named <= available, (
            f"the Dockerfile installs {sorted(named - available)}, which Alpine "
            "does not package — the image build will fail"
        )


class TestTheDecksOwnFace:

    def test_the_theme_body_font_is_read_from_the_template(self):
        assert theme_body_typeface(PptxReader(str(BASE_16_9))) == "Aptos"

    def test_the_builder_measures_with_it(self):
        pres = PowerpointPresentation([{"type": "content", "title": "C", "body": "- x"}], "16:9")

        assert pres._typeface == "Aptos"

    def test_a_presentation_without_a_theme_does_not_fail(self):
        presentation = PptxReader(str(BASE_16_9))
        with patch.object(type(presentation), "slide_masters",
                          property(lambda self: [])):
            assert theme_body_typeface(presentation) is None


class TestItStillDrivesTheDeck:

    def test_overfull_text_is_shrunk_and_reported(self):
        pres = PowerpointPresentation([{
            "type": "content", "title": "C",
            "body": "\n".join(f"- {'WWWW ' * 20}" for _ in range(20)),
        }], "16:9")

        assert any(w.code == "text_overflow" for w in pres.warnings)

    def test_text_that_fits_is_left_alone(self):
        pres = PowerpointPresentation([{
            "type": "content", "title": "C", "body": "- one\n- two",
        }], "16:9")

        assert pres.warnings == []
