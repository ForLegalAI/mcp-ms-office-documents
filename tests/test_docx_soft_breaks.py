"""Soft line breaks (issue #109): one rule, both entry points.

A soft break is a ``<w:br/>`` inside one Word paragraph. It has two spellings —
``<br>`` and a line ending in two spaces — and they must mean the same thing in
the base markdown tool and in a dynamic template's placeholder, in prose, in a
quote, inside an alignment block, in a list item and in a table cell.

The failures pinned here were all live before the fix:

* a soft break swallowed the block on the next line (its first list item, a
  table's header row, a ``---`` page break) and rendered it as literal text;
* a heading line ending in two spaces silently dropped everything after it;
* a block line that happened to end in two spaces stopped being a block;
* the two marker spaces were left in the run before the break;
* soft breaks were ignored inside ``<center>`` / ``<div align>``;
* a ``<br>`` inside a list item tore the item's second line into a paragraph;
* CRLF input lost every soft break.

Two more came out of the review on PR #110 and are pinned here too: a run
swallowed a blockquote line (the one block ``line_starts_block()`` did not
know), and the marker strip took a deliberate ``&nbsp;`` with it.

The deliberate limitation is pinned as well: ``<br>`` does not resume a numbered
list that began earlier in the document. Deciding that needs the renderer's own
running count, which does not exist yet when ``<br>`` is expanded; predicting it
drifted from the real one in seven distinct ways before the attempt was removed.
A continuation written with real newlines is unaffected.
"""
import sys
from pathlib import Path

import pytest
from docx import Document

project_root = Path(__file__).parent.parent
sys.path.insert(0, str(project_root))

from docx_tools.markdown_processor import process_markdown_content  # noqa: E402
from docx_tools.patterns import expand_br_to_block_breaks  # noqa: E402
from docx_tools.document_features import set_header_footer  # noqa: E402
from docx_tools.dynamic_docx_tools import (  # noqa: E402
    _replace_placeholders_in_document,
)


# ---------------------------------------------------------------------------
# helpers
# ---------------------------------------------------------------------------

def _render(markdown):
    """Render *markdown* through the base tool's pipeline."""
    doc = Document()
    process_markdown_content(doc, markdown)
    return doc


def _fill(value, template="{{body}}"):
    """Render *value* into a one-paragraph template through the placeholder path."""
    doc = Document()
    doc.add_paragraph(template)
    _replace_placeholders_in_document(doc, {"body": value})
    return doc


def _bodies(doc):
    return [p for p in doc.paragraphs if p.text.strip()]


def _breaks(paragraph):
    return paragraph._p.xml.count("<w:br")


def _signature(doc):
    """(style, text) for every non-empty paragraph — what the reader sees."""
    return [(p.style.name, p.text) for p in _bodies(doc)]


# ---------------------------------------------------------------------------
# the two spellings agree, on both entry points
# ---------------------------------------------------------------------------

@pytest.mark.parametrize("value", ["Line one  \nLine two", "Line one<br>Line two"])
def test_both_spellings_make_one_paragraph_with_one_break(value):
    para = _bodies(_render(value))[0]
    assert para.text == "Line one\nLine two"
    assert _breaks(para) == 1


@pytest.mark.parametrize("value", ["Line one  \nLine two", "Line one<br>Line two"])
def test_placeholder_matches_the_base_tool(value):
    assert _signature(_fill(value)) == _signature(_render(value))


def test_marker_spaces_are_not_drawn():
    # The two spaces are the marker, not content: they must not survive into the
    # run before the break (visible on centred/justified text and on copy-out).
    para = _bodies(_render("Line one  \nLine two"))[0]
    # The break itself reads back as a "\n" run; no drawn run keeps the spaces.
    assert [r.text for r in para.runs] == ["Line one", "\n", "Line two"]
    assert "  " not in para.text


def test_crlf_input_keeps_its_soft_breaks():
    para = _bodies(_render("Line one  \r\nLine two"))[0]
    assert para.text == "Line one\nLine two"
    assert _breaks(para) == 1
    assert "\r" not in para.text


def test_three_lines_make_two_breaks_in_one_paragraph():
    para = _bodies(_render("A  \nB  \nC"))[0]
    assert para.text == "A\nB\nC"
    assert _breaks(para) == 2


def test_trailing_marker_at_end_of_content_is_not_a_break():
    para = _bodies(_render("Only line  "))[0]
    assert para.text == "Only line"
    assert _breaks(para) == 0


# ---------------------------------------------------------------------------
# a soft break never swallows the block that follows it
# ---------------------------------------------------------------------------

def test_soft_break_before_list_keeps_the_first_item():
    doc = _render("Intro line  \n- item 1\n- item 2")
    assert _signature(doc) == [
        ("Normal", "Intro line"),
        ("List Bullet", "item 1"),
        ("List Bullet", "item 2"),
    ]


def test_soft_break_before_table_keeps_the_header_row():
    doc = _render("Intro  \n| A | B |\n|---|---|\n| 1 | 2 |")
    assert [p.text for p in _bodies(doc)] == ["Intro"]
    assert len(doc.tables) == 1
    assert [c.text for c in doc.tables[0].rows[0].cells] == ["A", "B"]


def test_soft_break_before_heading_renders_the_heading():
    doc = _render("Intro  \n# Heading")
    assert _signature(doc) == [("Normal", "Intro"), ("Heading 1", "Heading")]


def test_soft_break_before_page_break_does_not_print_the_dashes():
    doc = _render("Intro  \n---")
    assert [p.text for p in _bodies(doc)] == ["Intro"]


@pytest.mark.parametrize("value", ["Intro  \n> quote", "Intro<br>> quote"])
def test_soft_break_before_blockquote_keeps_the_quote(value):
    # A quote is a block like any other: the run stops before it instead of
    # printing its ">" as literal text (PR #110 review, finding 1).
    doc = _render(value)
    assert _signature(doc) == [("Normal", "Intro"), ("Quote", "quote")]


def test_soft_break_before_directive_still_applies_the_style():
    doc = _render("Intro  \n<!-- style: Quote -->\nStyled")
    assert _signature(doc) == [("Normal", "Intro"), ("Quote", "Styled")]


def test_soft_break_before_continuation_item_keeps_the_numbering():
    # "3." continues the running count even though it is not locally genuine —
    # the run must stop before it rather than absorb it as prose.
    doc = _render("1. Prvni\n\n2. Druhy\n\nPoznamka  \n3. Treti")
    numbered = [p.text for p in doc.paragraphs if p.style.name.startswith("List Number")]
    assert numbered == ["Prvni", "Druhy", "Treti"]


def test_soft_break_run_stops_at_a_blank_line():
    doc = _render("A  \n\nB")
    assert _signature(doc) == [("Normal", "A"), ("Normal", "B")]


# ---------------------------------------------------------------------------
# a block line that ends in two spaces is still a block
# ---------------------------------------------------------------------------

def test_trailing_spaces_on_a_list_item_keep_the_list():
    doc = _render("- item one  \n- item two")
    assert _signature(doc) == [
        ("List Bullet", "item one"),
        ("List Bullet", "item two"),
    ]


def test_trailing_spaces_on_a_table_row_keep_the_table():
    doc = _render("| A | B |  \n|---|---|\n| 1 | 2 |")
    assert len(doc.tables) == 1
    assert len(doc.tables[0].rows) == 2


def test_trailing_spaces_on_an_alignment_tag_still_centre_the_line():
    doc = _render("<center>Title</center>  \nnext line")
    title, following = _bodies(doc)
    assert title.text == "Title" and title.alignment is not None
    assert following.text == "next line"


def test_trailing_spaces_on_a_heading_do_not_drop_the_next_line():
    doc = _render("# Heading  \ncontinuation text")
    assert _signature(doc) == [
        ("Heading 1", "Heading"),
        ("Normal", "continuation text"),
    ]


def test_trailing_spaces_on_a_code_fence_keep_the_block_verbatim():
    doc = _render("```  \n# not a heading\n```")
    assert _signature(doc) == [("Normal", "# not a heading")]


# ---------------------------------------------------------------------------
# soft breaks work inside the other block contexts
# ---------------------------------------------------------------------------

@pytest.mark.parametrize("value", [
    "<center>\nLine one  \nLine two\n</center>",
    "<center>\nLine one<br>Line two\n</center>",
])
def test_soft_break_inside_an_alignment_block(value):
    para = _bodies(_render(value))[0]
    assert para.text == "Line one\nLine two"
    assert _breaks(para) == 1
    assert para.alignment is not None


def test_alignment_block_close_tag_is_not_swallowed():
    doc = _render("<center>\nLine one  \n</center>\nAfter")
    assert _signature(doc) == [("Normal", "Line one"), ("Normal", "After")]
    assert _bodies(doc)[1].alignment is None, "the close tag ended the block"


@pytest.mark.parametrize("value", ["> quoted one  \n> quoted two",
                                   "> quoted one<br>quoted two"])
def test_soft_break_inside_a_block_quote(value):
    para = _bodies(_render(value))[0]
    assert para.style.name == "Quote"
    assert para.text == "quoted one\nquoted two"
    assert _breaks(para) == 1


def test_br_inside_a_list_item_stays_in_the_item():
    doc = _render("- item one<br>second line\n- item two")
    assert _signature(doc) == [
        ("List Bullet", "item one\nsecond line"),
        ("List Bullet", "item two"),
    ]
    assert _breaks(_bodies(doc)[0]) == 1


def test_br_inside_a_heading_stays_in_the_heading():
    doc = _render("# Heading<br>continuation")
    assert _signature(doc) == [("Heading 1", "Heading\ncontinuation")]


# ---------------------------------------------------------------------------
# what <br> promotion deliberately does not do
# ---------------------------------------------------------------------------

def test_br_does_not_continue_a_run_that_started_earlier():
    """The documented limitation (#110): a <br>-joined number does not resume a
    numbered list that began earlier in the document, because deciding that
    would mean predicting the renderer's own count before parsing. Written with
    a real newline, the same text continues the list.
    """
    body = "1. Prvni\n\n2. Druhy\n\nPoznamka{}3. Treti"

    def numbered(markdown):
        doc = _render(markdown)
        return [p.text for p in doc.paragraphs
                if p.style.name.startswith("List Number")]

    # <br>: one paragraph, the number kept as text.
    assert _signature(_render(body.format("<br>")))[-1] == ("Normal", "Poznamka\n3. Treti")
    assert numbered(body.format("<br>")) == ["Prvni", "Druhy"]
    # real newline: the renderer's own running count continues the list.
    assert numbered(body.format("\n\n")) == ["Prvni", "Druhy", "Treti"]


@pytest.mark.parametrize("markdown,expected", [
    # The general continuation feature (#95) is untouched by any of this: a run
    # resumes across anything written between its items, written with newlines.
    ("1. Prvni\n\n2. Druhy\n\nNejaky odstavec.\n\n3. Treti",
     ["Prvni", "Druhy", "Treti"]),
    ("1. Prvni\n\n2. Druhy\n\n## II\n\n3. Treti", ["Prvni", "Druhy", "Treti"]),
    ("1. Prvni\n\n2. Druhy\n\n> Citace\n\n3. Treti", ["Prvni", "Druhy", "Treti"]),
    ("1. Prvni\n\n<center>**II.**</center>\n\n2. Druhy", ["Prvni", "Druhy"]),
    ("1. Prvni\n\n*Dukaz:*\n\n- Smlouva\n\n2. Druhy", ["Prvni", "Druhy"]),
])
def test_newline_list_continuation_still_works(markdown, expected):
    doc = _render(markdown)
    assert [p.text for p in doc.paragraphs
            if p.style.name.startswith("List Number")] == expected



_RUN_PREFIXES = [
    "", "1. A\n", "1. A\n2. B\n", "1. A\n\n2. B\n", "1. A\n5. B\n",
    "1. A\n\n5. B\n", "1. A\n2. B\n\n## H\n", "1. A\n\n23. brezna 2026\n",
    "1. A\n2. B\n\n## H\n\n23. brezna 2026\n", "```\n3. code\n```\n",
    "1. A\n   2. nested\n", "1. A\n   2. nested\n5. B\n", "- x\n- y\n",
    # Nested and same-indent bullet detours around a numbered run.
    "1. A\n  - nested bullet\n9. B\n", "1. A\n  - nested bullet\n",
    "1. A\n2. B\n\n- bullet\n",
    "- x\n   1. nested\n", "- x\n   1. nested\n5. B\n", "1. A\n      9. deep\n",
    "> q\n", "Just prose\n", "1. A\n2. B\n\n1. X\n",
    # Indented, nested and restarted lists: none of them may sway a decision
    # that is supposed to depend on the split line alone.
    "  1. A\n", "  1. A\n  2. B\n", "  1. A\n1. X\n", "1. A\n\n  1. X\n",
    "- x\n   1. n1\n   2. n2\n", "1. A\n2. B\n\nprose\n",
    # Lookalikes that never become a list: each is prose on its own.
    "4. X\n\n5. Y\n", "4. X\n\n5. Y\n\n6. Z\n", "9. X\n",
    "23. brezna 2026\n\n24. dubna 2026\n", "1. A\n2. B\n\n9. X\n\n10. Y\n",
]


@pytest.mark.parametrize("prefix", _RUN_PREFIXES)
@pytest.mark.parametrize("number", [1, 2, 5, 6, 9, 10, 11, 24, 25])
def test_a_br_line_is_promoted_only_when_the_number_becomes_an_item(prefix, number):
    """Promotion and rendering must agree.

    Splitting a <br> line is only right if the segment really renders as a list
    item; if it does not, the line must stay one paragraph with a soft break.
    A promotion that no list takes turns a soft break into a paragraph break —
    the failure mode a count-only check does not see, because no list item
    appears either way.
    """
    markdown = prefix + f"\nNote<br>{number}. Tail"
    promoted = expand_br_to_block_breaks(markdown).splitlines()[-1].strip() == f"{number}. Tail"
    signature = _signature(_render(markdown))
    became_item = any(style.startswith("List Number") and text == "Tail"
                      for style, text in signature)
    assert promoted == became_item
    if not promoted:
        assert any(text == f"Note\n{number}. Tail" for _, text in signature)


def test_br_before_a_list_is_still_promoted():
    # The rescue that the narrowed promotion rule must keep: a list a model
    # separated from its lead-in with <br> is still a list.
    doc = _render("Intro:<br>1. A<br>2. B")
    assert _signature(doc) == [
        ("Normal", "Intro:"),
        ("List Number", "A"),
        ("List Number", "B"),
    ]


def test_nbsp_before_a_break_is_content_and_survives():
    # Only the marker characters are stripped from a breaking segment; a
    # deliberate &nbsp; is content (PR #110 review, finding 3).
    doc = Document()
    para = doc.add_paragraph()
    from docx_tools.inline_formatting import parse_inline_formatting
    parse_inline_formatting("a&nbsp;<br>b", para)
    assert [r.text for r in para.runs] == ["a\u00a0", "\n", "b"]


# ---------------------------------------------------------------------------
# table cells: <br> is a break, <br><br> is a paragraph
# ---------------------------------------------------------------------------

def test_cell_br_is_a_soft_break():
    doc = _render("| A | B |\n|---|---|\n| Street<br>City | x |")
    cell = doc.tables[0].rows[1].cells[0]
    assert [p.text for p in cell.paragraphs] == ["Street\nCity"]


def test_cell_double_br_is_a_new_paragraph():
    doc = _render("| A | B |\n|---|---|\n| First<br><br>Second | x |")
    cell = doc.tables[0].rows[1].cells[0]
    assert [p.text for p in cell.paragraphs] == ["First", "Second"]


def test_cell_and_cell_placeholder_agree():
    # The same address, written the same way, in a markdown cell and in a
    # template placeholder that sits in a cell.
    doc = _render("| A | B |\n|---|---|\n| Street<br>City | x |")
    md_cell = doc.tables[0].rows[1].cells[0]

    tmpl = Document()
    table = tmpl.add_table(rows=1, cols=1)
    table.cell(0, 0).paragraphs[0].add_run("{{body}}")
    _replace_placeholders_in_document(tmpl, {"body": "Street<br>City"})
    tmpl_cell = table.cell(0, 0)

    assert [p.text for p in md_cell.paragraphs] == [p.text for p in tmpl_cell.paragraphs]


# ---------------------------------------------------------------------------
# header / footer take the same spelling
# ---------------------------------------------------------------------------

@pytest.mark.parametrize("text", ["Line one<br>Line two", "Line one\nLine two",
                                  r"Line one\nLine two"])
def test_header_breaks_the_line(text):
    doc = Document()
    set_header_footer(doc, text, 'header')
    para = doc.sections[0].header.paragraphs[0]
    assert para.text == "Line one\nLine two"
    assert _breaks(para) == 1
    assert "<br" not in para.text


def test_footer_breaks_the_line():
    doc = Document()
    set_header_footer(doc, "Left<br>Right", 'footer')
    para = doc.sections[0].footer.paragraphs[0]
    assert para.text == "Left\nRight"


def test_header_page_field_still_works_alongside_a_break():
    doc = Document()
    set_header_footer(doc, "Title<br>Page {page} of {pages}", 'header')
    xml = doc.sections[0].header.paragraphs[0]._p.xml
    assert "PAGE" in xml and "NUMPAGES" in xml
    assert "<w:br" in xml


# ---------------------------------------------------------------------------
# the ambiguity that stays: a line-leading number is escaped, both spellings
# ---------------------------------------------------------------------------

@pytest.mark.parametrize("value", [
    "Advokat  \n1\\. maje 5  \n110 00 Praha",
    "Advokat<br>1\\. maje 5<br>110 00 Praha",
])
def test_escaped_line_leading_number_keeps_the_block_together(value):
    # An address line that looks like a day-1 date reads as a list item unless
    # its dot is escaped — the documented escape, and it now works identically
    # for both soft-break spellings.
    para = _bodies(_render(value))[0]
    assert para.text == "Advokat\n1. maje 5\n110 00 Praha"
    assert _breaks(para) == 2
