"""A title on a layout that has no title placeholder (#118).

`blank` is the case that always hit it: a Blank layout has no title
placeholder by definition, so every blank slide carrying a title warned and
lost it. The title is now drawn as a text box at the position — and in the
style — the template gives its titles, read from another layout.
"""
import sys
from pathlib import Path

project_root = Path(__file__).parent.parent
sys.path.insert(0, str(project_root))

import pytest
from pptx import Presentation as PptxReader
from pptx.oxml.ns import qn
from pptx.util import Inches

from pptx_tools.layouts import LayoutResolver
from pptx_tools.placeholder_style import read_title_style, title_placeholder
from pptx_tools.slide_builder import PowerpointPresentation
from pptx_tools.templates import TemplateSpec

BASE_16_9 = project_root / "default_templates" / "default_pptx_template_16_9.pptx"


def build(**slide):
    slide.setdefault("type", "blank")
    slide.setdefault("elements", [{"kind": "text", "text": "body", "x": 1, "y": 3, "w": 3}])
    pres = PowerpointPresentation([slide], "16:9")
    return pres, PptxReader(pres.save()).slides[0]


def title_boxes(slide, text):
    return [shape for shape in slide.shapes
            if shape.has_text_frame and shape.text_frame.text == text]


def strip_titles(prs):
    """Remove every title placeholder from the layouts and the master."""
    for container in [prs.slide_master, *prs.slide_layouts]:
        while True:
            placeholder = title_placeholder(container)
            if placeholder is None:
                break
            element = placeholder._element
            element.getparent().remove(element)


class TestTheTitleSurvives:

    def test_a_blank_slide_keeps_its_title(self):
        pres, slide = build(title="Kept")
        assert len(title_boxes(slide, "Kept")) == 1
        assert pres.warnings == []

    def test_no_title_draws_no_box(self):
        _, slide = build()
        assert [shape.text_frame.text for shape in slide.shapes if shape.has_text_frame] == ["body"]

    def test_an_empty_title_draws_no_box(self):
        _, slide = build(title="")
        assert [shape.text_frame.text for shape in slide.shapes if shape.has_text_frame] == ["body"]

    def test_the_title_is_drawn_first_so_elements_sit_on_top(self):
        _, slide = build(title="Kept", elements=[
            {"kind": "shape", "x": 0, "y": 0, "w": 2, "h": 2}])
        assert slide.shapes[0].text_frame.text == "Kept"


class TestItLooksLikeATitle:

    def test_it_lands_where_the_content_layout_puts_its_title(self):
        """Not a guessed band: the template's own title geometry."""
        reference = title_placeholder(PptxReader(str(BASE_16_9)).slide_layouts[1])
        _, slide = build(title="Kept")
        box = title_boxes(slide, "Kept")[0]
        assert (box.left, box.top, box.width, box.height) == (
            reference.left, reference.top, reference.width, reference.height)

    def test_it_inherits_the_template_title_font_and_theme_colour(self):
        _, slide = build(title="Kept")
        rPr = title_boxes(slide, "Kept")[0].text_frame.paragraphs[0].runs[0]._r.find(qn('a:rPr'))
        assert rPr.get('sz') == '4400'
        assert rPr.find(qn('a:latin')).get('typeface') == '+mj-lt'
        assert rPr.find(qn('a:solidFill')).find(qn('a:schemeClr')).get('val') == 'tx1'

    def test_every_line_of_a_multi_line_title_is_styled(self):
        """`TextFrame.text` splits on a newline; an unstyled second line shows."""
        _, slide = build(title="First line\nSecond line")
        box = [shape for shape in slide.shapes if shape.name == "Title"][0]
        sizes = [run._r.find(qn('a:rPr')).get('sz')
                 for paragraph in box.text_frame.paragraphs for run in paragraph.runs]
        assert sizes == ['4400', '4400']

    def test_a_long_title_shrinks_instead_of_resizing_the_box(self):
        """add_textbox() defaults to grow-to-fit; a title band is fixed."""
        _, slide = build(title="Kept " * 40)
        box = [shape for shape in slide.shapes if shape.name == "Title"][0]
        bodyPr = box.text_frame._txBody.find(qn('a:bodyPr'))
        assert bodyPr.find(qn('a:normAutofit')) is not None
        assert bodyPr.find(qn('a:spAutoFit')) is None


class TestStyleLookup:

    def test_a_layout_without_a_title_reads_as_none(self):
        prs = PptxReader(str(BASE_16_9))
        assert read_title_style(prs.slide_layouts[6]) is None

    def test_the_resolver_prefers_a_body_layout_over_the_cover(self):
        prs = PptxReader(str(BASE_16_9))
        style = LayoutResolver(prs).title_style()
        assert style.source == prs.slide_layouts[1].name

    def test_the_resolver_caches_the_answer(self):
        resolver = LayoutResolver(PptxReader(str(BASE_16_9)))
        assert resolver.title_style() is resolver.title_style()


class TestTemplateWithNoTitleAnywhere:

    @pytest.fixture
    def untitled_template(self, tmp_path):
        prs = PptxReader(str(BASE_16_9))
        strip_titles(prs)
        path = tmp_path / "untitled.pptx"
        prs.save(str(path))
        return path

    def test_the_title_is_still_drawn_and_the_guess_is_reported(self, untitled_template):
        pres = PowerpointPresentation(
            [{"type": "blank", "title": "Kept",
              "elements": [{"kind": "text", "text": "body", "x": 1, "y": 3, "w": 3}]}],
            "16:9",
            template_spec=TemplateSpec(name="untitled", path=untitled_template, aspect="16:9"),
        )
        assert any("no layout in this template has one" in w for w in pres.warnings)
        slide = PptxReader(pres.save()).slides[0]
        box = title_boxes(slide, "Kept")[0]
        assert box.top == Inches(0.3)

    def test_the_resolver_reports_no_style(self, untitled_template):
        prs = PptxReader(str(untitled_template))
        assert LayoutResolver(prs).title_style() is None
