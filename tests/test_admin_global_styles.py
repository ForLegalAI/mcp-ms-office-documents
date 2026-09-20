"""Editing the ``style_mapping`` that applies to every Word document (#161).

The setting with the widest blast radius on the server was the one the UI could
only read. Making it editable is not just a form: the mapping has three
consumers that must never disagree — the static ``markdown_to_word`` tool, every
dynamic Word template tool, and the preview — and two of them cached it in ways
that would have left an edit half-applied.

What the tests here are really pinning:

* a save reaches a **document**, not just a YAML file;
* the master ``docx_templates.yaml`` is never rewritten, and reverting brings
  it straight back;
* the reserved ``_global.yaml`` is not mistaken for a template, in either
  direction; and
* the page tells the truth about which of the two layers is in force.
"""
import io
import re
import sys
from pathlib import Path

project_root = Path(__file__).parent.parent
sys.path.insert(0, str(project_root))

import pytest
import yaml
from docx import Document
from starlette.testclient import TestClient

import admin.store as store_mod
import metrics
import template_utils as tu
from admin.store import TemplateStoreError, validate_name
from docx_tools.dynamic_docx_tools import (
    register_docx_template_tools_from_yaml, registered_docx_template_names,
    unregister_docx_template,
)
from admin.views.settings import UNSET
from config import Config
from template_registry import (
    GLOBAL_SETTINGS_FILENAME, gather_specs, global_config, read_spec_dir,
    spec_dir_for,
)

MASTER = "docx_templates.yaml"
SPEC_DIR = "docx_templates.d"


def _docx_bytes() -> bytes:
    doc = Document()
    doc.add_paragraph("Hi")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# The merge layer
# ---------------------------------------------------------------------------


def _write(path: Path, data) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(yaml.safe_dump(data), encoding="utf-8")
    return path


def test_the_d_file_overrides_the_master_key(tmp_path):
    master = _write(tmp_path / MASTER, {"style_mapping": {"heading_1": "Master H1"}})
    _write(tmp_path / SPEC_DIR / GLOBAL_SETTINGS_FILENAME,
           {"style_mapping": {"heading_1": "Managed H1"}})

    cfg = global_config(master, spec_dir_for(master))
    assert cfg["style_mapping"] == {"heading_1": "Managed H1"}


def test_the_override_replaces_the_mapping_rather_than_merging_into_it(tmp_path):
    """The rule the whole feature rests on.

    A key-wise merge would make a cleared key silently fall back to the
    master's value — the page would show "(use built-in)" while the master's
    name was still being applied, which is the exact lie #161 is about.
    """
    master = _write(tmp_path / MASTER,
                    {"style_mapping": {"heading_1": "Master H1", "quote": "Master Q"}})
    _write(tmp_path / SPEC_DIR / GLOBAL_SETTINGS_FILENAME,
           {"style_mapping": {"heading_1": "Managed H1"}})

    mapping = global_config(master, spec_dir_for(master))["style_mapping"]
    assert mapping == {"heading_1": "Managed H1"}, "quote must not survive"


def test_an_empty_stored_mapping_turns_the_masters_off(tmp_path):
    """Stored-and-empty and not-stored are different states.

    Clearing every dropdown has to mean "built-in", so the file is written
    with an explicit empty mapping rather than deleted.
    """
    master = _write(tmp_path / MASTER, {"style_mapping": {"heading_1": "Master H1"}})
    _write(tmp_path / SPEC_DIR / GLOBAL_SETTINGS_FILENAME, {"style_mapping": {}})

    assert global_config(master, spec_dir_for(master))["style_mapping"] == {}


def test_with_no_stored_file_the_master_applies(tmp_path):
    master = _write(tmp_path / MASTER, {"style_mapping": {"heading_1": "Master H1"}})
    assert global_config(master, spec_dir_for(master))["style_mapping"] == \
        {"heading_1": "Master H1"}


def test_the_override_cannot_inject_templates(tmp_path):
    """One file able to register tools by a back door is not a power this layer
    should have: `templates` belongs to the per-template files."""
    master = _write(tmp_path / MASTER, {"templates": [{"name": "real"}]})
    _write(tmp_path / SPEC_DIR / GLOBAL_SETTINGS_FILENAME,
           {"templates": [{"name": "smuggled", "docx_path": "x.docx"}]})

    specs, cfg = gather_specs(master, tmp_path / SPEC_DIR)
    assert [s["name"] for s in specs] == ["real"]
    assert [t["name"] for t in cfg.get("templates", [])] == ["real"]


def test_the_reserved_file_is_not_read_as_a_template(tmp_path):
    """It has no `name`, so reading it as a spec would log it as malformed on
    every load — and a `name` in it would register a tool."""
    d = tmp_path / SPEC_DIR
    _write(d / GLOBAL_SETTINGS_FILENAME,
           {"name": "sneaky", "docx_path": "x.docx", "style_mapping": {}})
    _write(d / "real.yaml", {"name": "real", "docx_path": "r.docx"})

    assert [s["name"] for s in read_spec_dir(d)] == ["real"]


def test_a_template_cannot_be_named_global():
    """`_NAME_RE` allows a leading underscore, so `_global` would be written to
    `_global.yaml` and take the server's global settings with it."""
    with pytest.raises(TemplateStoreError, match="reserved"):
        validate_name("_global")


def test_no_store_operation_can_reach_the_reserved_file(tmp_path):
    """One guard, but it has to cover every door.

    `validate_name` is called from `_spec_path`, which every read and write of
    a managed spec goes through — so this passes for the same reason each of
    them does, and fails the day one of them builds a path itself.
    """
    from admin.store import KIND_DOCX, FileTemplateStore

    store = FileTemplateStore(custom_dir=tmp_path / "c", config_dir=tmp_path / "g")
    store.save_global_settings(KIND_DOCX, {"style_mapping": {"heading_1": "Brand"}})
    store.save_spec(KIND_DOCX, {"name": "real", "docx_path": "real.docx"},
                    asset_bytes=_docx_bytes())

    doors = {
        "save_spec": lambda: store.save_spec(
            KIND_DOCX, {"name": "_global", "docx_path": "x.docx"},
            asset_bytes=_docx_bytes()),
        "rename_spec": lambda: store.rename_spec(KIND_DOCX, "real", "_global"),
        "clone_spec": lambda: store.clone_spec(KIND_DOCX, "real", "_global"),
        "set_enabled": lambda: store.set_enabled(KIND_DOCX, "_global", False),
        "get_spec": lambda: store.get_spec(KIND_DOCX, "_global"),
        "delete_spec": lambda: store.delete_spec(KIND_DOCX, "_global"),
    }
    for label, call in doors.items():
        with pytest.raises(TemplateStoreError, match="reserved"):
            call()

    assert store.global_settings(KIND_DOCX) == {"style_mapping": {"heading_1": "Brand"}}
    assert [s["name"] for s in store.list_specs(KIND_DOCX)] == ["real"]


def test_a_malformed_override_does_not_take_the_config_with_it(tmp_path, caplog):
    master = _write(tmp_path / MASTER, {"style_mapping": {"heading_1": "Master H1"}})
    path = tmp_path / SPEC_DIR / GLOBAL_SETTINGS_FILENAME
    path.parent.mkdir(parents=True)
    path.write_text("[not, a, mapping]", encoding="utf-8")

    assert global_config(master, spec_dir_for(master))["style_mapping"] == \
        {"heading_1": "Master H1"}


# ---------------------------------------------------------------------------
# The static Word tool reads it live
# ---------------------------------------------------------------------------


@pytest.fixture
def config_dir(tmp_path, monkeypatch):
    """Point docx_tools.style_map at a scratch config directory."""
    import docx_tools.style_map as sm

    cfg = tmp_path / "config"
    cfg.mkdir()
    monkeypatch.setattr(sm, "_CONFIG_PATHS", (cfg / MASTER,))
    return cfg


def test_the_static_word_tool_follows_an_edit_without_a_restart(config_dir):
    """It cached the map for the life of the process.

    So the one setting that restyles every document could be changed in the UI
    and go on producing documents in the old style until someone restarted the
    server — with the page showing the new value the whole time.
    """
    from docx_tools.style_map import load_global_style_map

    _write(config_dir / MASTER, {"style_mapping": {"heading_1": "Master H1"}})
    assert load_global_style_map().heading_style(1) == "Master H1"

    _write(config_dir / SPEC_DIR / GLOBAL_SETTINGS_FILENAME,
           {"style_mapping": {"heading_1": "Managed H1"}})
    assert load_global_style_map().heading_style(1) == "Managed H1", \
        "the second read must not come from a cache"

    (config_dir / SPEC_DIR / GLOBAL_SETTINGS_FILENAME).unlink()
    assert load_global_style_map().heading_style(1) == "Master H1"


def test_editing_a_stored_mapping_in_place_is_picked_up(config_dir):
    """Content changing under an unchanged filename.

    Creating and deleting the file moves its *existence*, which a fingerprint
    could notice while still being blind to an edit. This one rewrites the
    same file to the same length, so only reading its mtime catches it.
    """
    from docx_tools.style_map import invalidate_global_style_map, load_global_style_map

    path = config_dir / SPEC_DIR / GLOBAL_SETTINGS_FILENAME
    _write(path, {"style_mapping": {"quote": "AAAA"}})
    invalidate_global_style_map()
    assert load_global_style_map().quote == "AAAA"

    before = path.stat().st_size
    _write(path, {"style_mapping": {"quote": "BBBB"}})
    assert path.stat().st_size == before, "same length, so only mtime differs"

    assert load_global_style_map().quote == "BBBB"


def test_a_stored_mapping_applies_with_no_master_yaml_at_all(config_dir):
    """Resolution used to stop at the first *existing* master file; a
    deployment that never had one would ignore the managed mapping."""
    from docx_tools.style_map import load_global_style_map

    _write(config_dir / SPEC_DIR / GLOBAL_SETTINGS_FILENAME,
           {"style_mapping": {"quote": "Managed Quote"}})
    assert load_global_style_map().quote == "Managed Quote"


def test_an_unchanged_config_is_not_reparsed(config_dir, monkeypatch):
    """The cache has to actually cache, or `invalidate_global_style_map` is
    dead weight and every document pays for parsing the master file.

    Measured on the shipped `config/docx_templates.yaml` — a few hundred lines
    of worked examples — a full resolution is ~8 ms against ~0.02 ms for a hit.
    """
    import docx_tools.style_map as sm

    _write(config_dir / MASTER, {"style_mapping": {"heading_1": "Master H1"}})
    sm.invalidate_global_style_map()
    sm.load_global_style_map()

    parses = []
    real = sm.global_config
    monkeypatch.setattr(sm, "global_config",
                        lambda m, d: (parses.append(1), real(m, d))[1])

    assert sm.load_global_style_map().heading_style(1) == "Master H1"
    assert parses == [], "an unchanged config must not be re-parsed"

    sm.invalidate_global_style_map()
    sm.load_global_style_map()
    assert parses == [1], "and invalidating must make it re-parse"


def test_no_config_at_all_gives_the_built_in_map(config_dir):
    from docx_tools.style_map import DEFAULT_STYLE_MAP, load_global_style_map

    assert load_global_style_map() == DEFAULT_STYLE_MAP


# ---------------------------------------------------------------------------
# The page
# ---------------------------------------------------------------------------


@pytest.fixture
def admin_client(tmp_path, monkeypatch):
    custom = tmp_path / "custom"
    cfg = tmp_path / "config"
    custom.mkdir()
    cfg.mkdir()
    monkeypatch.setattr(store_mod, "_APP_CUSTOM_DIR", tmp_path / "nx1")
    monkeypatch.setattr(store_mod, "_APP_CONFIG_DIR", tmp_path / "nx2")
    monkeypatch.setattr(store_mod, "_LOCAL_CUSTOM_DIR", custom)
    monkeypatch.setattr(store_mod, "_LOCAL_CONFIG_DIR", cfg)
    monkeypatch.setattr(tu, "APP_CUSTOM_DIR", tmp_path / "nx1")
    monkeypatch.setattr(tu, "LOCAL_CUSTOM_DIR", custom)
    import docx_tools.style_map as sm
    monkeypatch.setattr(sm, "_CONFIG_PATHS", (cfg / MASTER,))
    import pptx_tools.templates as templates_mod
    monkeypatch.setattr(templates_mod, "_APP_CONFIG_DIR", tmp_path / "nx2")
    monkeypatch.setattr(templates_mod, "_LOCAL_CONFIG_DIR", cfg)
    templates_mod.clear_cache()
    monkeypatch.setenv("ADMIN_ENABLED", "true")
    monkeypatch.setenv("ADMIN_PASSWORD", "pw")
    monkeypatch.delenv("API_KEY", raising=False)

    from fastmcp import FastMCP
    from admin.app import build_combined_app

    mcp = FastMCP("test-styles")
    client = TestClient(build_combined_app(mcp, Config.from_env()))
    client.mcp = mcp          # the routes' AdminContext registers against this one
    metrics.reset()
    client.__enter__()
    client.post("/admin/login", data={"password": "pw"})
    yield client, custom, cfg
    client.__exit__(None, None, None)
    for name in list(registered_docx_template_names()):
        unregister_docx_template(mcp, name)
    metrics.reset()


def _csrf(client) -> str:
    html = client.get("/admin/styles").text
    tag = re.search(r'<input[^>]*name="csrf"[^>]*>', html)
    return re.search(r'value="([^"]*)"', tag.group(0)).group(1)


def _post(client, url, data=None):
    payload = dict(data or {})
    payload["csrf"] = _csrf(client)
    return client.post(url, data=payload)


def _stored(cfg: Path):
    path = cfg / SPEC_DIR / GLOBAL_SETTINGS_FILENAME
    return yaml.safe_load(path.read_text(encoding="utf-8")) if path.is_file() else None


def test_the_page_is_reachable_and_linked(admin_client):
    client, _custom, _cfg = admin_client
    assert 'href="/admin/styles"' in client.get("/admin/").text, "linked from the nav"
    assert client.get("/admin/styles").status_code == 200


def test_saving_writes_the_managed_file_and_leaves_the_master_alone(admin_client):
    client, _custom, cfg = admin_client
    master = _write(cfg / MASTER, {"style_mapping": {"heading_1": "Master H1"}})
    before = master.read_text(encoding="utf-8")

    r = _post(client, "/admin/styles/save",
              {"style_heading_1": "Brand Title", "style_quote": "__default__"})

    assert r.status_code == 200
    assert _stored(cfg) == {"style_mapping": {"heading_1": "Brand Title"}}
    assert master.read_text(encoding="utf-8") == before, \
        "the hand-written master file is never rewritten"


def test_a_saved_mapping_reaches_a_generated_document(admin_client):
    """The assertion that matters: not that a YAML key moved, but that the
    next Word document comes out styled the new way."""
    from docx_tools.base_docx_tool import _markdown_to_word_buffer

    client, custom, _cfg = admin_client
    _install_base_docx(custom, styles=["Brand Title"])

    _post(client, "/admin/styles/save", {"style_heading_1": "Brand Title"})

    buf, _warnings = _markdown_to_word_buffer("# Hello")
    doc = Document(buf)
    assert [p.style.name for p in doc.paragraphs if p.text == "Hello"] == ["Brand Title"]


def test_reverting_brings_the_master_back(admin_client):
    client, _custom, cfg = admin_client
    _write(cfg / MASTER, {"style_mapping": {"heading_1": "Master H1"}})
    _post(client, "/admin/styles/save", {"style_heading_1": "Brand Title"})
    assert _stored(cfg) is not None

    r = _post(client, "/admin/styles/revert", {})

    assert _stored(cfg) is None
    assert "Master H1" in r.text, "the master's mapping is on the page again"


def test_clearing_every_key_is_stored_rather_than_reverted(admin_client):
    """Otherwise "set nothing" and "revert" would be the same button, and an
    admin who cleared a key would silently get the master's value back."""
    client, _custom, cfg = admin_client
    _write(cfg / MASTER, {"style_mapping": {"heading_1": "Master H1"}})

    _post(client, "/admin/styles/save", {"style_heading_1": "__default__"})

    assert _stored(cfg) == {"style_mapping": {}}
    from docx_tools.style_map import load_global_style_map
    assert load_global_style_map().heading_style(1) == "Heading 1"


def _selected(html: str, key: str) -> str:
    """The value the ``style_<key>`` select is currently showing.

    Attribute order is FastHTML's business and it does not put ``value`` and
    ``selected`` in the order a hand-written assertion would guess, so this
    reads the select out rather than matching a fixed string.
    """
    m = re.search(rf'<select[^>]*name="style_{key}"[^>]*>(.*?)</select>', html, re.S)
    assert m, f"no select for {key}"
    for opt in re.findall(r"<option[^>]*>", m.group(1)):
        if re.search(r"\bselected\b", opt):
            value = re.search(r'value="([^"]*)"', opt).group(1)
            # The sentinel the form posts for "do not set this key"; reported
            # as "" so a test says what it means.
            return "" if value == UNSET else value
    raise AssertionError(f"no option is selected for {key}")


def test_the_page_shows_what_is_in_force_not_an_empty_form(admin_client):
    """An empty form here would read as "nothing is set" on a server where a
    global mapping has been in force all along."""
    client, _custom, cfg = admin_client
    _write(cfg / MASTER, {"style_mapping": {"heading_1": "Master H1",
                                            "quote": "Master Q"}})

    html = client.get("/admin/styles").text

    assert _selected(html, "heading_1") == "Master H1"
    assert _selected(html, "quote") == "Master Q"
    assert _selected(html, "heading_2") == "", "an unset key stays unset"


def test_a_master_key_the_override_drops_is_named_on_the_page(admin_client):
    """The override replaces the master's mapping, so a key it does not carry
    stops applying. An admin reading docx_templates.yaml on the volume would
    otherwise see a setting that is simply not in force."""
    client, _custom, cfg = admin_client
    _write(cfg / MASTER,
           {"style_mapping": {"heading_1": "Master H1", "quote": "Master Q"}})

    _post(client, "/admin/styles/save", {"style_heading_1": "Master H1"})
    html = client.get("/admin/styles").text

    assert "Not in force" in html
    assert "Master Q" in html, "the dropped key is named"


def test_a_configured_style_missing_from_the_base_template_is_kept_and_flagged(admin_client):
    """A value already in the config must never be dropped for being
    unrecognised — that would lose configuration by rendering a page."""
    client, custom, cfg = admin_client
    _install_base_docx(custom, styles=[])
    _write(cfg / MASTER, {"style_mapping": {"heading_1": "Ghost Style"}})

    html = client.get("/admin/styles").text

    assert 'value="Ghost Style"' in html
    assert _selected(html, "heading_1") == "Ghost Style", "and still selected"
    assert "not in the base Word template" in html


def test_the_marker_does_not_claim_more_than_it_knows(admin_client):
    """The dropdowns list the *base* template's styles, because that is what
    `markdown_to_word` renders onto. The mapping also reaches every Word
    template tool, each rendering onto its own document — which may well define
    the style. Reading the marker as "missing everywhere" is the shape of
    wrongness #161 is about, so the page says which it means."""
    client, custom, cfg = admin_client
    _install_base_docx(custom, styles=[])
    _write(cfg / MASTER, {"style_mapping": {"heading_1": "Ghost Style"}})

    html = client.get("/admin/styles").text

    assert "still uses it" in html, "a marked style is not necessarily wrong"
    assert "Ghost Style" in re.search(r"Marked styles.*?</p>", html, re.S).group(0)


def test_no_note_about_marked_styles_when_nothing_is_marked(admin_client):
    """A note that is always there is a note nobody reads."""
    client, custom, cfg = admin_client
    _install_base_docx(custom, styles=[])
    _write(cfg / MASTER, {"style_mapping": {"heading_1": "Heading 2"}})

    assert "Marked styles" not in client.get("/admin/styles").text


def test_a_key_the_renderer_ignores_is_flagged_and_stays_flagged(admin_client):
    """It does nothing — `style_map._normalize` drops it with a log — but it is
    in a file somebody wrote on purpose, and a page claiming to show what is
    set has to account for it rather than leave it to a log nobody reads.

    Still flagged after a save, because the master file still contains it: the
    override stops it being in the effective mapping, not on the volume.
    """
    client, _custom, cfg = admin_client
    _write(cfg / MASTER, {"style_mapping": {"heading_1": "H1", "sidebar": "Nope"}})

    assert "sidebar" in client.get("/admin/styles").text

    _post(client, "/admin/styles/save", {"style_heading_1": "H1"})
    assert "sidebar" in client.get("/admin/styles").text, \
        "it is still sitting in docx_templates.yaml looking like a setting"


def test_a_key_the_renderer_ignores_is_not_called_switched_off(admin_client):
    """"Not in force" means this page turned something off. A key the renderer
    never acted on was not in force to begin with, and listing it beside a real
    one invites an admin to go looking for the setting they lost."""
    client, _custom, cfg = admin_client
    _write(cfg / MASTER, {"style_mapping": {"heading_1": "Master H1", "sidebar": "Nope"}})

    _post(client, "/admin/styles/save", {"style_heading_1": "__default__"})
    html = client.get("/admin/styles").text

    row = re.search(r"Not in force.*?</div>", html, re.S)
    assert row and "Master H1" in row.group(0), "the real one is named"
    assert "sidebar" not in row.group(0)


def test_save_requires_the_csrf_token(admin_client):
    client, _custom, cfg = admin_client
    r = client.post("/admin/styles/save",
                    data={"style_heading_1": "Brand Title", "csrf": "wrong"})
    assert _stored(cfg) is None, f"saved anyway (status {r.status_code})"


def test_the_styles_routes_are_not_swallowed_by_the_kind_routes(admin_client):
    """`/styles/save` also fits `/{kind}/save`. Starlette matches in
    registration order, so this passes only while /styles/… is registered
    first — which is why the ordering carries a comment in admin/app.py."""
    client, _custom, cfg = admin_client
    r = _post(client, "/admin/styles/save", {"style_quote": "__default__"})
    assert r.status_code == 200
    assert _stored(cfg) == {"style_mapping": {}}, \
        "the save route ran, rather than /{kind}/save rejecting kind='styles'"


# ---------------------------------------------------------------------------
# Every consumer moves together
# ---------------------------------------------------------------------------


def _install_base_docx(custom: Path, styles=()):
    """Write a base Word template defining *styles* on top of the built-ins."""
    from admin.base_templates import slot
    from template_utils import find_file_in_template_dirs

    source = find_file_in_template_dirs("default_docx_template.docx")
    doc = Document(str(source)) if source else Document()
    for name in styles:
        try:
            doc.styles.add_style(name, 1)  # WD_STYLE_TYPE.PARAGRAPH
        except ValueError:
            pass
    buf = io.BytesIO()
    doc.save(buf)
    (custom / slot("docx").custom_name).write_bytes(buf.getvalue())


def test_saving_re_registers_the_dynamic_word_tools(admin_client, monkeypatch):
    """The global mapping is baked into each template tool at registration
    time, so without a re-registration a save would move the static tool and
    leave every template tool on the old mapping — a split an admin would
    never think to suspect."""
    client, custom, cfg = admin_client
    _install_base_docx(custom, styles=["Brand Title"])

    doc = Document()
    doc.add_paragraph("Hi {{who}}")
    buf = io.BytesIO()
    doc.save(buf)
    (custom / "letter.docx").write_bytes(buf.getvalue())
    _write(cfg / SPEC_DIR / "letter.yaml",
           {"name": "letter", "docx_path": "letter.docx",
            "args": [{"name": "who", "type": "string", "required": True}]})

    import docx_tools.dynamic_docx_tools as dyn
    seen = []
    real = dyn.register_docx_template_tools_from_yaml
    monkeypatch.setattr(dyn, "register_docx_template_tools_from_yaml",
                        lambda mcp, path: (seen.append(path), real(mcp, path))[1])

    _post(client, "/admin/styles/save", {"style_heading_1": "Brand Title"})

    assert seen, "the save must re-register the dynamic Word tools"
    assert seen[0].name == MASTER


def test_a_save_applies_even_when_the_files_look_unchanged(admin_client, monkeypatch):
    """The reason the save path invalidates explicitly rather than trusting
    the fingerprint.

    The cache is keyed on the config files' mtime and size. Two writes inside
    one filesystem timestamp tick that leave the size unchanged — a
    coarse-granularity mount, or simply a fast admin swapping one six-letter
    style name for another — would look identical, and the static Word tool
    would go on rendering with the old mapping. Freezing the fingerprint here
    is that filesystem, deterministically.
    """
    import docx_tools.style_map as sm

    client, _custom, cfg = admin_client
    _write(cfg / MASTER, {"style_mapping": {"heading_1": "Master"}})

    # Frozen *before* the first read, so the cached entry carries the frozen
    # fingerprint and the read after the save is a cache hit. Freezing it
    # afterwards would change the fingerprint and force a miss — which is the
    # bug passing itself off as the fix.
    monkeypatch.setattr(sm, "_stamp", lambda path: ("frozen",))
    sm.invalidate_global_style_map()
    assert sm.load_global_style_map().heading_style(1) == "Master"

    _post(client, "/admin/styles/save", {"style_heading_1": "Brandy"})

    assert sm.load_global_style_map().heading_style(1) == "Brandy"


def test_a_template_lost_to_the_re_registration_is_reported(admin_client):
    """Saving a style must not take a tool off the server in silence.

    `register_docx_template()` removes the existing tool *before* rebuilding
    it, and the loop logs a failure and moves on — so a template whose source
    file went missing since startup is gone, and from the admin's side the
    only thing that happened was saving a style mapping.
    """
    client, custom, cfg = admin_client
    (custom / "letter.docx").write_bytes(_docx_bytes())
    _write(cfg / SPEC_DIR / "letter.yaml",
           {"name": "letter", "docx_path": "letter.docx", "args": []})
    register_docx_template_tools_from_yaml(client.mcp, cfg / MASTER)
    assert "letter" in registered_docx_template_names()

    (custom / "letter.docx").unlink()   # deleted on the volume since startup

    r = _post(client, "/admin/styles/save", {"style_heading_1": "Heading 2"})

    assert "letter" not in registered_docx_template_names(), \
        "it really is off the server, not merely stale"
    assert "letter" in r.text and "no longer registered" in r.text
    assert 'class="flash flash-warn"' in r.text, "a lost tool is not a success"


def test_the_page_and_the_renderer_agree_on_what_is_in_force(admin_client):
    """One resolution path, so the card cannot describe a mapping the next
    document will not use."""
    client, _custom, cfg = admin_client
    _write(cfg / MASTER, {"style_mapping": {"heading_1": "Master H1"}})
    _post(client, "/admin/styles/save", {"style_heading_1": "Brand Title"})

    from docx_tools.style_map import load_global_style_map
    from admin.app import AdminContext
    from fastmcp import FastMCP

    ctx = AdminContext(FastMCP("x"), Config.from_env())
    assert ctx.global_style_mapping == {"heading_1": "Brand Title"}
    assert load_global_style_map().heading_style(1) == "Brand Title"
