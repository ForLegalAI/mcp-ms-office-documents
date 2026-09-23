"""Control characters in caller text no longer take the whole deck down.

Regression from the PowerPoint review: a BEL or a vertical tab (Office's own
soft line break, common in pasted text) in ``footer_text``, a section title or
a chart category raised ``XMLSyntaxError`` — those three are written without
python-pptx's escaping — and the caller got an error instead of a deck. Every
caller string is now cleaned once, after validation, and the removal is
reported as one info warning.
"""
import sys
from pathlib import Path

project_root = Path(__file__).parent.parent
sys.path.insert(0, str(project_root))

import pytest
from pptx import Presentation as PptxReader

from pptx_tools import warnings as W
from pptx_tools.helpers import clean_control_chars
from pptx_tools.slide_builder import PowerpointPresentation


def build(slides, **kwargs):
    pres = PowerpointPresentation(slides, "16:9", **kwargs)
    return pres, PptxReader(pres.save())


def codes(pres):
    return [warning.code for warning in pres.warnings]


@pytest.mark.parametrize("bad", ["\x07", "\x0b", "\x01", "￾"])
class TestEachFieldThatUsedToCrash:

    def test_footer_text(self, bad):
        pres, deck = build([{"type": "content", "title": "T", "body": "- a"}],
                           footer_text=f"Conf{bad}idential")
        assert W.CONTROL_CHARS_REMOVED in codes(pres)
        assert len(deck.slides) == 1

    def test_section_title(self, bad):
        pres, deck = build([{"type": "section", "title": f"Part{bad} one"},
                            {"type": "content", "title": "T", "body": "- a"}])
        assert W.CONTROL_CHARS_REMOVED in codes(pres)
        assert len(deck.slides) == 2

    def test_chart_category_and_series_name(self, bad):
        pres, deck = build([{"type": "chart", "title": "C", "chart_type": "column",
                             "categories": [f"a{bad}", "b"],
                             "series": [{"name": f"s{bad}", "values": [1, 2]}]}])
        assert W.CONTROL_CHARS_REMOVED in codes(pres)
        assert W.CHART_FAILED not in codes(pres)
        chart = [s.chart for s in deck.slides[0].shapes if s.has_chart][0]
        assert chart.plots[0].series[0].name.startswith("s")


def test_one_deck_level_warning_counts_every_removal():
    pres, deck = build(
        [{"type": "title", "title": "T\x07", "subtitle": "x\x01"},
         {"type": "section", "title": "S\x07"}],
        footer_text="F\x07", author="A\x02",
    )
    warnings = [w for w in pres.warnings if w.code == W.CONTROL_CHARS_REMOVED]
    assert len(warnings) == 1
    assert warnings[0].slide is None
    assert warnings[0].severity == "info"
    assert warnings[0].message.startswith("5 control character(s)")
    assert deck.core_properties.author == "A"


def test_clean_text_is_untouched_and_reports_nothing():
    pres, _ = build([{"type": "content", "title": "Tab\there", "body": "- line\nnext"}],
                    footer_text="Plain")
    assert W.CONTROL_CHARS_REMOVED not in codes(pres)


def test_vertical_tab_and_form_feed_become_line_breaks():
    assert clean_control_chars("a\x0bb\x0cc\x07d") == ("a\nb\ncd", 3)
