"""The admin UI and Word content controls.

The analyser lists the controls and the tags that bind them, proposes an
argument for each tag (a check-box group becomes a choice with its options, a
drop-down a choice of its items), reconciles tags with the declared arguments,
and the editor's new "Options" column carries ``enum`` through a save — which
it used to drop.
"""
import io
import sys
from pathlib import Path

project_root = Path(__file__).parent.parent
sys.path.insert(0, str(project_root))

from docx import Document  # noqa: E402
from fasthtml.common import to_xml  # noqa: E402
from starlette.datastructures import FormData  # noqa: E402

from admin.analysis import analyze_docx, propose_args_from_controls, reconcile  # noqa: E402
from admin.forms import parse_args_from_form, parse_enum  # noqa: E402
from admin.preview import render_docx_preview, sample_values, values_from_form  # noqa: E402
from admin.views.templates import _arg_rows, analysis_report, arg_row  # noqa: E402

from . import docx_content_control_helpers as cc  # noqa: E402


def _template():
    doc = Document()
    p = doc.add_paragraph()
    cc.add_inline(p, cc.checkbox("size=small", alias="Small", sdt_id=1))
    cc.add_inline(p, cc.checkbox("size=medium", sdt_id=2))
    cc.add_inline(p, cc.checkbox("consent", alias="Customer consent", sdt_id=3))
    cc.add_inline(doc.add_paragraph(), cc.dropdown("plan", sdt_id=4))
    cc.add_inline(doc.add_paragraph(), cc.text("full_name", alias="Full name", sdt_id=5))
    cc.add_inline(doc.add_paragraph(), cc.rich_text("notes", sdt_id=6))
    cc.add_inline(doc.add_paragraph(), cc.text(None, sdt_id=7))
    doc.add_paragraph("{{date}}")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


ARGS = [
    {"name": "size", "type": "string", "enum": ["small", "medium"], "required": False, "description": ""},
    {"name": "consent", "type": "bool", "required": False, "description": ""},
    {"name": "plan", "type": "string", "enum": ["basic", "pro"], "required": False, "description": ""},
    {"name": "full_name", "type": "string", "required": False, "description": ""},
    {"name": "date", "type": "string", "required": False, "description": ""},
]


# --- analysis ------------------------------------------------------------------

def test_analysis_lists_every_control():
    found = analyze_docx(_template()).content_controls
    assert [(f["tag"], f["kind"], f["supported"]) for f in found] == [
        ("size=small", "checkbox", True), ("size=medium", "checkbox", True),
        ("consent", "checkbox", True), ("plan", "dropdown", True),
        ("full_name", "text", True), ("notes", "richtext", False), ("", "text", True)]


def test_proposals_follow_what_each_control_is():
    proposals = propose_args_from_controls(analyze_docx(_template()))
    assert list(proposals) == ["size", "consent", "plan", "full_name"]
    assert proposals["size"]["enum"] == ["small", "medium"]
    assert proposals["size"]["description"] == ""      # "Small" names an option, not the group
    assert proposals["consent"] == {"name": "consent", "type": "bool", "required": False,
                                    "description": "Customer consent"}
    assert proposals["plan"]["enum"] == ["basic", "pro"]
    assert proposals["full_name"]["description"] == "Full name"
    assert all(p["required"] is False for p in proposals.values())


def test_reconcile_counts_tags_as_uses_and_reports_missing_ones():
    analysis = analyze_docx(_template())
    rec = reconcile(analysis, ARGS)
    assert rec.missing_args == [] and rec.orphan_args == []
    rec = reconcile(analysis, [a for a in ARGS if a["name"] != "plan"])
    assert rec.missing_args == ["plan"]


def test_reconcile_explains_tags_that_will_not_fill():
    args = [dict(a) for a in ARGS]
    args[0]["enum"] = ["small", "large"]            # no "medium" any more
    args[2]["enum"] = ["basic", "pro", "enterprise"]   # the drop-down has no enterprise item
    issues = reconcile(analyze_docx(_template()), args).control_issues
    assert any("size=medium" in i and "never be ticked" in i for i in issues)
    assert any("plan" in i and "enterprise" in i for i in issues)
    assert any("notes" in i and "not filled yet" in i for i in issues)


def test_the_editor_prefills_rows_from_the_controls():
    rows = _arg_rows({"args": []}, analyze_docx(_template()))
    html = "".join(to_xml(r) for r in rows)
    for name in ("date", "size", "consent", "plan", "full_name"):
        assert f'value="{name}"' in html
    assert 'value="small, medium"' in html


def test_the_report_shows_the_controls_table():
    html = to_xml(analysis_report(analyze_docx(_template()), {"args": ARGS}))
    assert "Content controls" in html and "drop-down list" in html
    assert "filled from plan" in html and "no tag" in html


# --- the Options column ----------------------------------------------------------

def test_options_column_round_trips_enum_through_a_save():
    html = to_xml(arg_row(ARGS[0]))
    assert 'name="arg_enum"' in html and 'value="small, medium"' in html
    form = FormData([("arg_name", "size"), ("arg_type", "string"), ("arg_required", "false"),
                     ("arg_default", ""), ("arg_desc", "Size"), ("arg_enum", "small, medium")])
    assert parse_args_from_form(form) == [{"name": "size", "type": "string", "required": False,
                                           "description": "Size", "enum": ["small", "medium"]}]


def test_a_form_without_the_options_column_still_parses():
    form = FormData([("arg_name", "x"), ("arg_type", "string"), ("arg_required", "true"),
                     ("arg_default", ""), ("arg_desc", "")])
    assert parse_args_from_form(form) == [{"name": "x", "type": "string", "required": True,
                                           "description": ""}]


def test_numeric_options_stay_numbers_and_list_defaults_are_lists():
    assert parse_enum("int", "1, 2, 2, x") == [1, 2, "x"]
    form = FormData([("arg_name", "channels"), ("arg_type", "list"), ("arg_required", "false"),
                     ("arg_default", "email"), ("arg_desc", ""), ("arg_enum", "email, post")])
    arg = parse_args_from_form(form)[0]
    assert arg["enum"] == ["email", "post"] and arg["default"] == ["email"]


# --- preview -----------------------------------------------------------------------

def test_samples_pick_a_valid_option():
    args = ARGS + [{"name": "channels", "type": "list", "enum": ["email", "post"]}]
    s = sample_values(args)
    assert s["size"] == "small" and s["plan"] == "basic" and s["channels"] == ["email"]


def test_a_multi_select_with_nothing_chosen_is_an_empty_list():
    args = [{"name": "channels", "type": "list", "enum": ["email", "post"]}]
    assert values_from_form(FormData([("preview_values", "1")]), args)["channels"] == []
    form = FormData([("preview_values", "1"), ("value_channels", "post")])
    assert values_from_form(form, args)["channels"] == ["post"]


def test_preview_fills_the_controls_like_a_real_build():
    values = {"size": "medium", "consent": True, "plan": "pro", "full_name": "Jan", "date": "today"}
    out = Document(io.BytesIO(render_docx_preview(_template(), {"name": "t", "args": ARGS}, values)))
    sdts = list(out.element.body.iter(f"{{{cc.W}}}sdt"))
    assert [cc.checked(b)[0] for b in sdts[:3]] == [False, True, True]
    assert cc.content_text(sdts[3]) == "Pro plan"
    assert cc.content_text(sdts[4]) == "Jan"
    assert "today" in [p.text for p in out.paragraphs]
