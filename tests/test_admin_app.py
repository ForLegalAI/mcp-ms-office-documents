"""Integration tests for the FastHTML admin UI and combined ASGI app.

The template store and template-resolution directories are redirected to a
temp path so tests never touch the repo's custom_templates/config.
"""
import io
import re
import sys
from pathlib import Path

project_root = Path(__file__).parent.parent
sys.path.insert(0, str(project_root))

import pytest
from docx import Document
from starlette.testclient import TestClient

import admin.store as store_mod
import template_utils as tu
import metrics
from config import Config


@pytest.fixture
def admin_client(tmp_path, monkeypatch):
    """A logged-in TestClient over the combined app, isolated to tmp dirs."""
    custom = tmp_path / "custom"
    cfg = tmp_path / "config"
    custom.mkdir()
    cfg.mkdir()

    # Redirect the storage layer to the temp directories.
    monkeypatch.setattr(store_mod, "_APP_CUSTOM_DIR", tmp_path / "noexist_app_custom")
    monkeypatch.setattr(store_mod, "_APP_CONFIG_DIR", tmp_path / "noexist_app_config")
    monkeypatch.setattr(store_mod, "_LOCAL_CUSTOM_DIR", custom)
    monkeypatch.setattr(store_mod, "_LOCAL_CONFIG_DIR", cfg)
    # Redirect template resolution so live registration finds the tmp asset.
    monkeypatch.setattr(tu, "APP_CUSTOM_DIR", tmp_path / "noexist_app_custom")
    monkeypatch.setattr(tu, "LOCAL_CUSTOM_DIR", custom)

    monkeypatch.setenv("ADMIN_ENABLED", "true")
    monkeypatch.setenv("ADMIN_PASSWORD", "pw")
    monkeypatch.delenv("API_KEY", raising=False)

    # Import here so the patched module globals are in effect.
    from fastmcp import FastMCP
    from admin.app import build_combined_app

    mcp = FastMCP("test-admin")
    app = build_combined_app(mcp, Config.from_env())
    metrics.reset()
    client = TestClient(app)
    client.__enter__()
    # Authenticate.
    client.post("/admin/login", data={"password": "pw"})
    yield client, mcp
    client.__exit__(None, None, None)
    metrics.reset()


def _docx_with_placeholders(*paragraphs) -> bytes:
    doc = Document()
    for p in paragraphs:
        doc.add_paragraph(p)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


async def _tool_names(mcp):
    return [t.name for t in await mcp.list_tools()]


def _csrf(client) -> str:
    """Extract the session CSRF token from a page that always renders a form."""
    html = client.get("/admin/new/docx").text
    tag = re.search(r'<input[^>]*name="csrf"[^>]*>', html)
    assert tag, "no CSRF input found on page"
    return re.search(r'value="([^"]*)"', tag.group(0)).group(1)


def _post(client, url, data=None, files=None, **kwargs):
    """POST to an admin route with the session CSRF token injected."""
    payload = dict(data or {})
    payload["csrf"] = _csrf(client)
    return client.post(url, data=payload, files=files, **kwargs)


def test_login_required(tmp_path, monkeypatch):
    monkeypatch.setattr(store_mod, "_LOCAL_CUSTOM_DIR", tmp_path / "c")
    monkeypatch.setattr(store_mod, "_LOCAL_CONFIG_DIR", tmp_path / "cfg")
    monkeypatch.setenv("ADMIN_ENABLED", "true")
    monkeypatch.setenv("ADMIN_PASSWORD", "pw")
    from fastmcp import FastMCP
    from admin.app import build_combined_app
    app = build_combined_app(FastMCP("t"), Config.from_env())
    with TestClient(app) as c:
        r = c.get("/admin/", follow_redirects=False)
        assert r.status_code == 303
        assert r.headers["location"].endswith("/admin/login")
        # Wrong password stays on the login page.
        r = c.post("/admin/login", data={"password": "wrong"}, follow_redirects=False)
        assert r.status_code == 200
        assert "Incorrect" in r.text


def test_index_lists_sections(admin_client):
    client, _ = admin_client
    r = client.get("/admin/")
    assert r.status_code == 200
    assert "Word templates" in r.text
    assert "Email templates" in r.text


def test_draft_detects_placeholders(admin_client):
    client, _ = admin_client
    data = _docx_with_placeholders("Dear {{recipient}},", "{{body}}")
    r = _post(
        client, "/admin/docx/draft",
        data={"name": "letter_x"},
        files={"file": ("letter_x.docx", data, "application/octet-stream")},
    )
    assert r.status_code == 200
    assert "recipient" in r.text
    assert "body" in r.text


@pytest.mark.asyncio
async def test_create_makes_tool_live(admin_client):
    client, mcp = admin_client
    data = _docx_with_placeholders("Dear {{recipient}},", "{{body}}")
    _post(
        client, "/admin/docx/draft",
        data={"name": "live_letter"},
        files={"file": ("live_letter.docx", data, "application/octet-stream")},
    )
    save = {
        "kind": "docx", "asset_filename": "live_letter.docx", "name": "live_letter",
        "title": "Letter", "description": "desc",
        "arg_name": ["recipient", "body"], "arg_type": ["string", "string"],
        "arg_required": ["true", "true"], "arg_default": ["", ""],
        "arg_desc": ["who", "the body"],
    }
    r = _post(client, "/admin/docx/save", data=save)
    assert r.status_code == 200
    assert "now live" in r.text
    assert "live_letter" in await _tool_names(mcp)
    # Persisted as a managed spec file.
    assert store_mod.FileTemplateStore.from_config().get_spec("docx", "live_letter")["name"] == "live_letter"


@pytest.mark.asyncio
async def test_edit_then_delete(admin_client):
    client, mcp = admin_client
    data = _docx_with_placeholders("{{body}}")
    _post(client, "/admin/docx/draft", data={"name": "tmp_tpl"},
          files={"file": ("tmp_tpl.docx", data, "application/octet-stream")})
    save = {
        "kind": "docx", "asset_filename": "tmp_tpl.docx", "name": "tmp_tpl",
        "title": "T", "description": "d",
        "arg_name": ["body"], "arg_type": ["string"], "arg_required": ["true"],
        "arg_default": [""], "arg_desc": [""],
    }
    _post(client, "/admin/docx/save", data=save)
    assert "tmp_tpl" in await _tool_names(mcp)

    # Edit page renders with the saved spec.
    r = client.get("/admin/docx/tmp_tpl/edit")
    assert r.status_code == 200 and "tmp_tpl" in r.text

    # Delete unregisters and removes the managed spec.
    r = _post(client, "/admin/docx/tmp_tpl/delete", follow_redirects=False)
    assert r.status_code == 303
    assert "tmp_tpl" not in await _tool_names(mcp)
    assert store_mod.FileTemplateStore.from_config().get_spec("docx", "tmp_tpl") is None


def test_preview_returns_docx(admin_client):
    client, _ = admin_client
    data = _docx_with_placeholders("Dear {{recipient}},", "{{body}}")
    _post(client, "/admin/docx/draft", data={"name": "prev_tpl"},
          files={"file": ("prev_tpl.docx", data, "application/octet-stream")})
    pv = {
        "kind": "docx", "asset_filename": "prev_tpl.docx", "name": "prev_tpl",
        "title": "P", "description": "d",
        "arg_name": ["recipient", "body"], "arg_type": ["string", "string"],
        "arg_required": ["true", "true"], "arg_default": ["", ""], "arg_desc": ["", ""],
    }
    r = _post(client, "/admin/docx/preview", data=pv)
    assert r.status_code == 200
    assert r.content[:2] == b"PK"  # a real .docx (zip)
    assert "attachment" in r.headers.get("content-disposition", "")


def test_invalid_name_rejected(admin_client):
    client, _ = admin_client
    data = _docx_with_placeholders("{{body}}")
    r = _post(client, "/admin/docx/draft", data={"name": "1 bad name"},
              files={"file": ("x.docx", data, "application/octet-stream")})
    assert r.status_code == 200
    assert "Invalid template name" in r.text


def test_ui_theme_and_controls_present(admin_client):
    """The self-contained theme and dynamic-row controls are wired in.

    That the theme reaches for nothing external is a property of every page,
    not just this one, and is tested in tests/test_admin_assets.py.
    """
    client, _ = admin_client
    # Index ships the inline theme + topbar.
    idx = client.get("/admin/").text
    assert "Template Admin" in idx
    assert "--brand" in idx  # inline CSS variables

    data = _docx_with_placeholders("Dear {{recipient}},", "{{#if ps}}", "{{note}}", "{{/if}}")
    r = _post(client, "/admin/docx/draft", data={"name": "ux_tpl"},
              files={"file": ("ux_tpl.docx", data, "application/octet-stream")})
    body = r.text
    assert "+ Add argument" in body                 # dynamic add-row control
    assert "window.__ARG_ROW_HTML__" in body        # client-side row template
    assert 'class="chip"' in body                    # detected placeholder chips
    assert "adminRemoveRow" in body                  # per-row remove handler


def test_status_page_renders(admin_client):
    client, _ = admin_client
    r = client.get("/admin/status")
    assert r.status_code == 200
    assert "Uptime" in r.text
    assert "Upload backend" in r.text
    assert "Template usage" in r.text
    # The errors-only filter is a valid view too.
    assert client.get("/admin/status?level=error").status_code == 200


@pytest.mark.asyncio
async def test_status_reflects_tool_calls(admin_client):
    client, mcp = admin_client
    data = _docx_with_placeholders("Hi {{name}}")
    _post(client, "/admin/docx/draft", data={"name": "metric_tpl"},
          files={"file": ("metric_tpl.docx", data, "application/octet-stream")})
    _post(client, "/admin/docx/save", data={
        "kind": "docx", "asset_filename": "metric_tpl.docx", "name": "metric_tpl",
        "title": "M", "description": "d", "arg_name": ["name"], "arg_type": ["string"],
        "arg_required": ["true"], "arg_default": [""], "arg_desc": [""],
    })
    await mcp.call_tool("metric_tpl", {"data": {"name": "World"}})
    # Recorded regardless of whether the LOCAL upload succeeds in this env.
    st = metrics.get_tool_stat("metric_tpl")
    assert st is not None and (st.calls + st.errors) >= 1
    r = client.get("/admin/status")
    assert "metric_tpl" in r.text


def test_reupload_rescans_new_placeholder(admin_client):
    client, _ = admin_client
    _post(client, "/admin/docx/draft", data={"name": "reup_tpl"},
          files={"file": ("reup_tpl.docx", _docx_with_placeholders("Hi {{name}}"),
                          "application/octet-stream")})
    _post(client, "/admin/docx/save", data={
        "kind": "docx", "asset_filename": "reup_tpl.docx", "name": "reup_tpl",
        "title": "R", "description": "d", "arg_name": ["name"], "arg_type": ["string"],
        "arg_required": ["true"], "arg_default": [""], "arg_desc": [""],
    })
    # Replace the document with one that has an extra placeholder.
    r = _post(client, "/admin/docx/reup_tpl/reupload",
              files={"file": ("reup_tpl.docx",
                              _docx_with_placeholders("Hi {{name}}", "Ref {{case_no}}"),
                              "application/octet-stream")})
    assert r.status_code == 200
    assert "Re-scanned" in r.text
    assert "case_no" in r.text  # newly detected placeholder surfaced


ANALYSIS_CARD = "What we found in the document"


def _saved_docx_template(client, name="reup_err_tpl"):
    """A saved, live docx template with one placeholder."""
    _post(client, f"/admin/docx/draft", data={"name": name},
          files={"file": (f"{name}.docx", _docx_with_placeholders("Hi {{name}}"),
                          "application/octet-stream")})
    _post(client, "/admin/docx/save", data={
        "kind": "docx", "asset_filename": f"{name}.docx", "name": name,
        "title": "R", "description": "d", "arg_name": ["name"], "arg_type": ["string"],
        "arg_required": ["true"], "arg_default": [""], "arg_desc": [""],
    })
    return name


@pytest.mark.parametrize("payload, expected", [
    (b"", "choose a file"),
    (b"not a docx at all", "could not open"),
])
def test_reupload_rejects_bad_file(admin_client, payload, expected):
    """A rejected re-upload re-renders the editor with the error."""
    client, _ = admin_client
    name = _saved_docx_template(client)
    r = _post(client, f"/admin/docx/{name}/reupload",
              files={"file": (f"{name}.docx", payload, "application/octet-stream")})
    assert r.status_code == 200
    assert expected in r.text.lower()
    assert "Re-scanned" not in r.text
    # The arguments are still editable, and the replace card is still offered.
    assert 'name="arg_name"' in r.text
    assert "Replace document" in r.text


def test_reupload_oversized_is_rejected(admin_client):
    client, _ = admin_client
    from admin.app import MAX_UPLOAD_BYTES
    name = _saved_docx_template(client, "reup_big_tpl")
    r = _post(client, f"/admin/docx/{name}/reupload",
              files={"file": (f"{name}.docx", b"x" * (MAX_UPLOAD_BYTES + 1),
                              "application/octet-stream")})
    assert r.status_code == 200
    assert "too large" in r.text.lower()
    assert "Re-scanned" not in r.text


def test_reupload_rejection_keeps_the_installed_file(admin_client):
    """A rejected re-upload must not overwrite what is on disk."""
    client, _ = admin_client
    name = _saved_docx_template(client, "reup_keep_tpl")
    before = (store_mod.FileTemplateStore.from_config()
              .read_asset("docx", f"{name}.docx"))
    _post(client, f"/admin/docx/{name}/reupload",
          files={"file": (f"{name}.docx", b"junk", "application/octet-stream")})
    after = (store_mod.FileTemplateStore.from_config()
             .read_asset("docx", f"{name}.docx"))
    assert after == before


def test_reupload_rejection_does_not_report_on_the_old_file(admin_client):
    """The rejected page must not describe the file still on disk.

    "What we found in the document" means the file just uploaded everywhere
    else it appears; on this path that file was refused, so showing the card
    would describe a different document under the same heading.
    """
    client, _ = admin_client
    name = _saved_docx_template(client, "reup_report_tpl")
    # The card is on the ordinary edit page...
    assert ANALYSIS_CARD in client.get(f"/admin/docx/{name}/edit").text
    # ...and on a successful re-scan...
    ok = _post(client, f"/admin/docx/{name}/reupload",
               files={"file": (f"{name}.docx", _docx_with_placeholders("Hi {{name}}"),
                               "application/octet-stream")})
    assert ANALYSIS_CARD in ok.text
    # ...but not when the upload was refused.
    bad = _post(client, f"/admin/docx/{name}/reupload",
                files={"file": (f"{name}.docx", b"junk", "application/octet-stream")})
    assert ANALYSIS_CARD not in bad.text


def test_login_form_carries_no_csrf_field(admin_client):
    """Login authenticates by password; there is no session to hold a token.

    Rendering an empty one would imply a protection this route does not have.
    """
    client, _ = admin_client
    html = client.get("/admin/login", follow_redirects=False).text
    if "password" not in html.lower():  # already authed -> redirected
        client.get("/admin/logout")
        html = client.get("/admin/login").text
    assert 'name="password"' in html
    assert 'name="csrf"' not in html


def test_every_kind_has_a_reachable_create_link(admin_client):
    """#157: a kind's create page must not disappear once it has a template.

    The empty state used to be the only place that linked to it. Word and
    Email were covered by the top bar; PowerPoint was not, so making one
    PowerPoint template removed the only route to the page that made it.
    """
    from admin.kinds import KINDS

    client, _ = admin_client
    empty = client.get("/admin/").text
    for kind in KINDS:
        assert f"/admin/new/{kind}" in empty, f"{kind}: no create link when empty"

    # Give every kind a template, so no section is in its empty state.
    _post(client, "/admin/docx/draft", data={"name": "reach_docx"},
          files={"file": ("reach_docx.docx", _docx_with_placeholders("{{a}}"),
                          "application/octet-stream")})
    _post(client, "/admin/email/draft", data={"name": "reach_email"},
          files={"file": ("reach_email.html", b"<p>{{a}}</p>", "text/html")})
    pptx = (project_root / "default_templates" / "default_pptx_template_16_9.pptx")
    _post(client, "/admin/pptx/draft", data={"name": "reach_deck"},
          files={"file": ("reach_deck.pptx", pptx.read_bytes(),
                          "application/octet-stream")})

    populated = client.get("/admin/").text
    for kind in KINDS:
        assert f"/admin/new/{kind}" in populated, (
            f"{kind}: create page unreachable once a template exists")

    # Belt and braces: the top bar lists every kind too, so the table is not
    # the only way in. Checked separately because either one alone would
    # satisfy the reachability assertion above.
    nav = re.search(r"<nav>(.*?)</nav>", populated, re.S)
    assert nav, "no nav rendered"
    for kind in KINDS:
        assert f"/admin/new/{kind}" in nav.group(1), f"{kind}: missing from the top bar"


class _FakeRequest:
    """Just enough request for the `before` gate: it only reads url.path."""

    def __init__(self, path):
        self.url = type("U", (), {"path": path})()


def test_auth_gate_has_no_static_file_exemption():
    """#159: the gate used to let any path ending .css or .ico through.

    Tested against the gate itself rather than through the app: the admin UI
    serves no static files, so no route matches such a path and an HTTP
    request 404s before `before` is ever consulted. That is why the exemption
    leaked nothing — and also why only a direct test can show it is gone.
    """
    from admin import auth

    before = auth.make_before("/admin/login")
    signed_out = {}

    # The login endpoint is the one public path.
    assert before(_FakeRequest("/admin/login"), signed_out) is None
    assert before(_FakeRequest("/admin/login/"), signed_out) is None

    # Everything else redirects — including what used to be exempt.
    for path in ("/admin/", "/admin/status", "/admin/theme.css",
                 "/admin/favicon.ico", "/admin/nested/thing.css"):
        result = before(_FakeRequest(path), signed_out)
        assert getattr(result, "status_code", None) == 303, f"{path} was not gated"

    # An authenticated session still passes, and gets a CSRF token.
    signed_in = {auth.SESSION_KEY: True}
    assert before(_FakeRequest("/admin/theme.css"), signed_in) is None
    assert signed_in.get(auth.CSRF_KEY)


def _delete_fixture(client, name="del_tpl"):
    """A saved docx template; returns (name, asset filename)."""
    _post(client, "/admin/docx/draft", data={"name": name},
          files={"file": (f"{name}.docx", _docx_with_placeholders("Hi {{who}}"),
                          "application/octet-stream")})
    _post(client, "/admin/docx/save", data={
        "kind": "docx", "asset_filename": f"{name}.docx", "name": name,
        "title": "T", "description": "d", "arg_name": ["who"], "arg_type": ["string"],
        "arg_required": ["true"], "arg_default": [""], "arg_desc": [""],
    })
    return name, f"{name}.docx"


def test_delete_confirmation_offers_the_asset_choice(admin_client):
    """#158: deleting used to always keep the file, and never said so clearly."""
    client, _ = admin_client
    name, asset = _delete_fixture(client, "del_confirm")
    r = client.get(f"/admin/docx/{name}/delete")
    assert r.status_code == 200
    assert asset in r.text
    assert 'name="delete_asset"' in r.text
    # Nothing is destroyed by looking at the page.
    assert store_mod.FileTemplateStore.from_config().get_spec("docx", name) is not None


def test_delete_keeps_the_asset_by_default(admin_client):
    client, _ = admin_client
    name, asset = _delete_fixture(client, "del_keep")
    store = store_mod.FileTemplateStore.from_config()
    _post(client, f"/admin/docx/{name}/delete")
    assert store.get_spec("docx", name) is None
    assert store.asset_exists("docx", asset), "asset should be kept when not asked for"


def test_delete_can_remove_the_asset_too(admin_client):
    client, _ = admin_client
    name, asset = _delete_fixture(client, "del_asset")
    store = store_mod.FileTemplateStore.from_config()
    _post(client, f"/admin/docx/{name}/delete", data={"delete_asset": "1"})
    assert store.get_spec("docx", name) is None
    assert not store.asset_exists("docx", asset), "asset should be gone when asked for"


def test_delete_never_removes_an_asset_another_template_uses(admin_client):
    """The option is withheld on screen, and ignored if posted anyway."""
    client, _ = admin_client
    name, asset = _delete_fixture(client, "del_shared")
    store = store_mod.FileTemplateStore.from_config()
    # A second managed template pointing at the same file.
    store.save_spec("docx", {"name": "del_sharer", "description": "d",
                             "docx_path": asset, "args": []})

    page = client.get(f"/admin/docx/{name}/delete").text
    assert "also used by" in page.lower()
    assert 'name="delete_asset"' not in page

    _post(client, f"/admin/docx/{name}/delete", data={"delete_asset": "1"})
    assert store.get_spec("docx", name) is None
    assert store.asset_exists("docx", asset), "shared asset must survive"
    assert store.get_spec("docx", "del_sharer") is not None


def test_delete_unknown_template_is_not_found(admin_client):
    client, _ = admin_client
    assert "Not found" in client.get("/admin/docx/no_such_tpl/delete").text


def test_post_without_csrf_is_rejected(admin_client):
    client, _ = admin_client
    # A save POST with no CSRF token is refused.
    r = client.post("/admin/docx/save", data={
        "kind": "docx", "asset_filename": "x.docx", "name": "nope",
        "arg_name": [], "arg_type": [], "arg_required": [], "arg_default": [], "arg_desc": [],
    })
    assert r.status_code == 403


def test_oversized_upload_rejected(admin_client):
    client, _ = admin_client
    from admin.app import MAX_UPLOAD_BYTES
    big = b"x" * (MAX_UPLOAD_BYTES + 1)
    r = _post(client, "/admin/docx/draft", data={"name": "big_tpl"},
              files={"file": ("big_tpl.docx", big, "application/octet-stream")})
    assert r.status_code == 200
    assert "too large" in r.text.lower()


def test_logout_get_does_not_clear_session(admin_client):
    client, _ = admin_client
    # A GET to /logout only shows a confirm page; the session stays authed.
    r = client.get("/admin/logout")
    assert r.status_code == 200
    assert "Sign out" in r.text
    assert client.get("/admin/", follow_redirects=False).status_code == 200

    # A CSRF-validated POST actually logs out.
    r = _post(client, "/admin/logout", follow_redirects=False)
    assert r.status_code == 303
    assert client.get("/admin/", follow_redirects=False).status_code == 303  # back to login


def test_mcp_endpoint_still_works(admin_client):
    """The MCP endpoint is reachable through the combined app."""
    client, _ = admin_client
    headers = {"Accept": "application/json, text/event-stream", "Content-Type": "application/json"}
    r = client.post("/mcp", json={
        "jsonrpc": "2.0", "id": 1, "method": "initialize",
        "params": {"protocolVersion": "2025-06-18", "capabilities": {},
                   "clientInfo": {"name": "t", "version": "1"}},
    }, headers=headers)
    assert r.status_code == 200
    assert "mcp-session-id" in r.headers
