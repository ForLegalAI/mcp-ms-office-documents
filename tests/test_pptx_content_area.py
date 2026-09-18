"""Slides on a layout with no body placeholder (#119).

A Blank layout is not an empty canvas — templates keep a logo, a header rule
and a footer band on it — and Title Only has no body placeholder either. Both
used to draw into a hardcoded band starting 1.5 inches from the top of the
slide, which on a template with a deep title band or a header rule landed on
the decoration. They now use the rectangle the template itself reserves for
content, read from a layout that has a body placeholder.
"""
import sys
from pathlib import Path

project_root = Path(__file__).parent.parent
sys.path.insert(0, str(project_root))

import pytest
from pptx import Presentation as PptxReader
from pptx.util import Inches

from pptx_tools.layouts import LayoutResolver
from pptx_tools.placeholder_style import read_content_rect
from pptx_tools.slide_builder import PowerpointPresentation
from pptx_tools.templates import TemplateSpec, validate_templates

BASE_16_9 = project_root / "default_templates" / "default_pptx_template_16_9.pptx"


def content_rect_of(template=BASE_16_9):
    """The body placeholder rectangle of the template's content layout."""
    return read_content_rect(PptxReader(str(template)).slide_layouts[1])


def build(slides):
    pres = PowerpointPresentation(slides, "16:9")
    return pres, PptxReader(pres.save())


def strip_content_placeholders(prs):
    """Remove every body/object placeholder from the layouts and the master."""
    from pptx_tools.placeholder_style import content_placeholders
    for container in [prs.slide_master, *prs.slide_layouts]:
        for placeholder in content_placeholders(container):
            element = placeholder._element
            element.getparent().remove(element)


class TestDrawingStaysInsideTheContentArea:

    def test_an_untitled_quote_on_the_blank_layout(self):
        """The reproduction from #119: it used to start at a hardcoded 1.5in."""
        _, doc = build([{"type": "quote", "text": "Citace bez nadpisu.", "attribution": "§ 1"}])
        slide = doc.slides[0]
        rect = content_rect_of()

        assert slide.slide_layout.name == "Prázdný"
        box = slide.shapes[0]
        assert (box.left, box.top, box.width) == (rect.left, rect.top, rect.width)
        assert box.top > Inches(1.5)     # below the old hardcoded band

    @pytest.mark.parametrize("slide_data", [
        {"type": "kpi", "title": "K", "items": [{"value": "12", "label": "a"},
                                                {"value": "7", "label": "b"}]},
        {"type": "timeline", "title": "T", "steps": [{"label": "one"}, {"label": "two"}]},
    ])
    def test_slide_types_on_a_title_only_layout(self, slide_data):
        """Title Only has no body placeholder either, and drew into the title band."""
        _, doc = build([slide_data])
        rect = content_rect_of()
        drawn = [shape for shape in doc.slides[0].shapes if not shape.is_placeholder]

        assert drawn
        for shape in drawn:
            assert shape.top >= rect.top
            assert shape.top + shape.height <= rect.top + rect.height + Inches(0.01)

    def test_a_layout_with_a_body_placeholder_still_uses_it(self):
        """Unchanged: the slide's own placeholder wins over the template-wide area."""
        _, doc = build([{"type": "table", "title": "T", "rows": [["A"], ["1"]]}])
        slide = doc.slides[0]
        body = [ph for ph in slide.slide_layout.placeholders
                if ph.placeholder_format.idx == 1][0]
        table = [shape for shape in slide.shapes if shape.has_table][0]
        assert table.left == body.left


class TestReadingTheArea:

    def test_the_resolver_reads_the_content_layout(self):
        resolver = LayoutResolver(PptxReader(str(BASE_16_9)))
        area = resolver.content_area()
        assert area.as_tuple() == content_rect_of().as_tuple()
        assert area.source == "Nadpis a obsah"

    def test_it_is_cached(self):
        resolver = LayoutResolver(PptxReader(str(BASE_16_9)))
        assert resolver.content_area() is resolver.content_area()

    def test_the_largest_body_placeholder_wins(self):
        """On Comparison the first placeholder is the heading strip."""
        comparison = PptxReader(str(BASE_16_9)).slide_layouts[4]
        rect = read_content_rect(comparison)
        heights = [ph.height for ph in comparison.placeholders
                   if ph.placeholder_format.idx in (1, 2, 3, 4)]
        assert rect.height == max(heights)

    def test_a_blank_layout_reserves_nothing(self):
        assert read_content_rect(PptxReader(str(BASE_16_9)).slide_layouts[6]) is None


class TestTemplateWithNoBodyPlaceholderAnywhere:

    @pytest.fixture
    def bare_template(self, tmp_path):
        prs = PptxReader(str(BASE_16_9))
        strip_content_placeholders(prs)
        path = tmp_path / "bare.pptx"
        prs.save(str(path))
        return path

    def test_the_hardcoded_band_is_the_last_resort(self, bare_template):
        pres = PowerpointPresentation(
            [{"type": "quote", "text": "q"}], "16:9",
            template_spec=TemplateSpec(name="bare", path=bare_template, aspect="16:9"),
        )
        assert LayoutResolver(PptxReader(str(bare_template))).content_area() is None
        box = PptxReader(pres.save()).slides[0].shapes[0]
        assert box.top == Inches(1.5)


class TestItIsReportedToCallers:

    def test_the_template_report_carries_the_area_as_percentages(self):
        """What an author of a `blank` slide needs to position against."""
        areas = {report["name"]: report.get("content_area") for report in validate_templates()}
        assert areas["16_9"] == {
            "layout": "Nadpis a obsah", "x": "6.9%", "y": "26.6%", "w": "86.2%", "h": "63.4%",
        }
