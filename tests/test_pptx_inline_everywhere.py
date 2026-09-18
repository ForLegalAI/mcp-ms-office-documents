"""Inline markdown in *any* text field, as the tool has always claimed.

The tool description says "INLINE MARKDOWN in any text field" and "LINKS:
[label](url) works in any text field"; docs/powerpoint-slides.md says "every
text field". Ten of them rendered the markers literally instead, because the
text was assigned with `paragraph.text` rather than written through the
inline renderer. Every field a caller can put text in now goes through
`inline_formatting.write_text()`.
"""
import base64
import sys
from pathlib import Path

project_root = Path(__file__).parent.parent
sys.path.insert(0, str(project_root))

import pytest
from pptx import Presentation as PptxReader

from pptx_tools.slide_builder import PowerpointPresentation

MARKUP = "**bold** and [link](https://example.com)"
PNG_DATA_URI = "data:image/png;base64," + base64.b64encode(base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8z8BQDwAEhQGAhKmMIQAAAABJRU5ErkJggg=="
)).decode()

# Every field that takes free text, with the slide that carries it.
FIELDS = {
    "title": {"type": "content", "title": MARKUP, "body": "- x"},
    "section title": {"type": "section", "title": MARKUP},
    "blank title": {"type": "blank", "title": MARKUP,
                    "elements": [{"kind": "text", "text": "x", "x": 1, "y": 3, "w": 2}]},
    "subtitle": {"type": "title", "title": "T", "subtitle": MARKUP},
    "content body": {"type": "content", "title": "T", "body": f"- {MARKUP}"},
    "two_column heading": {"type": "two_column", "title": "T",
                           "left": {"heading": MARKUP, "body": "- l"},
                           "right": {"body": "- r"}},
    "two_column body": {"type": "two_column", "title": "T",
                        "left": {"body": f"- {MARKUP}"}, "right": {"body": "- r"}},
    "table cell": {"type": "table", "title": "T", "rows": [["H"], [MARKUP]]},
    "quote text": {"type": "quote", "title": "T", "text": MARKUP},
    "quote attribution": {"type": "quote", "title": "T", "text": "q", "attribution": MARKUP},
    "kpi value": {"type": "kpi", "title": "T", "items": [{"value": MARKUP, "label": "l"}]},
    "kpi label": {"type": "kpi", "title": "T", "items": [{"value": "1", "label": MARKUP}]},
    "kpi delta": {"type": "kpi", "title": "T",
                  "items": [{"value": "1", "label": "l", "delta": MARKUP}]},
    "timeline label": {"type": "timeline", "title": "T",
                       "steps": [{"label": MARKUP}, {"label": "b"}]},
    "timeline detail": {"type": "timeline", "title": "T",
                        "steps": [{"label": "a", "detail": MARKUP}, {"label": "b"}]},
    "agenda item": {"type": "agenda", "title": "T", "items": [MARKUP]},
    "image caption": {"type": "image", "title": "T", "source": PNG_DATA_URI, "caption": MARKUP},
    "closing subtitle": {"type": "closing", "title": "T", "subtitle": MARKUP},
    "closing contact": {"type": "closing", "title": "T", "contact": [MARKUP]},
    "blank text element": {"type": "blank", "title": "T", "elements": [
        {"kind": "text", "text": MARKUP, "x": 1, "y": 2, "w": 4}]},
    "blank shape text": {"type": "blank", "title": "T", "elements": [
        {"kind": "shape", "x": 1, "y": 2, "w": 4, "h": 1, "text": MARKUP}]},
}

# Chart text lives in the chart part, where PowerPoint does not follow a
# hyperlink. The emphasis still applies; the link renders as its label.
CHART_FIELDS = {
    "chart title": {"type": "chart", "title": "T", "chart_type": "column",
                    "categories": ["a"], "series": [{"name": "s", "values": [1]}],
                    "chart_title": MARKUP},
    "x axis title": {"type": "chart", "title": "T", "chart_type": "column",
                     "categories": ["a"], "series": [{"name": "s", "values": [1]}],
                     "x_title": MARKUP},
    "y axis title": {"type": "chart", "title": "T", "chart_type": "column",
                     "categories": ["a"], "series": [{"name": "s", "values": [1]}],
                     "y_title": MARKUP},
}


def frames_of(slide):
    """Every text frame a caller's string can land in, chart parts included."""
    frames = [shape.text_frame for shape in slide.shapes if shape.has_text_frame]
    for shape in slide.shapes:
        if getattr(shape, "has_table", False):
            frames += [cell.text_frame for row in shape.table.rows for cell in row.cells]
        if getattr(shape, "has_chart", False):
            chart = shape.chart
            if chart.has_title:
                frames.append(chart.chart_title.text_frame)
            for axis_name in ("category_axis", "value_axis"):
                try:
                    axis = getattr(chart, axis_name)
                except (ValueError, NotImplementedError):
                    continue
                if axis.has_title:
                    frames.append(axis.axis_title.text_frame)
    return frames


def runs_of(slide):
    return [run for frame in frames_of(slide) for paragraph in frame.paragraphs
            for run in paragraph.runs]


def render(slide_data):
    pres = PowerpointPresentation([slide_data], "16:9")
    return runs_of(PptxReader(pres.save()).slides[0])


@pytest.mark.parametrize("field", sorted(FIELDS))
class TestEveryTextField:

    def test_the_markers_are_not_shown_on_the_slide(self, field):
        assert not [r for r in render(FIELDS[field]) if "**" in r.text]

    def test_the_emphasis_is_applied(self, field):
        """The emphasised word is bold.

        Not "the only bold run": a KPI figure, a quote's attribution and a
        timeline label are bold by design, so there the markup adds nothing
        visible — what matters is that it was read rather than printed.
        """
        emphasised = [r for r in render(FIELDS[field]) if r.text == "bold"]

        assert emphasised and all(r.font.bold for r in emphasised)

    def test_the_link_is_a_hyperlink(self, field):
        addresses = [r.hyperlink.address for r in render(FIELDS[field]) if r.hyperlink.address]
        assert addresses == ["https://example.com"]


@pytest.mark.parametrize("field", sorted(CHART_FIELDS))
class TestChartText:

    def test_the_markers_are_not_shown_on_the_chart(self, field):
        assert not [r for r in render(CHART_FIELDS[field]) if "**" in r.text]

    def test_the_emphasis_is_applied(self, field):
        emphasised = [r for r in render(CHART_FIELDS[field]) if r.text == "bold"]

        assert emphasised and all(r.font.bold for r in emphasised)

    def test_the_link_renders_as_its_label(self, field):
        """PowerPoint does not follow a link inside chart text; the label stays."""
        texts = [r.text for r in render(CHART_FIELDS[field])]

        assert "link" in texts
        assert "https://example.com" not in " ".join(texts)


class TestWhatTheGrammarLeavesAlone:

    def test_prose_that_merely_contains_a_marker_is_untouched(self):
        """'5 * 3 * 2' is arithmetic, not emphasis — in a title as anywhere."""
        runs = render({"type": "content", "title": "5 * 3 * 2 = 30", "body": "- x"})
        titles = [r.text for r in runs if "5 *" in r.text]

        assert titles == ["5 * 3 * 2 = 30"]
        assert not any(r.font.bold for r in runs)

    def test_an_escaped_marker_loses_its_backslash(self):
        runs = render({"type": "kpi", "title": "T",
                       "items": [{"value": "1", "label": r"price \* qty"}]})

        assert [r.text for r in runs if "price" in r.text] == ["price * qty"]

    def test_plain_text_keeps_the_styling_the_builder_asked_for(self):
        """The plain path must still apply what the formatted path inherits."""
        from pptx.util import Pt

        pres = PowerpointPresentation(
            [{"type": "kpi", "title": "T", "items": [{"value": "€4.2M", "label": "ARR"}]}], "16:9")
        slide = PptxReader(pres.save()).slides[0]
        figure = [p for frame in frames_of(slide) for p in frame.paragraphs
                  if p.text == "€4.2M"][0]

        assert figure.font.size == Pt(40)
        assert figure.font.bold is True
