"""A ``<br>`` at the end of a line, followed by a newline, is one soft break.

Models write a signature block or an address as::

    _______________<br>
    ACME s.r.o.

The ``<br>`` says "break here"; the newline is only where the writer pressed
Enter. Before the fix the block parser saw two lines, made two paragraphs —
with the style's paragraph spacing between them — and the inline layer dropped
the ``<br>`` as a trailing break, all without a warning. A line ending in
``<br>`` now continues the paragraph exactly as a line ending in two spaces
does, on every entry point. A blank line after the ``<br>`` is still a new
paragraph.
"""
import sys
from pathlib import Path

import pytest
from docx import Document

project_root = Path(__file__).parent.parent
sys.path.insert(0, str(project_root))

from docx_tools.markdown_processor import process_markdown_content  # noqa: E402
from docx_tools.document_features import set_header_footer  # noqa: E402
from docx_tools.dynamic_docx_tools import (  # noqa: E402
    _replace_placeholders_in_document,
)
from docx_tools.patterns import br_to_newlines, ends_with_soft_break  # noqa: E402


def _render(markdown):
    doc = Document()
    process_markdown_content(doc, markdown)
    return doc


def _fill(value, template="{{body}}"):
    doc = Document()
    doc.add_paragraph(template)
    _replace_placeholders_in_document(doc, {"body": value})
    return doc


def _signature(doc):
    """(style, text) for every non-empty paragraph — what the reader sees."""
    return [(p.style.name, p.text) for p in doc.paragraphs if p.text.strip()]


def _breaks(paragraph):
    return paragraph._p.xml.count("<w:br")


SIGNATURE = "_______________\nACME s.r.o."


# ---------------------------------------------------------------------------
# paragraphs
# ---------------------------------------------------------------------------

@pytest.mark.parametrize("value", [
    "_______________<br>\nACME s.r.o.",
    "_______________<br/>\nACME s.r.o.",
    "_______________<BR />\nACME s.r.o.",
    "_______________<br>  \nACME s.r.o.",     # both markers at once
    "_______________<br>\r\nACME s.r.o.",      # CRLF
    "_______________<br>\\nACME s.r.o.",       # a literal backslash-n
])
def test_br_then_newline_is_one_paragraph_with_one_break(value):
    doc = _render(value)
    assert _signature(doc) == [("Normal", SIGNATURE)]
    assert _breaks(doc.paragraphs[0]) == 1


def test_br_then_blank_line_is_a_new_paragraph():
    doc = _render("_______________<br>\n\nACME s.r.o.")
    assert _signature(doc) == [("Normal", "_______________"), ("Normal", "ACME s.r.o.")]
    assert _breaks(doc.paragraphs[0]) == 0


def test_double_br_then_newline_keeps_both_breaks():
    # Only the one <br> the newline duplicates is absorbed: <br><br> + Enter is
    # an empty line inside the paragraph, as <br><br> mid-line is.
    para = _render("A<br><br>\nB").paragraphs[0]
    assert para.text == "A\n\nB"
    assert _breaks(para) == 2


def test_a_chain_of_br_lines_is_one_paragraph():
    para = _render("ACME s.r.o.<br>\nVodičkova 1<br>\n110 00 Praha 1").paragraphs[0]
    assert para.text == "ACME s.r.o.\nVodičkova 1\n110 00 Praha 1"
    assert _breaks(para) == 2


def test_br_line_and_trailing_spaces_line_mix():
    para = _render("A  \nB<br>\nC").paragraphs[0]
    assert para.text == "A\nB\nC"


def test_br_on_the_last_line_is_not_a_break():
    para = _render("Only line<br>").paragraphs[0]
    assert para.text == "Only line"
    assert _breaks(para) == 0


def test_br_in_a_quote_continues_the_quote():
    doc = _render("> Quoted<br>\n> still quoted")
    assert _signature(doc) == [("Quote", "Quoted\nstill quoted")]


def test_br_inside_an_alignment_block():
    doc = _render("<center>\n_______________<br>\nACME s.r.o.\n</center>")
    assert _signature(doc) == [("Normal", SIGNATURE)]


# ---------------------------------------------------------------------------
# a trailing <br> never swallows the block on the next line
# ---------------------------------------------------------------------------

@pytest.mark.parametrize("markdown, expected", [
    ("Intro<br>\n- item", [("Normal", "Intro"), ("List Bullet", "item")]),
    ("Intro<br>\n1. item", [("Normal", "Intro"), ("List Number", "item")]),
    ("Intro<br>\n# Title", [("Normal", "Intro"), ("Heading 1", "Title")]),
    ("Intro<br>\n> quote", [("Normal", "Intro"), ("Quote", "quote")]),
])
def test_br_before_a_block_keeps_the_block(markdown, expected):
    assert _signature(_render(markdown)) == expected


def test_br_before_a_table_keeps_the_table():
    doc = _render("Intro<br>\n| A | B |\n|---|---|\n| 1 | 2 |")
    assert _signature(doc) == [("Normal", "Intro")]
    assert len(doc.tables) == 1


def test_br_before_a_continuing_number_keeps_the_numbering():
    doc = _render("1. one\n2. two\n\nNote<br>\n3. three")
    assert _signature(doc)[-1] == ("List Number", "three")


def test_heading_is_still_one_line():
    # A heading is one line in markdown; a <br> ending it does not pull the next
    # line in, and it leaves no stray break behind.
    doc = _render("# Title<br>\nBody")
    assert _signature(doc) == [("Heading 1", "Title"), ("Normal", "Body")]
    assert _breaks(doc.paragraphs[0]) == 0


# ---------------------------------------------------------------------------
# list items
# ---------------------------------------------------------------------------

@pytest.mark.parametrize("marker", ["<br>", "  "])
def test_list_item_continues_across_a_soft_break(marker):
    doc = _render(f"- Seller:{marker}\nACME s.r.o.\n- Buyer")
    assert _signature(doc) == [
        ("List Bullet", "Seller:\nACME s.r.o."),
        ("List Bullet", "Buyer"),
    ]


def test_ordered_item_continues_and_keeps_counting():
    doc = _render("1. Seller:<br>\n   ACME s.r.o.\n2. Buyer")
    assert _signature(doc) == [
        ("List Number", "Seller:\nACME s.r.o."),
        ("List Number", "Buyer"),
    ]


def test_list_item_br_before_the_next_item_keeps_the_item():
    doc = _render("1. one<br>\n2. two")
    assert _signature(doc) == [("List Number", "one"), ("List Number", "two")]
    assert _breaks(doc.paragraphs[0]) == 0


def test_list_item_br_before_a_nested_item_keeps_the_nesting():
    doc = _render("- parent<br>\n  - child\n- sibling")
    assert _signature(doc) == [
        ("List Bullet", "parent"),
        ("List Bullet 2", "child"),
        ("List Bullet", "sibling"),
    ]


def test_list_item_br_then_blank_line_ends_the_list():
    doc = _render("- item<br>\n\nAfter")
    assert _signature(doc) == [("List Bullet", "item"), ("Normal", "After")]


def test_list_item_br_before_a_table_keeps_the_table():
    doc = _render("- item<br>\n| A | B |\n|---|---|\n| 1 | 2 |")
    assert _signature(doc) == [("List Bullet", "item")]
    assert len(doc.tables) == 1


# ---------------------------------------------------------------------------
# dynamic template placeholders and header/footer
# ---------------------------------------------------------------------------

def test_whole_paragraph_placeholder_matches_the_base_tool():
    value = "_______________<br>\nACME s.r.o."
    assert _signature(_fill(value)) == _signature(_render(value)) == [("Normal", SIGNATURE)]


def test_whole_paragraph_placeholder_blank_line_is_a_new_paragraph():
    doc = _fill("_______________<br>\n\nACME s.r.o.")
    assert _signature(doc) == [("Normal", "_______________"), ("Normal", "ACME s.r.o.")]


def test_mid_paragraph_placeholder_gets_one_break():
    doc = _fill("_______________<br>\nACME s.r.o.", template="Za kupujícího: {{body}}")
    para = doc.paragraphs[0]
    assert para.text == "Za kupujícího: " + SIGNATURE
    assert _breaks(para) == 1


def test_header_gets_one_break():
    doc = Document()
    set_header_footer(doc, "ACME s.r.o.<br>\nVodičkova 1", "header")
    para = doc.sections[0].header.paragraphs[0]
    assert para.text == "ACME s.r.o.\nVodičkova 1"
    assert _breaks(para) == 1


# ---------------------------------------------------------------------------
# the helpers
# ---------------------------------------------------------------------------

@pytest.mark.parametrize("line, expected", [
    ("text  ", True),
    ("text<br>", True),
    ("text<BR/> ", True),
    ("text<br />\t", True),
    ("text", False),
    ("text<br>more", False),
    ("text ", False),
])
def test_ends_with_soft_break(line, expected):
    assert ends_with_soft_break(line) is expected


@pytest.mark.parametrize("text, expected", [
    ("a<br>\nb", "a\nb"),
    ("a<br> \nb", "a\nb"),
    ("a<br><br>\nb", "a\n\nb"),
    ("a<br>b", "a\nb"),
    ("a\nb", "a\nb"),
])
def test_br_to_newlines(text, expected):
    assert br_to_newlines(text) == expected
