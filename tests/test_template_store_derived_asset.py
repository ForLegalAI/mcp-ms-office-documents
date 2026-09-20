"""`save_spec` must not adopt a leftover file it merely guessed the name of (#184).

The filename comes from `asset_filename`, else the spec's path key, else the
template's name. That last branch is a guess, and the guard beneath it read
"the asset must exist" — when what a caller in that branch means is "the asset
must be the one I just wrote".

So a file left behind by a template that was deleted with its source kept
would be adopted by the next template of the same name, silently: a working
template built on a document nobody uploaded in this session, with nothing
saying so.
"""
import io
import sys
from pathlib import Path

project_root = Path(__file__).parent.parent
sys.path.insert(0, str(project_root))

import pytest
from docx import Document

from admin.store import FileTemplateStore, TemplateStoreError


@pytest.fixture
def store(tmp_path):
    custom = tmp_path / "custom"
    config = tmp_path / "config"
    custom.mkdir()
    config.mkdir()
    return FileTemplateStore(custom_dir=custom, config_dir=config)


def _docx(text="hello"):
    doc = Document()
    doc.add_paragraph(text)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _spec(name="report", **extra):
    return {"name": name, "description": "d", "args": [], **extra}


def test_a_derived_name_landing_on_a_leftover_is_refused(store):
    """The scenario from the issue, end to end at the store."""
    (store.custom_dir / "report.docx").write_bytes(_docx("STALE LEFTOVER"))

    with pytest.raises(TemplateStoreError) as excinfo:
        store.save_spec("docx", _spec())

    assert "report.docx" in str(excinfo.value), "the message names the file"
    assert store.get_spec("docx", "report") is None, "and nothing is written"


def test_the_refusal_says_what_to_do_instead(store):
    (store.custom_dir / "report.docx").write_bytes(_docx())

    with pytest.raises(TemplateStoreError) as excinfo:
        store.save_spec("docx", _spec())

    message = str(excinfo.value)
    assert "Upload one" in message
    assert "name the file explicitly" in message


def test_naming_the_file_explicitly_still_reuses_it(store):
    """Reuse stays possible — it just has to be meant rather than guessed."""
    (store.custom_dir / "shared.docx").write_bytes(_docx("SHARED"))

    stored = store.save_spec("docx", _spec(), asset_filename="shared.docx")

    assert stored["docx_path"] == "shared.docx"
    assert store.get_spec("docx", "report")["docx_path"] == "shared.docx"


def test_a_spec_carrying_its_own_path_key_still_reuses_it(store):
    """Which is how every edit of an existing template saves."""
    (store.custom_dir / "report.docx").write_bytes(_docx())
    store.save_spec("docx", _spec(), asset_filename="report.docx")

    again = store.save_spec("docx", _spec(description="edited",
                                          docx_path="report.docx"))

    assert again["description"] == "edited"


def test_supplying_the_document_makes_a_derived_name_fine(store):
    """Deriving is only a guess when there is nothing to write."""
    (store.custom_dir / "report.docx").write_bytes(_docx("OLD"))

    store.save_spec("docx", _spec(), asset_bytes=_docx("NEW"))

    written = Document(str(store.custom_dir / "report.docx"))
    assert written.paragraphs[0].text == "NEW", "the upload wins, as before"


def test_a_derived_name_with_no_file_still_says_what_is_missing(store):
    """The pre-existing error for this branch must not be swallowed."""
    with pytest.raises(TemplateStoreError) as excinfo:
        store.save_spec("docx", _spec())

    assert "does not exist yet" in str(excinfo.value)


def test_the_two_refusals_are_distinguishable(store):
    """A missing file and an unexpected one are different problems."""
    missing = None
    try:
        store.save_spec("docx", _spec(name="absent"))
    except TemplateStoreError as e:
        missing = str(e)

    (store.custom_dir / "present.docx").write_bytes(_docx())
    unexpected = None
    try:
        store.save_spec("docx", _spec(name="present"))
    except TemplateStoreError as e:
        unexpected = str(e)

    assert missing and unexpected and missing != unexpected
