"""The Word tool's warnings channel (#114).

Word used to hand back a URL and nothing else, however much of the caller's
markdown had gone missing on the way: a block whose rendering raised was
logged and skipped, an image that would not load became a bracketed
placeholder, a style the template does not define fell back to Normal. The
caller — a model that could have fixed its markdown — was never told.

Each test here names one of those losses and asserts it comes back as data:
a stable code, a severity, and the source line to look at.
"""
import sys
from pathlib import Path

project_root = Path(__file__).parent.parent
sys.path.insert(0, str(project_root))

import pytest
from docx import Document

from docx_tools import _markdown_to_word_buffer
from docx_tools import warnings as W
from docx_tools.block_elements import add_image_to_doc, add_table_to_doc
from docx_tools.markdown_processor import process_markdown_content
from docx_tools.style_map import apply_style
from warning_channel import SEVERITIES, SEVERITY_ERROR, SEVERITY_WARNING

# A syntactically valid data URI that carries nothing decodable: the image
# fails without the test touching the network.
BROKEN_IMAGE = "data:image/png;base64,nope!"


def render(markdown):
    """Render *markdown* and return the warnings it produced."""
    warnings = W.channel()
    process_markdown_content(Document(), markdown, warnings=warnings)
    return warnings


def codes(warnings):
    return [warning.code for warning in warnings]


def only(warnings, code):
    matched = [w for w in warnings if w.code == code]
    assert matched, f"expected a {code} warning, got {codes(warnings)}"
    return matched[0]


class TestTheCodeTable:

    def test_every_code_has_a_severity(self):
        codes_defined = {
            value for name, value in vars(W).items()
            if name.isupper() and isinstance(value, str) and not name.startswith("SEVERITY")
        }
        assert codes_defined == set(W.WARNING_SEVERITY)

    def test_every_severity_is_one_of_the_three(self):
        assert set(W.WARNING_SEVERITY.values()) <= set(SEVERITIES)


class TestContentThatDidNotArrive:

    def test_a_block_that_cannot_be_rendered_is_reported(self, monkeypatch):
        """The renderer swallows the exception so one bad block does not cost
        the whole document — which is exactly why it has to say so."""
        import docx_tools.markdown_processor as mp

        def explode(text, paragraph, *args, **kwargs):
            if "boom" in text:
                raise RuntimeError("synthetic failure")
            return None

        monkeypatch.setattr(mp, "parse_inline_formatting", explode)

        warning = only(render("fine\n\nboom\n\nalso fine\n"), W.BLOCK_FAILED)
        assert warning.severity == SEVERITY_ERROR
        assert warning.location["line"] == 3
        assert "synthetic failure" in warning.message

    def test_an_image_that_will_not_load_is_reported(self):
        warning = only(render(f"![a diagram]({BROKEN_IMAGE})\n"), W.IMAGE_FAILED)

        assert warning.severity == SEVERITY_ERROR
        assert warning.location["line"] == 1
        assert BROKEN_IMAGE in warning.message

    def test_a_table_that_cannot_be_created_is_reported(self, monkeypatch):
        doc = Document()
        monkeypatch.setattr(doc, "add_table",
                            lambda *a, **k: (_ for _ in ()).throw(RuntimeError("no table")))
        warnings = W.channel()

        result = add_table_to_doc([["A", "B"]], doc, warnings=warnings, line=12)

        assert result is None
        warning = only(warnings, W.TABLE_FAILED)
        assert warning.location["line"] == 12
        assert "no table" in warning.message

    def test_a_cell_that_cannot_be_written_is_reported(self, monkeypatch):
        import docx_tools.block_elements as be
        monkeypatch.setattr(be, "parse_inline_formatting",
                            lambda *a, **k: (_ for _ in ()).throw(ValueError("bad cell")))
        warnings = W.channel()

        add_table_to_doc([["A"], ["1"]], Document(), warnings=warnings, line=4)

        warning = only(warnings, W.TABLE_CELL_FAILED)
        assert warning.severity == SEVERITY_ERROR
        assert warning.location["line"] == 4
        assert "row 1, column 1" in warning.message

    def test_each_failed_cell_is_reported_separately(self, monkeypatch):
        import docx_tools.block_elements as be
        monkeypatch.setattr(be, "parse_inline_formatting",
                            lambda *a, **k: (_ for _ in ()).throw(ValueError("bad cell")))
        warnings = W.channel()

        add_table_to_doc([["A", "B"]], Document(), warnings=warnings)

        assert len([w for w in warnings if w.code == W.TABLE_CELL_FAILED]) == 2


class TestInstructionsNotFollowed:

    def test_pipe_markup_that_is_not_a_table_is_reported(self):
        """Nothing is lost — it renders as prose — but the caller asked for a
        table and will find pipes. Excel calls the same input an error,
        because there the line has nowhere to fall through to."""
        warning = only(render("intro\n\n| Item | Qty |\n"), W.TABLE_NOT_RECOGNISED)

        assert warning.severity == SEVERITY_WARNING
        assert warning.location["line"] == 3
        assert "| Item | Qty |" in warning.message
        assert "separator row" in warning.message

    def test_the_text_is_still_in_the_document(self):
        doc = Document()
        process_markdown_content(doc, "| Item | Qty |\n")

        assert "| Item | Qty |" in doc.paragraphs[-1].text

    def test_a_well_formed_table_reports_nothing(self):
        assert list(render("| A | B |\n|---|---|\n| 1 | 2 |\n")) == []


    def test_a_style_the_template_lacks_is_reported(self):
        warning = only(render("<!-- style: Definitely Not A Style -->\nhello\n"),
                       W.STYLE_MISSING)

        assert warning.severity == SEVERITY_WARNING
        assert "Definitely Not A Style" in warning.message
        assert warning.location["line"] == 1

    def test_a_table_without_a_separator_row_is_reported(self):
        """Word makes the same call Excel does — first row is the header —
        and the caller did not ask for it either."""
        warning = only(render("| Item | Qty |\n| A    | 1   |\n"),
                       W.TABLE_SEPARATOR_MISSING)

        assert warning.severity == SEVERITY_WARNING
        assert warning.location["line"] == 1
        assert "used as the header" in warning.message

    def test_a_separator_row_anywhere_else_does_not_count(self):
        warnings = render("| A | 1 |\n| B | 2 |\n|---|---|\n")

        assert codes(warnings) == [W.TABLE_SEPARATOR_MISSING]

    def test_a_blank_row_is_not_a_separator_row(self):
        """`|  |  |` carries no dashes. Treating it as a separator swallowed
        the caller's blank row and counted the table as properly formed."""
        doc = Document()
        warnings = W.channel()
        process_markdown_content(
            doc, "| Item | Qty |\n|  |  |\n| A | 1 |\n", warnings=warnings)

        assert len(doc.tables[0].rows) == 3          # the blank row survives
        assert codes(warnings) == [W.TABLE_SEPARATOR_MISSING]

    def test_a_real_separator_row_still_is_one(self):
        doc = Document()
        warnings = W.channel()
        process_markdown_content(
            doc, "| A | B |\n|:--|--:|\n| 1 | 2 |\n", warnings=warnings)

        assert len(doc.tables[0].rows) == 2          # the separator is not a row
        assert list(warnings) == []

    def test_the_table_is_still_built(self):
        """A warning, not an error: every row the caller wrote is in it."""
        doc = Document()
        process_markdown_content(doc, "| Item | Qty |\n| A    | 1   |\n")

        rows = doc.tables[0].rows
        assert [c.text for c in rows[0].cells] == ["Item", "Qty"]
        assert [c.text for c in rows[1].cells] == ["A", "1"]

    def test_a_missing_style_is_reported_once_not_once_per_item(self):
        """A template without 'List Number' would otherwise warn per item."""
        from docx_tools.style_map import build_style_map

        warnings = W.channel()
        process_markdown_content(
            Document(), "1. one\n2. two\n3. three\n4. four\n",
            style_map=build_style_map({"list_number": "Nonexistent List"}),
            warnings=warnings,
        )

        assert codes(warnings).count(W.STYLE_MISSING) == 1

    def test_a_missing_fallback_style_is_reported_too(self):
        warnings = W.channel()
        apply_style(Document().add_paragraph(), "Not A Style",
                    fallback="Not A Fallback Either", warnings=warnings)

        assert codes(warnings) == [W.STYLE_MISSING, W.STYLE_FALLBACK_MISSING]

    def test_a_malformed_widths_directive_is_reported(self):
        markdown = "<!-- widths: wide narrow -->\n| A | B |\n|---|---|\n| 1 | 2 |\n"

        warning = only(render(markdown), W.WIDTHS_INVALID)
        assert warning.location["line"] == 2      # the table the directive modifies
        assert "wide narrow" in warning.message


class TestTheEntryPoint:

    def test_clean_markdown_produces_a_buffer_and_no_warnings(self):
        buffer, warnings = _markdown_to_word_buffer("# Title\n\nSome prose.\n")

        assert buffer.tell() == 0
        assert warnings == []
        buffer.close()

    def test_the_buffer_still_holds_a_document_when_something_went_wrong(self):
        """A warning is not an error: the file is delivered either way."""
        buffer, warnings = _markdown_to_word_buffer(
            f"# Title\n\n![x]({BROKEN_IMAGE})\n")

        assert len(buffer.read()) > 0
        assert codes(warnings) == [W.IMAGE_FAILED]
        buffer.close()

    def test_warnings_are_records_not_sentences(self):
        _, warnings = _markdown_to_word_buffer(f"![x]({BROKEN_IMAGE})\n")

        assert warnings[0].as_dict() == {
            "code": W.IMAGE_FAILED,
            "severity": SEVERITY_ERROR,
            "message": warnings[0].message,
            "line": 1,
        }

    def test_rendering_without_a_channel_still_works(self):
        """Dynamic template tools render through the same code and pass none."""
        assert process_markdown_content(Document(), f"![x]({BROKEN_IMAGE})\n") == []


class TestToolBoundary:
    """The structure only matters if it survives to the caller."""

    @staticmethod
    async def call(monkeypatch, **arguments):
        import main
        from fastmcp import Client

        async def fake_upload(file_buffer, extension, file_name, user_context,
                              message, **kwargs):
            return "https://example.invalid/doc.docx"

        monkeypatch.setattr(main, "upload_and_format_response", fake_upload)
        async with Client(main.mcp) as client:
            return await client.call_tool("create_word_from_markdown",
                                          arguments, raise_on_error=False)

    async def test_warnings_reach_the_caller_as_objects(self, monkeypatch):
        result = await self.call(monkeypatch,
                                 markdown_content=f"![x]({BROKEN_IMAGE})\n")

        assert not result.is_error
        assert result.data["file"] == "https://example.invalid/doc.docx"
        warning = result.data["warnings"][0]
        assert warning["code"] == W.IMAGE_FAILED
        assert warning["severity"] == "error"
        assert warning["line"] == 1
        assert "could not be loaded" in warning["message"]

    async def test_a_clean_document_still_returns_a_bare_url(self, monkeypatch):
        result = await self.call(monkeypatch, markdown_content="# Title\n\nProse.\n")

        assert not result.is_error
        assert result.data == "https://example.invalid/doc.docx"
