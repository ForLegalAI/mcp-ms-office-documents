"""`image` slides on a template's own picture layout (#120).

`ROLE_IMAGE_TEXT` was defined, classified and given a positional fallback, but
nothing ever asked for it and no builder ever filled a `PICTURE` placeholder —
so a template's photo layouts were dead weight and every image was scaled into
a rectangle the builder invented. An `image` slide now prefers a picture
layout when the template has one, and puts the picture in its placeholder.
"""
import base64
import sys
from pathlib import Path

project_root = Path(__file__).parent.parent
sys.path.insert(0, str(project_root))

import pytest
from pptx import Presentation as PptxReader
from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER
from pptx.oxml.ns import qn

from pptx_tools.layouts import ROLE_IMAGE_TEXT, LayoutResolver, role_for_slide
from pptx_tools.schema import coerce_slides
from pptx_tools.slide_builder import PowerpointPresentation
from pptx_tools.templates import TemplateSpec

BASE_16_9 = project_root / "default_templates" / "default_pptx_template_16_9.pptx"
PICTURE_LAYOUT = "Obrázek s titulkem"

PNG_DATA_URI = "data:image/png;base64," + base64.b64encode(base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8z8BQDwAEhQGAhKmMIQAAAABJRU5ErkJggg=="
)).decode()


def build(slide, spec=None):
    slide = {"type": "image", "source": PNG_DATA_URI, **slide}
    pres = PowerpointPresentation([slide], "16:9", template_spec=spec)
    return pres, PptxReader(pres.save()).slides[0]


def pictures(slide):
    return [shape for shape in slide.shapes if getattr(shape, "image", None) is not None]


def text_of(slide):
    return "\n".join(shape.text_frame.text for shape in slide.shapes if shape.has_text_frame)


@pytest.fixture
def photo_only(tmp_path):
    """A template whose picture layout has a title and a picture, no text."""
    prs = PptxReader(str(BASE_16_9))
    layout = prs.slide_layouts[8]
    layout.name = "Foto"
    for placeholder in list(layout.placeholders):
        if placeholder.placeholder_format.type in (PP_PLACEHOLDER.BODY, PP_PLACEHOLDER.OBJECT):
            element = placeholder._element
            element.getparent().remove(element)
    path = tmp_path / "photo_only.pptx"
    prs.save(str(path))
    return TemplateSpec(name="photo_only", path=path, aspect="16:9")


@pytest.fixture
def no_picture_layout(tmp_path):
    """The shipped template with its picture placeholders removed."""
    prs = PptxReader(str(BASE_16_9))
    for layout in prs.slide_layouts:
        for placeholder in list(layout.placeholders):
            if placeholder.placeholder_format.type == PP_PLACEHOLDER.PICTURE:
                element = placeholder._element
                element.getparent().remove(element)
    path = tmp_path / "no_picture.pptx"
    prs.save(str(path))
    return TemplateSpec(name="no_picture", path=path, aspect="16:9")


class TestTheRoleIsReachable:

    def test_an_image_slide_asks_for_the_picture_layout(self):
        resolver = LayoutResolver(PptxReader(str(BASE_16_9)))
        slide = coerce_slides([{"type": "image", "source": PNG_DATA_URI}])[0]
        assert role_for_slide(slide, resolver) == ROLE_IMAGE_TEXT

    def test_without_one_it_stays_on_the_content_layout(self, no_picture_layout):
        resolver = LayoutResolver(PptxReader(str(no_picture_layout.path)))
        slide = coerce_slides([{"type": "image", "source": PNG_DATA_URI}])[0]
        assert not resolver.provides(ROLE_IMAGE_TEXT)
        assert role_for_slide(slide, resolver) == "content"

    def test_a_configured_layout_counts_as_provided(self):
        """A template may name a layout the signature rules do not classify."""
        resolver = LayoutResolver(PptxReader(str(BASE_16_9)),
                                  {ROLE_IMAGE_TEXT: "Nadpis a svislý text"})
        assert resolver.provides(ROLE_IMAGE_TEXT)

    def test_the_builder_lands_on_it(self):
        _, slide = build({"title": "Chart"})
        assert slide.slide_layout.name == PICTURE_LAYOUT


class TestThePlaceholderIsFilled:

    def test_the_picture_is_the_layout_placeholder(self):
        """Not a shape at a computed position: the template's own frame."""
        _, slide = build({"title": "Chart"})
        picture = pictures(slide)[0]
        assert picture.is_placeholder
        assert picture.placeholder_format.type == PP_PLACEHOLDER.PICTURE

        layout_placeholder = [ph for ph in slide.slide_layout.placeholders
                              if ph.placeholder_format.type == PP_PLACEHOLDER.PICTURE][0]
        assert (picture.left, picture.top, picture.width, picture.height) == (
            layout_placeholder.left, layout_placeholder.top,
            layout_placeholder.width, layout_placeholder.height)

    def test_body_and_caption_go_in_the_text_placeholder(self):
        _, slide = build({"title": "Chart", "body": "- Up and to the right",
                          "caption": "Fig 1"})
        body = [shape for shape in slide.shapes
                if shape.is_placeholder
                and shape.placeholder_format.type in (PP_PLACEHOLDER.BODY, PP_PLACEHOLDER.OBJECT)][0]
        assert "Up and to the right" in body.text_frame.text
        assert "Fig 1" in body.text_frame.text
        # Nothing drawn beside or under the picture.
        assert not [s for s in slide.shapes if s.shape_type == MSO_SHAPE_TYPE.TEXT_BOX]

    def test_a_caption_alone_fills_the_text_placeholder(self):
        _, slide = build({"title": "Chart", "caption": "Fig 1"})
        assert "Fig 1" in text_of(slide)
        assert not [s for s in slide.shapes if s.shape_type == MSO_SHAPE_TYPE.TEXT_BOX]

    def test_an_unused_text_placeholder_is_removed(self):
        """An empty one shows 'Click to add text' the moment anyone edits."""
        _, slide = build({"title": "Chart"})
        kinds = [shape.placeholder_format.type for shape in slide.placeholders]
        assert PP_PLACEHOLDER.BODY not in kinds and PP_PLACEHOLDER.OBJECT not in kinds

    def test_overfull_text_is_shrunk_and_reported(self):
        """The layout's text area is small; overflow there needs saying."""
        pres, slide = build({"title": "Chart",
                             "body": "\n".join(f"- {'word ' * 12}" for _ in range(20))})
        body = [shape for shape in slide.shapes
                if shape.is_placeholder
                and shape.placeholder_format.type in (PP_PLACEHOLDER.BODY, PP_PLACEHOLDER.OBJECT)][0]
        bodyPr = body.text_frame._txBody.find(qn('a:bodyPr'))
        assert bodyPr.find(qn('a:normAutofit')) is not None
        assert any("space available" in w for w in pres.warning_messages)

    def test_a_failed_image_leaves_no_empty_placeholder(self):
        pres, slide = build({"title": "Chart", "source": "data:image/png;base64,not-base64!"})
        assert not pictures(slide)
        assert PP_PLACEHOLDER.PICTURE not in [ph.placeholder_format.type
                                              for ph in slide.placeholders]
        assert "Image could not be loaded" in text_of(slide)
        assert any("image could not be loaded" in w for w in pres.warning_messages)


class TestFallingBack:

    def test_no_picture_layout_keeps_the_computed_rectangle(self, no_picture_layout):
        _, slide = build({"title": "Chart"}, spec=no_picture_layout)
        picture = pictures(slide)[0]
        assert slide.slide_layout.name == "Nadpis a obsah"
        assert not picture.is_placeholder

    def test_a_named_picture_layout_without_room_for_the_text(self, photo_only):
        """Text with nowhere to go falls back rather than being dropped."""
        _, slide = build({"title": "Chart", "body": "- Beside it",
                          "layout": "Foto"}, spec=photo_only)
        assert "Beside it" in text_of(slide)
        picture = pictures(slide)[0]
        assert not picture.is_placeholder
        assert PP_PLACEHOLDER.PICTURE not in [ph.placeholder_format.type
                                              for ph in slide.placeholders]
