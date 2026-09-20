"""Previewing with your own values rather than generated samples (#168).

`sample_values()` substitutes `[recipient_name]` for every string. That proves
the placeholders are wired up. It cannot prove the template *works*, because
the one thing these templates are for — Markdown inside a placeholder value —
is never exercised: `[body]` is a plain string, so the preview renders a plain
string and the admin ships a template having never seen a list, a heading or a
bold run come out of it.

Conditionals are the other half, and are subtler than #168 puts it. A sample
does not force every flag on: it uses the declared default and only falls back
to True when there is none. So an undeclared or required flag always previews
*on* and can never be seen off, while an optional one always previews *off*
and can never be seen on — `build_spec()` writes `default: false` for a blank
default. Either way it is a guess, which is the argument for typing the value.
"""
import io
import html as html_mod
import json
import re
import sys
from pathlib import Path

project_root = Path(__file__).parent.parent
sys.path.insert(0, str(project_root))

import pytest
from docx import Document
from starlette.testclient import TestClient

import admin.store as store_mod
import metrics
import template_utils as tu
from admin.forms import carried_spec
from admin.preview import has_submitted_values, sample_values, values_from_form
from config import Config


@pytest.fixture
def admin_client(tmp_path, monkeypatch):
    custom = tmp_path / "custom"
    cfg = tmp_path / "config"
    custom.mkdir()
    cfg.mkdir()
    monkeypatch.setattr(store_mod, "_APP_CUSTOM_DIR", tmp_path / "nx1")
    monkeypatch.setattr(store_mod, "_APP_CONFIG_DIR", tmp_path / "nx2")
    monkeypatch.setattr(store_mod, "_LOCAL_CUSTOM_DIR", custom)
    monkeypatch.setattr(store_mod, "_LOCAL_CONFIG_DIR", cfg)
    monkeypatch.setattr(tu, "APP_CUSTOM_DIR", tmp_path / "nx1")
    monkeypatch.setattr(tu, "LOCAL_CUSTOM_DIR", custom)
    monkeypatch.setenv("ADMIN_ENABLED", "true")
    monkeypatch.setenv("ADMIN_PASSWORD", "pw")
    monkeypatch.delenv("API_KEY", raising=False)

    from fastmcp import FastMCP
    from admin.app import build_combined_app

    client = TestClient(build_combined_app(FastMCP("test-preview"), Config.from_env()))
    metrics.reset()
    client.__enter__()
    client.post("/admin/login", data={"password": "pw"})
    yield client, custom
    client.__exit__(None, None, None)
    metrics.reset()


def _template_bytes(*paragraphs):
    doc = Document()
    for text in paragraphs:
        doc.add_paragraph(text)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _csrf(client) -> str:
    html = client.get("/admin/new/docx").text
    tag = re.search(r'<input[^>]*name="csrf"[^>]*>', html)
    return re.search(r'value="([^"]*)"', tag.group(0)).group(1)


def _post(client, url, data=None, files=None):
    payload = dict(data or {})
    payload["csrf"] = _csrf(client)
    return client.post(url, data=payload, files=files)


#: A template exercising both halves: a body placeholder and a conditional.
BODY_AND_FLAG = ("{{body}}", "{{#if extras}}", "OPTIONAL BLOCK", "{{/if}}")

FIELDS = dict(
    kind="docx", asset_filename="rpt.docx", name="rpt", original_name="rpt",
    title="R", description="d",
    arg_name=["body", "extras"], arg_type=["string", "bool"],
    arg_required=["true", "false"], arg_default=["", ""],
    arg_desc=["The body", "Show extras"],
)


def _saved(client, *paragraphs):
    _post(client, "/admin/docx/draft", data={"name": "rpt"},
          files={"file": ("rpt.docx", _template_bytes(*paragraphs),
                          "application/octet-stream")})
    _post(client, "/admin/docx/save", data=FIELDS)


def _carried(html: str) -> str:
    """The spec the values form is carrying, as the browser would send it."""
    tag = re.search(r'<input[^>]*name="spec_json"[^>]*>', html)
    assert tag, "the values form must carry the spec it was built from"
    return html_mod.unescape(re.search(r"""value=(["'])(.*?)\1""", tag.group(0)).group(2))


def _paragraphs(response):
    doc = Document(io.BytesIO(response.content))
    return [(p.style.name, p.text, any(r.bold for r in p.runs))
            for p in doc.paragraphs if p.text.strip()]


# ---------------------------------------------------------------------------
# The two things a generated sample can never show
# ---------------------------------------------------------------------------


def test_markdown_in_a_value_renders_as_document_structure(admin_client):
    """The whole point of #168.

    `[body]` is a plain string, so the sample preview renders a plain string.
    Only a real value proves the block-level Markdown promised by
    docs/templates.md actually comes out of a body placeholder.
    """
    client, _custom = admin_client
    _saved(client, *BODY_AND_FLAG)
    form = _post(client, "/admin/docx/preview/values", data=FIELDS).text

    r = _post(client, "/admin/docx/preview", data={
        "spec_json": _carried(form),
        "value_body": "# A Heading\n\n- first\n- second\n\nSome **bold** text.",
        "value_extras": "on",
    })

    styles = [style for style, _text, _bold in _paragraphs(r)]
    assert "Heading 1" in styles, "a heading must come out as a heading"
    assert styles.count("List Bullet") == 2, "a list must come out as a list"
    assert any(bold for _s, _t, bold in _paragraphs(r)), "bold must come out bold"


def test_a_conditional_can_be_previewed_switched_off(admin_client):
    """Previewing a flag in the state the sample does not pick.

    `extras` is optional here, so its sample is False; the point is that both
    states are now reachable on demand rather than whichever one the declared
    default happens to give.
    """
    client, _custom = admin_client
    _saved(client, *BODY_AND_FLAG)
    form = _post(client, "/admin/docx/preview/values", data=FIELDS).text
    carried = _carried(form)

    on = _post(client, "/admin/docx/preview", data={
        "spec_json": carried, "value_body": "x", "value_extras": "on"})
    off = _post(client, "/admin/docx/preview", data={
        "spec_json": carried, "value_body": "x"})  # unticked sends nothing

    assert any("OPTIONAL BLOCK" in t for _s, t, _b in _paragraphs(on))
    assert not any("OPTIONAL BLOCK" in t for _s, t, _b in _paragraphs(off)), \
        "an unticked conditional must render the document without the block"


# ---------------------------------------------------------------------------
# One click still works for anyone who does not care
# ---------------------------------------------------------------------------


def test_preview_without_values_still_uses_samples(admin_client):
    client, _custom = admin_client
    _saved(client, *BODY_AND_FLAG)

    r = _post(client, "/admin/docx/preview", data=FIELDS)

    texts = [t for _s, t, _b in _paragraphs(r)]
    assert any("[body]" in t for t in texts), "the sample placeholder"


@pytest.mark.parametrize("declared, shown", [
    # A conditional found in the document but never declared as an argument.
    (None, True),
    # Declared and required: no default is written, so it falls back to True.
    ("required", True),
    # Declared optional: build_spec writes `default: false` for a blank
    # default, and the sample respects it — so the block previews hidden.
    ("optional", False),
])
def test_what_a_sampled_conditional_actually_does(admin_client, declared, shown):
    """Pinning the sample behaviour, which is subtler than "flags are True".

    #168 describes samples as forcing every boolean on, so a block can never
    be previewed off. That holds for an undeclared or required flag; for an
    optional one the opposite was true and it could never be previewed *on*.
    Both are guesses, which is the argument for typing the value yourself.
    """
    client, _custom = admin_client
    _post(client, "/admin/docx/draft", data={"name": "rpt"},
          files={"file": ("rpt.docx", _template_bytes(*BODY_AND_FLAG),
                          "application/octet-stream")})
    fields = dict(FIELDS, arg_name=["body"], arg_type=["string"],
                  arg_required=["true"], arg_default=[""], arg_desc=[""])
    if declared:
        fields = dict(FIELDS, arg_required=["true",
                      "true" if declared == "required" else "false"])
    _post(client, "/admin/docx/save", data=fields)

    r = _post(client, "/admin/docx/preview", data=fields)
    texts = [t for _s, t, _b in _paragraphs(r)]
    assert any("OPTIONAL BLOCK" in t for t in texts) is shown


def test_has_submitted_values_distinguishes_the_two_posts():
    assert has_submitted_values({"value_body": "x"}) is True
    assert has_submitted_values({"name": "rpt", "arg_name": "body"}) is False


# ---------------------------------------------------------------------------
# The form itself
# ---------------------------------------------------------------------------


def test_the_form_is_prefilled_with_the_samples(admin_client):
    """So it stays one click for anyone who does not want to type."""
    client, _custom = admin_client
    _saved(client, *BODY_AND_FLAG)

    html = _post(client, "/admin/docx/preview/values", data=FIELDS).text
    assert "[body]" in html


def test_a_string_gets_a_textarea_so_markdown_is_possible(admin_client):
    """A single-line input cannot hold the list this feature exists to test."""
    client, _custom = admin_client
    _saved(client, *BODY_AND_FLAG)

    html = _post(client, "/admin/docx/preview/values", data=FIELDS).text
    body_control = re.search(r"<textarea[^>]*name=\"value_body\"", html)
    assert body_control, "a string value needs a textarea, not an input"


def test_a_boolean_gets_a_checkbox(admin_client):
    client, _custom = admin_client
    _saved(client, *BODY_AND_FLAG)

    html = _post(client, "/admin/docx/preview/values", data=FIELDS).text
    tag = re.search(r'<input[^>]*name="value_extras"[^>]*>', html)
    assert tag and 'type="checkbox"' in tag.group(0)


def test_an_undeclared_conditional_still_gets_a_control(admin_client):
    """Otherwise a conditional found in the document could not be turned off."""
    client, _custom = admin_client
    _post(client, "/admin/docx/draft", data={"name": "rpt"},
          files={"file": ("rpt.docx", _template_bytes(*BODY_AND_FLAG),
                          "application/octet-stream")})
    only_body = dict(FIELDS, arg_name=["body"], arg_type=["string"],
                     arg_required=["true"], arg_default=[""], arg_desc=[""])
    _post(client, "/admin/docx/save", data=only_body)

    html = _post(client, "/admin/docx/preview/values", data=only_body).text
    assert 'name="value_extras"' in html, \
        "a conditional detected in the document needs a control even undeclared"


def test_the_values_form_is_offered_from_the_edit_page(admin_client):
    client, _custom = admin_client
    _saved(client, *BODY_AND_FLAG)

    html = client.get("/admin/docx/rpt/edit").text
    assert "/admin/docx/preview/values" in html


def test_powerpoint_is_not_offered_a_values_form(admin_client):
    """A presentation template declares no arguments; there is nothing to fill."""
    client, _custom = admin_client
    source = project_root / "default_templates" / "default_pptx_template_16_9.pptx"
    _post(client, "/admin/pptx/draft", data={"name": "deck"},
          files={"file": ("deck.pptx", source.read_bytes(),
                          "application/octet-stream")})
    _post(client, "/admin/pptx/save", data={
        "kind": "pptx", "asset_filename": "deck.pptx", "name": "deck",
        "original_name": "deck", "description": "d"})

    html = client.get("/admin/pptx/deck/edit").text
    assert "/admin/pptx/preview/values" not in html


def test_the_values_form_needs_the_csrf_token(admin_client):
    client, _custom = admin_client
    _saved(client, *BODY_AND_FLAG)

    r = client.post("/admin/docx/preview/values", data=FIELDS)
    assert "Preview rpt" not in r.text


# ---------------------------------------------------------------------------
# Carrying the spec between the two posts
# ---------------------------------------------------------------------------


def test_unsaved_edits_reach_the_preview(admin_client):
    """The form is reached from the editor, so it must carry what is on screen."""
    client, _custom = admin_client
    _saved(client, *BODY_AND_FLAG)

    edited = dict(FIELDS, description="not saved anywhere")
    html = _post(client, "/admin/docx/preview/values", data=edited).text

    assert json.loads(_carried(html))["description"] == "not saved anywhere"


def test_an_apostrophe_survives_the_round_trip(admin_client):
    """The carried JSON is an HTML attribute; "the client's letter" is ordinary."""
    client, _custom = admin_client
    _saved(client, *BODY_AND_FLAG)

    edited = dict(FIELDS, description="the client's letter")
    html = _post(client, "/admin/docx/preview/values", data=edited).text

    assert json.loads(_carried(html))["description"] == "the client's letter"


@pytest.mark.parametrize("raw", ["", "not json", "[]", '"a string"', "{}",
                                 '{"description": "no name"}'])
def test_a_junk_carried_spec_is_ignored(raw):
    assert carried_spec({"spec_json": raw}) is None


def test_a_valid_carried_spec_is_used():
    spec = {"name": "rpt", "args": []}
    assert carried_spec({"spec_json": json.dumps(spec)}) == spec


# ---------------------------------------------------------------------------
# Value coercion
# ---------------------------------------------------------------------------


def test_values_are_coerced_to_their_declared_types():
    args = [{"name": "n", "type": "int"}, {"name": "f", "type": "float"},
            {"name": "s", "type": "string"}]
    values = values_from_form(
        {"value_n": "7", "value_f": "1.5", "value_s": "text"}, args)
    assert values == {"n": 7, "f": 1.5, "s": "text"}


def test_an_unparseable_number_falls_back_to_the_sample():
    """A half-filled form should still render: the point is to look at it."""
    args = [{"name": "n", "type": "int"}]
    assert values_from_form({"value_n": "not a number"}, args) == \
        sample_values(args)


def test_a_missing_value_falls_back_to_the_sample():
    args = [{"name": "a", "type": "string"}, {"name": "b", "type": "string"}]
    values = values_from_form({"value_a": "given"}, args)
    assert values["a"] == "given"
    assert values["b"] == "[b]"


def test_an_empty_string_is_honoured_rather_than_replaced(admin_client):
    """Clearing a field means "render it empty", not "put the sample back"."""
    args = [{"name": "a", "type": "string"}]
    assert values_from_form({"value_a": ""}, args) == {"a": ""}


# ---------------------------------------------------------------------------
# The carried spec is client-supplied, so its asset path is still gated
# ---------------------------------------------------------------------------


@pytest.mark.parametrize("attempt", [
    "../secret.docx",
    "../../secret.docx",
    "/etc/passwd",
    "../config/docx_templates.yaml",
    "sub/../../secret.docx",
])
def test_a_carried_spec_cannot_read_outside_the_uploads_directory(admin_client,
                                                                  tmp_path,
                                                                  attempt):
    """The spec rides through the form, so its path key is attacker-shaped.

    `asset_exists()` resolves through `asset_path()`, which runs
    `validate_asset_filename()` and refuses anything but a bare filename — so
    a crafted path never reaches `read_asset()`.
    """
    client, _custom = admin_client
    (tmp_path / "secret.docx").write_bytes(b"SENSITIVE")

    spec = {"name": "x", "description": "d", "docx_path": attempt, "args": []}
    r = _post(client, "/admin/docx/preview",
              data={"spec_json": json.dumps(spec), "value_a": "x"})

    assert r.status_code == 400
    assert b"SENSITIVE" not in r.content


def test_the_body_and_cell_distinction_becomes_visible(admin_client):
    """Answering my own question: is the third invisible thing now testable?

    docs/templates.md promises block-level Markdown in a body placeholder but
    inline-only inside a table cell. With a generated `[body]` sample neither
    treatment is exercised, so the difference is invisible. Typing the same
    Markdown into both shows it: the body becomes a real bullet list, the
    cell does not gain list paragraphs.
    """
    from docx import Document as Doc

    client, _custom = admin_client
    doc = Doc()
    doc.add_paragraph("{{body}}")
    table = doc.add_table(rows=1, cols=1)
    table.cell(0, 0).paragraphs[0].text = "{{cell}}"
    buf = io.BytesIO()
    doc.save(buf)

    _post(client, "/admin/docx/draft", data={"name": "rpt"},
          files={"file": ("rpt.docx", buf.getvalue(), "application/octet-stream")})
    fields = dict(FIELDS, arg_name=["body", "cell"], arg_type=["string", "string"],
                  arg_required=["true", "true"], arg_default=["", ""],
                  arg_desc=["", ""])
    _post(client, "/admin/docx/save", data=fields)
    form = _post(client, "/admin/docx/preview/values", data=fields).text

    markdown = "- one\n- two"
    r = _post(client, "/admin/docx/preview", data={
        "spec_json": _carried(form),
        "value_body": markdown, "value_cell": markdown})

    out = Document(io.BytesIO(r.content))
    body_styles = [p.style.name for p in out.paragraphs if p.text.strip()]
    cell_styles = [p.style.name for p in out.tables[0].cell(0, 0).paragraphs
                   if p.text.strip()]
    cell_text = "\n".join(p.text for p in out.tables[0].cell(0, 0).paragraphs)
    assert "List Bullet" in body_styles, \
        "a body placeholder must render block-level Markdown"
    assert "one" in cell_text, \
        "the cell must actually be rendered, or the next assertion is vacuous"
    assert "List Bullet" not in cell_styles, \
        "a table cell is inline-only, and now you can see that"


def test_a_template_of_only_flags_can_have_them_all_unticked(admin_client):
    """The narrowest case, and the one #168 exists for.

    An unticked checkbox sends nothing at all. A template whose arguments are
    all booleans therefore submits *no* `value_` key when every box is off —
    and "no values" used to mean "generate samples", so the admin's explicit
    off came back as a sample on. The form carries a marker so that the
    submission is recognised whatever the controls happen to send.
    """
    client, _custom = admin_client
    _post(client, "/admin/docx/draft", data={"name": "flagonly"},
          files={"file": ("flagonly.docx",
                          _template_bytes("{{#if extras}}", "OPTIONAL BLOCK",
                                          "{{/if}}", "Always here"),
                          "application/octet-stream")})
    fields = dict(FIELDS, name="flagonly", original_name="flagonly",
                  asset_filename="flagonly.docx",
                  arg_name=["extras"], arg_type=["bool"],
                  arg_required=["true"], arg_default=[""], arg_desc=[""])
    _post(client, "/admin/docx/save", data=fields)
    form = _post(client, "/admin/docx/preview/values", data=fields).text

    from admin.preview import VALUES_MARKER
    assert f'name="{VALUES_MARKER}"' in form, "the form must mark its submission"

    # Exactly what a browser sends with the only box unticked.
    r = _post(client, "/admin/docx/preview",
              data={"spec_json": _carried(form), VALUES_MARKER: "1"})

    texts = [t for _s, t, _b in _paragraphs(r)]
    assert any("Always here" in t for t in texts), "the document must render"
    assert not any("OPTIONAL BLOCK" in t for t in texts), \
        "an unticked flag must stay off even when it is the only argument"


def test_the_marker_alone_is_enough_to_mean_values_were_submitted():
    from admin.preview import VALUES_MARKER

    assert has_submitted_values({VALUES_MARKER: "1"}) is True
    assert has_submitted_values({"name": "rpt"}) is False
