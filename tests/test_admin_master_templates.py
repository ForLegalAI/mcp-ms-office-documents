"""Inspecting and adopting hand-written master-YAML templates (#167).

A template defined in `config/<kind>_templates.yaml` was a bare row: the name,
a dash, "● Live", and the badge "from master YAML — read-only". That was the
entire affordance — no description, no arguments, no style mapping, no way to
see which file it used — even though `gather_specs()` had returned the whole
spec to build that row. For anyone whose templates predate the admin UI, that
was most of their templates, permanently second-class.

Adopting copies the entry into the managed `*.d` layer, where the merge order
makes it win. The master file is never modified: that is the property worth
guarding, because it is someone's hand-written, heavily commented config.
"""
import io
import re
import sys
from pathlib import Path

project_root = Path(__file__).parent.parent
sys.path.insert(0, str(project_root))

import pytest
import yaml
from docx import Document
from starlette.testclient import TestClient

import admin.store as store_mod
import metrics
import template_utils as tu
from config import Config


@pytest.fixture
def admin_client(tmp_path, monkeypatch):
    custom = tmp_path / "custom"
    cfg = tmp_path / "config"
    custom.mkdir()
    cfg.mkdir()
    monkeypatch.setattr(store_mod, "_APP_CUSTOM_DIR", tmp_path / "nx1")
    monkeypatch.setattr(store_mod, "_APP_CONFIG_DIR", tmp_path / "nx2")
    monkeypatch.setattr(store_mod, "_LOCAL_CUSTOM_DIR", custom)
    monkeypatch.setattr(store_mod, "_LOCAL_CONFIG_DIR", cfg)
    monkeypatch.setattr(tu, "APP_CUSTOM_DIR", tmp_path / "nx1")
    monkeypatch.setattr(tu, "LOCAL_CUSTOM_DIR", custom)
    import pptx_tools.templates as templates_mod
    monkeypatch.setattr(templates_mod, "_APP_CONFIG_DIR", tmp_path / "nx2")
    monkeypatch.setattr(templates_mod, "_LOCAL_CONFIG_DIR", cfg)
    templates_mod.clear_cache()
    monkeypatch.setenv("ADMIN_ENABLED", "true")
    monkeypatch.setenv("ADMIN_PASSWORD", "pw")
    monkeypatch.delenv("API_KEY", raising=False)

    from fastmcp import FastMCP
    from admin.app import build_combined_app

    mcp = FastMCP("test-master")
    client = TestClient(build_combined_app(mcp, Config.from_env()))
    metrics.reset()
    client.__enter__()
    client.post("/admin/login", data={"password": "pw"})
    yield client, mcp, custom, cfg
    client.__exit__(None, None, None)
    metrics.reset()


def _docx_bytes(text="Dear {{who}}"):
    doc = Document()
    doc.add_paragraph(text)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _csrf(client) -> str:
    html = client.get("/admin/new/docx").text
    tag = re.search(r'<input[^>]*name="csrf"[^>]*>', html)
    return re.search(r'value="([^"]*)"', tag.group(0)).group(1)


def _post(client, url, data=None, files=None):
    payload = dict(data or {})
    payload["csrf"] = _csrf(client)
    return client.post(url, data=payload, files=files)


def _store():
    return store_mod.FileTemplateStore.from_config()


#: A master entry rich enough that losing any part of it would show.
MASTER_SPEC = {
    "name": "legacy_letter",
    "description": "Predates the admin UI",
    "docx_path": "legacy_letter.docx",
    "args": [{"name": "who", "type": "string", "required": True,
              "description": "Who it is addressed to"}],
    "style_mapping": {"heading_1": "My Heading 1"},
}


def _write_master(cfg: Path, *specs, kind="docx"):
    path = cfg / f"{kind}_templates.yaml"
    path.write_text(yaml.safe_dump({"templates": list(specs)}), encoding="utf-8")
    return path


def _row(html, name):
    m = re.search(r"<tr>((?:(?!</tr>).)*" + re.escape(name) + r".*?)</tr>", html, re.S)
    return " ".join(re.sub(r"<[^>]+>", " ", m.group(1)).split()) if m else ""


# ---------------------------------------------------------------------------
# Inspect
# ---------------------------------------------------------------------------


def test_a_master_template_links_to_its_detail_page(admin_client):
    client, _mcp, custom, cfg = admin_client
    _write_master(cfg, MASTER_SPEC)
    (custom / "legacy_letter.docx").write_bytes(_docx_bytes())

    html = client.get("/admin/word/templates").text
    assert "/admin/docx/legacy_letter/master" in html
    assert "Inspect" in _row(html, "legacy_letter")


def _above_the_yaml(html: str) -> str:
    """The page without its "as written" YAML dump.

    The dump serialises the whole spec, so asserting that a description or a
    style name appears *somewhere* on the page passes even when the cards
    meant to present it are gone. Everything worth checking has to be checked
    above it.
    """
    marker = "The YAML this is stored as"
    assert marker in html, "the page should still end with the YAML block"
    return html.split(marker)[0]


def test_the_detail_page_shows_what_the_row_never_could(admin_client):
    client, _mcp, custom, cfg = admin_client
    _write_master(cfg, MASTER_SPEC)
    (custom / "legacy_letter.docx").write_bytes(_docx_bytes())

    html = _above_the_yaml(client.get("/admin/docx/legacy_letter/master").text)

    assert "Predates the admin UI" in html, "the description"
    assert "Who it is addressed to" in html, "the argument's description"
    assert "legacy_letter.docx" in html, "the source file"
    assert "Style mapping" in html and "My Heading 1" in html, "the style mapping"
    assert "What we found in the document" in html, "the analysis"


def test_the_detail_page_offers_no_way_to_edit_the_spec(admin_client):
    """Read-only: the master YAML is hand-written and tooling never rewrites it."""
    client, _mcp, custom, cfg = admin_client
    _write_master(cfg, MASTER_SPEC)
    (custom / "legacy_letter.docx").write_bytes(_docx_bytes())

    html = client.get("/admin/docx/legacy_letter/master").text

    assert "/admin/docx/save" not in html, "no save form"
    assert 'name="arg_name"' not in html, "no argument editor"


def test_the_detail_page_says_where_the_file_lives(admin_client):
    """A master entry usually points at a file that ships with the server."""
    client, _mcp, _custom, cfg = admin_client
    _write_master(cfg, dict(MASTER_SPEC, docx_path="default_docx_template.docx"))

    html = _above_the_yaml(client.get("/admin/docx/legacy_letter/master").text)
    assert "default_templates/" in html


def test_a_disabled_master_template_is_still_listed(admin_client):
    """Listing from the YAML, not from the live tool names.

    Keyed off the live tools, a master template that is disabled or that
    failed to register simply vanished from the page — with no way to inspect
    it and no way to adopt it.
    """
    client, _mcp, custom, cfg = admin_client
    _write_master(cfg, dict(MASTER_SPEC, enabled=False))
    (custom / "legacy_letter.docx").write_bytes(_docx_bytes())

    row = _row(client.get("/admin/word/templates").text, "legacy_letter")
    assert row, "a disabled master template must still appear"
    assert "Disabled" in row
    assert client.get("/admin/docx/legacy_letter/master").status_code == 200


def test_an_unknown_master_template_is_not_found(admin_client):
    client, _mcp, _custom, _cfg = admin_client
    r = client.get("/admin/docx/ghost/master")
    assert "not found" in r.text.lower()


def test_a_managed_template_is_not_listed_twice(admin_client):
    """Once adopted, it is a managed template and nothing else."""
    client, _mcp, custom, cfg = admin_client
    _write_master(cfg, MASTER_SPEC)
    (custom / "legacy_letter.docx").write_bytes(_docx_bytes())
    _post(client, "/admin/docx/legacy_letter/adopt", data={})

    html = client.get("/admin/word/templates").text
    assert html.count("/admin/docx/legacy_letter/master") == 0, \
        "the master row must go once the managed spec overrides it"
    assert "/admin/docx/legacy_letter/edit" in html


# ---------------------------------------------------------------------------
# Adopt
# ---------------------------------------------------------------------------


def test_adopting_copies_the_spec_into_the_managed_layer(admin_client):
    client, _mcp, custom, cfg = admin_client
    _write_master(cfg, MASTER_SPEC)
    (custom / "legacy_letter.docx").write_bytes(_docx_bytes())

    r = _post(client, "/admin/docx/legacy_letter/adopt", data={})

    adopted = _store().get_spec("docx", "legacy_letter")
    assert adopted is not None
    assert adopted["description"] == "Predates the admin UI"
    assert [a["description"] for a in adopted["args"]] == ["Who it is addressed to"]
    assert adopted["style_mapping"] == {"heading_1": "My Heading 1"}
    assert "Edit legacy_letter" in r.text, "it should land on the edit page"


def test_adopting_never_modifies_the_master_yaml(admin_client):
    """It is someone's hand-written, commented config. Adopting is additive."""
    client, _mcp, custom, cfg = admin_client
    master = _write_master(cfg, MASTER_SPEC)
    (custom / "legacy_letter.docx").write_bytes(_docx_bytes())
    before = master.read_text()

    _post(client, "/admin/docx/legacy_letter/adopt", data={})

    assert master.read_text() == before, "the master file must be byte-identical"


def test_the_adopted_copy_wins_over_the_master_entry(admin_client):
    """The merge order is what makes leaving the master file alone safe."""
    from template_registry import gather_specs

    client, _mcp, custom, cfg = admin_client
    master = _write_master(cfg, MASTER_SPEC)
    (custom / "legacy_letter.docx").write_bytes(_docx_bytes())
    _post(client, "/admin/docx/legacy_letter/adopt", data={})

    _post(client, "/admin/docx/save", data={
        "kind": "docx", "asset_filename": "legacy_letter.docx",
        "name": "legacy_letter", "original_name": "legacy_letter",
        "title": "T", "description": "edited here",
        "arg_name": ["who"], "arg_type": ["string"], "arg_required": ["true"],
        "arg_default": [""], "arg_desc": ["Who it is addressed to"],
    })

    merged, _ = gather_specs(master, cfg / "docx_templates.d")
    entries = [s for s in merged if s.get("name") == "legacy_letter"]
    assert len(entries) == 1, "the two must merge, not both appear"
    assert entries[0]["description"] == "edited here"


def test_adopting_copies_a_shipped_file_into_the_writable_directory(admin_client):
    """Otherwise you could edit the template but never replace its document.

    A master entry usually points at a file in `default_templates/`, which is
    read-only as far as the admin UI is concerned. `broadcast_email_style_1`
    ships there and is not one of the base-template names.
    """
    client, _mcp, custom, cfg = admin_client
    _write_master(cfg, {"name": "broadcast", "description": "d",
                        "html_path": "broadcast_email_style_1.html", "args": []},
                  kind="email")
    assert not (custom / "broadcast_email_style_1.html").exists()

    _post(client, "/admin/email/broadcast/adopt", data={})

    assert (custom / "broadcast_email_style_1.html").exists(), \
        "the document must land where Replace can overwrite it"
    assert _store().get_spec("email", "broadcast")["html_path"] == \
        "broadcast_email_style_1.html"


def test_adopting_never_takes_over_a_base_template_filename(admin_client):
    """The dangerous case: a master entry pointing at a base-template file.

    `template_utils` searches custom_templates/ first, so copying
    `default_docx_template.docx` in there shadows the base Word template —
    and then replacing this one template's document silently restyles every
    Word document the server generates. The adopted template gets a private
    copy under its own name instead.
    """
    client, _mcp, custom, cfg = admin_client
    before = tu.find_docx_template()
    _write_master(cfg, dict(MASTER_SPEC, docx_path="default_docx_template.docx"))

    _post(client, "/admin/docx/legacy_letter/adopt", data={})

    assert not (custom / "default_docx_template.docx").exists(), \
        "a base-template filename must never appear in the uploads directory"
    assert _store().get_spec("docx", "legacy_letter")["docx_path"] == \
        "legacy_letter.docx"
    assert (custom / "legacy_letter.docx").exists()
    assert tu.find_docx_template() == before, \
        "the base Word template must still resolve to the shipped file"


def test_replacing_an_adopted_documents_file_cannot_restyle_everything(admin_client):
    """The consequence the rename above exists to prevent, end to end."""
    client, _mcp, custom, cfg = admin_client
    _write_master(cfg, dict(MASTER_SPEC, docx_path="default_docx_template.docx"))
    _post(client, "/admin/docx/legacy_letter/adopt", data={})

    _post(client, "/admin/docx/legacy_letter/reupload",
          files={"file": ("x.docx", _docx_bytes("ONLY FOR THIS TEMPLATE"),
                          "application/octet-stream")})

    base = tu.find_docx_template()
    assert "custom" not in str(Path(base).parent), \
        "the base Word template must not have been replaced"
    assert "ONLY FOR THIS TEMPLATE" not in \
        Document(str(base)).paragraphs[0].text


def test_adopting_leaves_the_admins_own_file_alone(admin_client):
    """Adoption must never replace a file the admin already put there.

    Note this cannot distinguish the `asset_exists` guard from its absence:
    `_candidate_dirs()` searches custom_templates/ first, so "recopying"
    resolves to the same file and rewrites it byte-for-byte. The guard is an
    I/O optimisation; this test guards the property that matters, and would
    catch a change that resolved from the shipped defaults first.
    """
    client, _mcp, custom, cfg = admin_client
    _write_master(cfg, MASTER_SPEC)
    (custom / "legacy_letter.docx").write_bytes(_docx_bytes("MINE {{who}}"))

    _post(client, "/admin/docx/legacy_letter/adopt", data={})

    assert "MINE" in Document(str(custom / "legacy_letter.docx")).paragraphs[0].text


def test_adopting_an_already_managed_template_is_refused(admin_client):
    client, _mcp, custom, cfg = admin_client
    _write_master(cfg, MASTER_SPEC)
    (custom / "legacy_letter.docx").write_bytes(_docx_bytes())
    _post(client, "/admin/docx/legacy_letter/adopt", data={})

    r = _post(client, "/admin/docx/legacy_letter/adopt", data={})
    assert "already managed" in r.text


def test_adopting_a_template_whose_file_is_missing_is_refused(admin_client):
    client, _mcp, _custom, cfg = admin_client
    _write_master(cfg, dict(MASTER_SPEC, docx_path="nowhere.docx"))

    r = _post(client, "/admin/docx/legacy_letter/adopt", data={})

    assert "Cannot find" in r.text
    assert _store().get_spec("docx", "legacy_letter") is None


def test_adopting_needs_the_csrf_token(admin_client):
    client, _mcp, custom, cfg = admin_client
    _write_master(cfg, MASTER_SPEC)
    (custom / "legacy_letter.docx").write_bytes(_docx_bytes())

    client.post("/admin/docx/legacy_letter/adopt", data={})
    assert _store().get_spec("docx", "legacy_letter") is None


def test_adopting_is_a_post_not_a_link(admin_client):
    client, _mcp, custom, cfg = admin_client
    _write_master(cfg, MASTER_SPEC)
    (custom / "legacy_letter.docx").write_bytes(_docx_bytes())

    r = client.get("/admin/docx/legacy_letter/adopt", follow_redirects=False)
    assert r.status_code in (404, 405)
    assert _store().get_spec("docx", "legacy_letter") is None


def test_adopting_a_disabled_master_template_keeps_it_disabled(admin_client):
    client, _mcp, custom, cfg = admin_client
    _write_master(cfg, dict(MASTER_SPEC, enabled=False))
    (custom / "legacy_letter.docx").write_bytes(_docx_bytes())

    _post(client, "/admin/docx/legacy_letter/adopt", data={})

    assert _store().get_spec("docx", "legacy_letter")["enabled"] is False


# ---------------------------------------------------------------------------
# What adoption does *not* promise: the copy is keyed by name, and that is all
# ---------------------------------------------------------------------------


def test_disabling_the_master_entry_after_adopting_has_no_effect(admin_client):
    """The override replaces the whole entry, so the master's flag is ignored.

    Pinned as a decision rather than left to be discovered: once adopted, the
    managed copy is the template, and the master entry of that name is dead
    config. Editing it — including turning it off — changes nothing.
    """
    from template_registry import gather_specs

    client, _mcp, custom, cfg = admin_client
    master = _write_master(cfg, MASTER_SPEC)
    (custom / "legacy_letter.docx").write_bytes(_docx_bytes())
    _post(client, "/admin/docx/legacy_letter/adopt", data={})

    _write_master(cfg, dict(MASTER_SPEC, enabled=False))

    merged, _ = gather_specs(master, cfg / "docx_templates.d")
    assert [s["name"] for s in merged] == ["legacy_letter"], \
        "the adopted copy still wins, and it is still enabled"


def test_renaming_the_master_entry_after_adopting_leaves_both(admin_client):
    """A name-keyed merge cannot follow a rename it was never told about.

    Hand-renaming an adopted entry in the master YAML gives two templates:
    the renamed master entry, no longer overridden, and the adopted copy,
    now standing alone. Inherent to keying the merge on the name — adoption
    records no provenance — but surprising enough to pin so that a change in
    this behaviour is deliberate.
    """
    from template_registry import gather_specs

    client, _mcp, custom, cfg = admin_client
    master = _write_master(cfg, MASTER_SPEC)
    (custom / "legacy_letter.docx").write_bytes(_docx_bytes())
    _post(client, "/admin/docx/legacy_letter/adopt", data={})

    _write_master(cfg, dict(MASTER_SPEC, name="legacy_letter_v2"))

    merged, _ = gather_specs(master, cfg / "docx_templates.d")
    assert sorted(s["name"] for s in merged) == ["legacy_letter", "legacy_letter_v2"]


def test_the_yaml_block_does_not_claim_the_ui_wrote_the_file(admin_client):
    """The one page whose point is that the file belongs to the admin."""
    client, _mcp, custom, cfg = admin_client
    _write_master(cfg, MASTER_SPEC)
    (custom / "legacy_letter.docx").write_bytes(_docx_bytes())

    html = client.get("/admin/docx/legacy_letter/master").text
    assert "what the UI wrote" not in html
    assert "written by you" in html


def test_a_managed_templates_yaml_block_still_says_the_ui_wrote_it(admin_client):
    """And the wording must not flip for templates the UI really did write."""
    client, _mcp, custom, cfg = admin_client
    _write_master(cfg, MASTER_SPEC)
    (custom / "legacy_letter.docx").write_bytes(_docx_bytes())
    _post(client, "/admin/docx/legacy_letter/adopt", data={})

    html = client.get("/admin/docx/legacy_letter/edit").text
    assert "what the UI wrote" in html


# ---------------------------------------------------------------------------
# The Overview tab counts it (#194 follow-up)
# ---------------------------------------------------------------------------


def _install_and_register(mcp, custom, cfg, spec=None):
    """Write a master-YAML template and register it, as startup would.

    The admin app does not register dynamic tools — ``main.py`` does that once
    at startup — so a test that only writes the YAML has a template the
    registry has never heard of, and "live" is legitimately 0. The deployed
    server has it registered, which is the shape this bug appeared in, so the
    test has to do what startup does or it is measuring something else.
    """
    from docx_tools.dynamic_docx_tools import register_docx_template_tools_from_yaml

    master = _write_master(cfg, spec or MASTER_SPEC)
    (custom / "legacy_letter.docx").write_bytes(_docx_bytes())
    register_docx_template_tools_from_yaml(mcp, master)
    return master


def test_the_overview_counts_a_master_yaml_template(admin_client):
    """A hand-written template is as live as a managed one, and says so.

    Found on the deployed instance, which has exactly this shape: one Word
    template declared in config/docx_templates.yaml and none managed by the
    UI. Three pages then disagreed about the same fact — the Templates tab
    showed it Live, Server > Status counted 1 live Word tool, and the Word
    Overview said 0 with the tool missing from the table headed "Tools the AI
    can call". The counts came from `store.list_specs()` alone, which is the
    managed half only; `template_table()` walks both, which is why it was the
    one page that was right.
    """
    client, mcp, custom, cfg = admin_client
    _install_and_register(mcp, custom, cfg)

    overview = client.get("/admin/word").text
    assert "legacy_letter" in overview, (
        "a live tool is missing from the table of tools the AI can call")
    assert "From master YAML" in overview, (
        "it is live but not editable here; the table should say which it is")

    # The dashboard tile counts it too — it reads from the same record. Match
    # the tile, not the nav link of the same href, which carries no counts.
    tile = re.search(r'<a href="/admin/word" class="tile">.*?</a>',
                     client.get("/admin/").text, re.S)
    assert tile, "no Word tile on the dashboard"
    assert "<b>1</b><span>live</span>" in re.sub(r"\s+", "", tile.group(0)), (
        "the Word tile must count the master-YAML template as live")


def test_the_overview_and_the_server_status_agree_on_live_word_tools(admin_client):
    """The two counts are of the same thing and must not contradict.

    Server > Status reads the registry through `live_names()`, the Overview
    reads the specs. One counting a template the other does not is the shape
    of the bug, whichever way round it happens.
    """
    client, mcp, custom, cfg = admin_client
    _install_and_register(mcp, custom, cfg)

    status = client.get("/admin/server/status").text
    on_status = re.search(
        r'<div class="num">(\d+)</div>\s*<div class="lbl">Live Word tools</div>',
        status)
    assert on_status, "the Status tab no longer reports live Word tools"
    assert on_status.group(1) == "1", "the fixture did not register the tool"

    # +1: create_word_document is always live and is not a template.
    overview = client.get("/admin/word").text
    live_rows = len(re.findall(r"badge badge-live", overview))
    assert live_rows == int(on_status.group(1)) + 1, (
        "Overview and Status disagree about how many Word tools are live")


def test_a_disabled_master_template_is_counted_as_disabled(admin_client):
    """`enabled: false` in the master YAML is a real state, not an absence."""
    client, mcp, custom, cfg = admin_client
    _install_and_register(mcp, custom, cfg, dict(MASTER_SPEC, enabled=False))

    overview = client.get("/admin/word").text
    assert "legacy_letter" in overview
    # One template, not live, so the Templates card's Disabled count is 1.
    assert re.search(r'<div class="num warn-text">1</div>\s*'
                     r'<div class="lbl">Disabled</div>', overview), (
        "a disabled master-YAML template must show as disabled, not missing")
